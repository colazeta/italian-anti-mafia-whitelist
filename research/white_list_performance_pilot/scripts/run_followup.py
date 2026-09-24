#!/usr/bin/env python3
"""Research checkpoint v0.3: no canonical writes, no public row-level output.

The new Word adapter is deliberately research-only. Its structural diagnostics
are exported and it is not registered as a production parser. The existing PDF
parser remains unchanged. Raw evidence stays on the ephemeral runner.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import os
import re
import subprocess
import tempfile
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

import run_source_pilot as base

VERSION = '0.3.0-source-diagnostic'
DOCX_VERSION = 'research-cosenza-seven-column-docx-v1'


def suffix(url, content_type, data):
    if data.startswith(b'%PDF'):
        return '.pdf'
    ext = Path(urlparse(url).path).suffix.lower()
    if ext in {'.xlsx', '.xls', '.docx', '.doc', '.html', '.json'}:
        return ext
    by_type = {
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': '.xlsx',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document': '.docx',
        'text/html': '.html', 'application/json': '.json',
    }
    return by_type.get(content_type, '.bin')


def extra_metrics(records):
    result = base.analyse(records)
    all_pending = []
    dated_pending = 0
    for r in records:
        if r['status'] != 'pending':
            continue
        if r['app'] and r['reference'] and r['app'] <= r['reference']:
            dated_pending += 1
            all_pending.append((r['reference'] - r['app']).days)
    result['observed_pending_age_all_source_rows'] = base.quantiles(all_pending)
    result['sensitivity_interpretation'] = (
        'All-source-row age is a literal published-row diagnostic, not distinct procedures. '
        'The strict linked subset additionally excludes ambiguous or multiple identifiers. '
        'Differences measure sensitivity to research eligibility, not source error.'
    )
    result['pending_age_histogram_all_source_rows'] = {
        '0_to_90': sum(x <= 90 for x in all_pending),
        '91_to_180': sum(90 < x <= 180 for x in all_pending),
        '181_to_365': sum(180 < x <= 365 for x in all_pending),
        '366_to_730': sum(365 < x <= 730 for x in all_pending),
        '731_to_1825': sum(730 < x <= 1825 for x in all_pending),
        'over_1825': sum(x > 1825 for x in all_pending),
    }
    assert sum(result['pending_age_histogram_all_source_rows'].values()) == dated_pending
    return result


def current_source(cfg, work, parsers):
    report = {k: cfg.get(k) for k in ('source_key', 'authority_key', 'register_key',
                                    'reference_date', 'population_scope', 'parser', 'resource_url')}
    report['reference_date_basis'] = 'frozen project source configuration; document-date validation separate'
    report['expected_sha256'] = cfg.get('sha256')
    try:
        data, capture = base.fetch(cfg['resource_url'])
        report.update(capture)
        if not base.HEX.fullmatch(str(cfg.get('sha256') or '')) or capture['sha256'] != cfg['sha256']:
            report['status'] = 'blocked_unapproved_bytes'
            return report, []
        ext = suffix(capture['final_url'], capture['content_type'], data)
        path = work / (cfg['source_key'] + ext)
        path.write_bytes(data)
        report['local_format'] = ext
        if cfg['parser'] == 'cosenza_combined_v2':
            records = base.cosenza_records(path)
        else:
            batch = parsers[cfg['parser']](path, cfg)
            records = batch.records
        if len(records) != cfg['expected_source_rows']:
            raise ValueError('Approved source-row count mismatch')
        adapted = [base.normalise(r, cfg) for r in records]
        report['status'] = 'approved_parser_executed'
        report['metrics'] = extra_metrics(adapted)
        return report, adapted
    except Exception as exc:
        report['status'] = 'capture_or_parse_failed'
        report['error_class'] = type(exc).__name__
        return report, []


def is_header(cells):
    if len(cells) != 7:
        return False
    c = [' '.join(x.upper().split()) for x in cells]
    return ('RAGIONE' in c[0] and 'SEDE' in c[1] and 'SEDE' in c[2]
            and ('FISCALE' in c[3] or 'IVA' in c[3])
            and ('ATTIVIT' in c[4] or 'ISCRIZIONE' in c[4])
            and ('ISTANZA' in c[5] or 'PRESENTAZIONE' in c[5]) and 'ESITO' in c[6])


def parse_cosenza_docx(path):
    from docx import Document
    from white_list_archive.parsers.cosenza_combined_v2 import _parse_dates, _parse_identifiers, _parse_outcome
    doc = Document(path)
    records = []
    diagnostics = {'adapter_version': DOCX_VERSION, 'table_count': len(doc.tables),
                   'matched_header_tables': 0, 'header_rows': 0, 'blank_rows': 0,
                   'unparsed_nonempty_rows': 0, 'body_cell_counts': {},
                   'row_locators_retained_in_memory_only': True}
    widths = Counter()
    for table_index, table in enumerate(doc.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        headers = [i for i, cells in enumerate(rows[:5]) if is_header(cells)]
        if not headers:
            diagnostics['unparsed_nonempty_rows'] += sum(any(x for x in cells) for cells in rows)
            continue
        diagnostics['matched_header_tables'] += 1
        for row_index, cells in enumerate(rows):
            if is_header(cells):
                diagnostics['header_rows'] += 1
                continue
            if not any(cells):
                diagnostics['blank_rows'] += 1
                continue
            widths[len(cells)] += 1
            if row_index < headers[0] or len(cells) != 7:
                diagnostics['unparsed_nonempty_rows'] += 1
                continue
            name, office, secondary, identifier, activities, application, outcome_raw = cells
            if not name:
                diagnostics['unparsed_nonempty_rows'] += 1
                continue
            _, ids = _parse_identifiers(identifier.splitlines())
            outcome = _parse_outcome(outcome_raw)
            records.append({'operator_name_raw': name, 'identifiers_json': json.dumps(ids),
                            'application_dates_json': json.dumps(_parse_dates(application)),
                            'observed_listing_date': outcome.get('observed_listing_date') or '',
                            'source_status': outcome['status'],
                            'source_locator': f'table:{table_index + 1}/row:{row_index + 1}'})
    diagnostics['body_cell_counts'] = dict(widths)
    diagnostics['parsed_rows'] = len(records)
    if not diagnostics['matched_header_tables'] or not records:
        raise ValueError('No explicitly matched seven-column source table')
    return records, diagnostics


def history_source(edition, work):
    from white_list_archive.acquisition.cosenza_capture import resolve_combined_resource
    report = {'edition_key': edition['edition_key'], 'reference_date': edition['reference_date'],
              'page_url': edition['edition_page_url'], 'source_origin': edition['source_origin']}
    try:
        url, _ = resolve_combined_resource(edition['edition_page_url'])
        data, capture = base.fetch(url)
        report.update(capture)
        report['resource_url'] = url
        ext = suffix(capture['final_url'], capture['content_type'], data)
        path = work / (edition['edition_key'] + ext)
        path.write_bytes(data)
        report['local_format'] = ext
        if ext == '.pdf':
            text = subprocess.run(['pdftotext', '-f', '1', '-l', '1', '-layout', str(path), '-'],
                                  check=True, capture_output=True, text=True).stdout
            folded = ' '.join(text.upper().split())
            headers = ['RAGIONE', 'SEDE', 'ISTANZA', 'ESITO']
            report['required_header_tokens_found'] = sum(x in folded for x in headers)
            if not all(x in folded for x in headers):
                report['status'] = 'captured_schema_requires_review'
                return report, []
            records = base.cosenza_records(path)
            report['adapter_version'] = 'existing-cosenza-pdf-v2'
        elif ext == '.docx':
            records, diagnostics = parse_cosenza_docx(path)
            report['structural_diagnostics'] = diagnostics
            report['adapter_version'] = DOCX_VERSION
        else:
            report['status'] = 'captured_format_requires_adapter'
            return report, []
        adapted = [base.normalise(r, {'source_key': 'cosenza-combined',
                                     'reference_date': edition['reference_date']}) for r in records]
        report['status'] = 'exploratory_parse_NOT_canonicalised'
        report['metrics'] = extra_metrics(adapted)
        report['validation_limit'] = 'Structural checks and parse execution are not an independent visual review or a completed canonical ingestion'
        return report, adapted
    except Exception as exc:
        report['status'] = 'capture_or_parse_failed'
        report['error_class'] = type(exc).__name__
        return report, []


def linked_lag_diagnostics(editions):
    """Source-linked durations; no survival estimator or true clearance inference."""
    appearances = defaultdict(list)
    for ref, records in sorted(editions):
        selected, _ = base.safe_records(records)
        for row in selected:
            if row['app'] and row['app'] <= row['reference']:
                appearances[(row['link_key'], row['app'])].append(row)
    lag_values, observed_widths, exclusions = [], [], Counter()
    persistent_pending = 0
    last_date = max(base.parse_date(d) for d, _ in editions) if editions else None
    cohorts = defaultdict(Counter)
    for key, observations in appearances.items():
        observations = sorted(observations, key=lambda x: x['reference'])
        if observations[-1]['reference'] == last_date and observations[-1]['status'] == 'pending':
            persistent_pending += sum(r['status'] == 'pending' for r in observations) >= 2
        pending = [r for r in observations if r['status'] == 'pending']
        listed = [r for r in observations if r['status'] == 'listed']
        if not pending or not listed:
            continue
        first_listed = min(r['reference'] for r in listed)
        earlier = [r for r in pending if r['reference'] < first_listed]
        if not earlier:
            exclusions['no_earlier_pending_observation'] += 1
            continue
        first_completion = next(r for r in listed if r['reference'] == first_listed)
        app = first_completion['app']
        listing = first_completion['listing']
        last_pending = max(r['reference'] for r in earlier)
        if not listing or listing < app or listing > first_listed:
            exclusions['missing_or_invalid_listing_date'] += 1
            continue
        if listing <= last_pending:
            exclusions['listing_at_or_before_last_published_pending'] += 1
            continue
        if any(r['status'] in {'renewal_requested', 'renewal_update_in_progress'}
               and r['reference'] <= first_listed for r in observations):
            exclusions['renewal_annotation_before_first_listing'] += 1
            continue
        lag_values.append((listing - app).days)
        observed_widths.append((first_listed - last_pending).days)
        cohort = f'{app.year}-Q{(app.month - 1) // 3 + 1}'
        cohorts[cohort]['observed_linked_listing_events'] += 1
    return {'unit': 'unique conservative entity-name-identifier/application-date source linkage',
            'qualifying_linked_application_to_listing_lag': base.quantiles(lag_values),
            'publication_observation_gap_days': base.quantiles(observed_widths),
            'exclusions': dict(exclusions),
            'last_edition_pending_seen_pending_in_multiple_editions': persistent_pending,
            'event_cohorts_NOT_complete_application_cohorts': dict(cohorts),
            'interpretation': 'Requires a prior published pending observation, consistent unique application date, later listing date within the observed gap, and no prior renewal annotation. Still conditional on observed linked successful cases; not an unbiased overall processing-time estimate.'}


def main():
    cli = argparse.ArgumentParser()
    cli.add_argument('--output-dir', type=Path, required=True)
    args = cli.parse_args()
    config = base.ROOT / 'data/publication/multi_prefecture_pilot.json'
    history_path = base.ROOT / 'data/source_registry/cosenza_historical_editions.csv'
    sources = [c for c in json.loads(config.read_text())['sources'] if c['authority_key'] in base.PILOT]
    history = list(csv.DictReader(history_path.open()))
    result = {'analysis_version': VERSION, 'code_commit': os.environ.get('GITHUB_SHA'),
              'started_at': datetime.now(timezone.utc).isoformat(),
              'input_sha256': {str(p.relative_to(base.ROOT)): base.digest(p.read_bytes()) for p in (config, history_path)},
              'current_sources': [], 'cosenza_history': [], 'canonical_database_writes': 0,
              'production_files_modified': 0}
    with tempfile.TemporaryDirectory(prefix='wl-research-followup-') as temp:
        work = Path(temp)
        parsers = base.parser_registry()
        with ThreadPoolExecutor(max_workers=2) as pool:
            outputs = list(pool.map(lambda cfg: current_source(cfg, work, parsers), sources))
        result['current_sources'] = [report for report, _ in outputs]
        editions = []
        for entry in history:
            report, rows = history_source(entry, work)
            result['cosenza_history'].append(report)
            if rows:
                editions.append((entry['reference_date'], rows))
        result['source_transitions'] = base.transition_diagnostics(sorted(editions))
        result['linked_lag_diagnostics'] = linked_lag_diagnostics(editions)
    result['run_summary'] = {
        'current_sources_requested': len(sources),
        'current_sources_parsed': sum(x['status'] == 'approved_parser_executed' for x in result['current_sources']),
        'historical_editions_requested': len(history),
        'historical_editions_exploratory_parsed': len(editions),
        'current_source_observations': sum(x.get('metrics', {}).get('source_observations', 0) for x in result['current_sources']),
        'historical_source_observations': sum(x.get('metrics', {}).get('source_observations', 0) for x in result['cosenza_history'])}
    result['completed_at'] = datetime.now(timezone.utc).isoformat()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / 'source_diagnostics_v3.json').write_text(json.dumps(result, indent=2) + '\n')
    print('RESEARCH_FOLLOWUP_SUMMARY=' + json.dumps(result['run_summary']))

if __name__ == '__main__':
    main()
