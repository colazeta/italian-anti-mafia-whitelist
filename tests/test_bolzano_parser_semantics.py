from pathlib import Path

from docx import Document

import white_list_archive.parsers.bolzano_docx as bz


LISTED_HEADER = [
    "Ragione Sociale",
    "Sede Legale",
    "Sede Secondaria con rappresentanza stabile in Italia",
    "Codice fiscale / Partita Iva",
    "Data di iscrizione",
    "Data validità iscrizione",
    "Aggiornamento in corso SI-JA",
]
APPLICANT_HEADER = [
    "Ragione Sociale",
    "Sede Legale",
    "Sede Secondaria con rappresentanza stabile in Italia",
    "Codice fiscale / Partita Iva",
    "Attività per cui è richiesta l'iscrizione",
    "Data di presentazione",
    "Esito",
]


def _cfg(source_key: str, scope: str, sha256: str) -> dict[str, str]:
    return {
        "source_key": source_key,
        "authority_key": "bolzano-bozen",
        "authority_name": "Prefettura di Bolzano",
        "register_key": "bolzano-bozen-ordinary",
        "register_name": "White List ordinaria",
        "population_scope": scope,
        "reference_date": "2026-09-11",
        "source_page_url": "https://example.invalid/bolzano",
        "resource_url": "https://example.invalid/source.docx",
        "sha256": sha256,
    }


def _put(cells, values):
    for cell, value in zip(cells, values):
        cell.text = value


def _listed_fixture(path: Path, rows_by_section: dict[int, list[list[str]]]) -> None:
    document = Document()
    for section in range(1, 11):
        table = document.add_table(rows=2, cols=7)
        _put(table.rows[0].cells, LISTED_HEADER)
        _put(table.rows[1].cells, [f"Sezione {section}", "", "", "", "", "", ""])
        for values in rows_by_section[section]:
            _put(table.add_row().cells, values)
    document.save(path)


def _applicant_fixture(path: Path, rows: list[list[str]]) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=7)
    _put(table.rows[0].cells, APPLICANT_HEADER)
    _put(table.rows[1].cells, ["", "", "", "", "", "", ""])
    for values in rows:
        _put(table.add_row().cells, values)
    document.save(path)


def test_bolzano_listed_exact_grouping_raw_dates_and_identifiers(tmp_path: Path, monkeypatch):
    path = tmp_path / "listed.docx"
    rows_by_section: dict[int, list[list[str]]] = {}
    for section in range(1, 11):
        listing = "29/6/2026" if section == 3 else "01/01/2026"
        identifier = "ATU 31173208" if section == 4 else f"000000000{section:02d}"
        rows_by_section[section] = [
            [f"IMPRESA {section} SRL", "Bolzano", "", identifier, listing, "01/01/2027", ""]
        ]
    repeated = ["RIPETUTA SRL", "Merano", "", "01234567890", "02/02/2026", "02/02/2027", ""]
    repeated_update = repeated.copy()
    repeated_update[-1] = "Si-JA"
    rows_by_section[1].append(repeated)
    rows_by_section[2].append(repeated_update)
    _listed_fixture(path, rows_by_section)

    monkeypatch.setattr(bz, "_EXPECTED_SECTION_ROWS", {1: 2, 2: 2, **{section: 1 for section in range(3, 11)}})
    monkeypatch.setattr(bz, "_EXPECTED_LISTED_SECTOR_ROWS", 12)
    monkeypatch.setattr(bz, "_EXPECTED_LISTED_RECORDS", 11)
    monkeypatch.setattr(bz, "_EXPECTED_LISTED_STATUS_COUNTS", {"listed": 10, "renewal_update_in_progress": 1})
    monkeypatch.setattr(bz, "_EXPECTED_LISTED_UPDATE_VALUES", {"": 11, "Si-JA": 1})

    batch = bz.parse_bolzano_listed(
        path, _cfg("bolzano-bozen-listed", "listed", bz._LISTED_SHA256)
    )
    assert batch.diagnostics["sector_rows"] == 12
    assert batch.diagnostics["public_records"] == 11
    assert batch.diagnostics["status_counts"] == {"listed": 10, "renewal_update_in_progress": 1}

    repeated_record = next(record for record in batch.records if record["name"] == "RIPETUTA SRL")
    assert repeated_record["source_status"] == "renewal_update_in_progress"
    assert repeated_record["requested_activities"] == ["Sezione 1", "Sezione 2"]

    malformed_date = next(record for record in batch.records if record["name"] == "IMPRESA 3 SRL")
    assert malformed_date["observed_listing_date"] == ""
    assert malformed_date["source_fields"]["listing_date_raw"] == "29/6/2026"

    foreign_identifier = next(record for record in batch.records if record["name"] == "IMPRESA 4 SRL")
    assert foreign_identifier["identifier_field_raw"] == "ATU 31173208"
    assert foreign_identifier["identifiers"] == []


def test_bolzano_applicants_stay_pending_and_preserve_malformed_date(tmp_path: Path, monkeypatch):
    path = tmp_path / "applicants.docx"
    _applicant_fixture(
        path,
        [
            ["RICHIEDENTE UNO SRL", "Bolzano", "", "01234567890", "Sezione 1", "01/09/2026", ""],
            ["RICHIEDENTE DUE SRL", "Merano", "", "C.F. ITALIANO (REA) 91066300210", "Sezione 2", "14/032025", ""],
        ],
    )
    monkeypatch.setattr(bz, "_EXPECTED_APPLICANT_ROWS", 2)
    monkeypatch.setattr(bz, "_EXPECTED_APPLICANT_RECORDS", 2)
    monkeypatch.setattr(bz, "_EXPECTED_APPLICANT_STATUS_COUNTS", {"pending": 2})

    batch = bz.parse_bolzano_applicants(
        path, _cfg("bolzano-bozen-applicants", "applicant", bz._APPLICANT_SHA256)
    )
    assert batch.diagnostics["public_records"] == 2
    assert batch.diagnostics["status_counts"] == {"pending": 2}
    assert all(record["source_status"] == "pending" for record in batch.records)

    malformed = next(record for record in batch.records if record["name"] == "RICHIEDENTE DUE SRL")
    assert malformed["application_date"] == ""
    assert malformed["source_fields"]["application_date_raw"] == "14/032025"
    assert malformed["identifier_field_raw"] == "C.F. ITALIANO (REA) 91066300210"
    assert malformed["identifiers"] == []


def test_bolzano_unknown_update_fails_closed(tmp_path: Path, monkeypatch):
    path = tmp_path / "listed.docx"
    rows_by_section = {
        section: [[f"IMPRESA {section}", "Bolzano", "", f"000000000{section:02d}", "01/01/2026", "01/01/2027", ""]]
        for section in range(1, 11)
    }
    rows_by_section[1][0][-1] = "STATUS NON AUDITATO"
    _listed_fixture(path, rows_by_section)

    try:
        bz.parse_bolzano_listed(path, _cfg("bolzano-bozen-listed", "listed", bz._LISTED_SHA256))
    except RuntimeError as exc:
        assert "unreviewed update status" in str(exc)
    else:
        raise AssertionError("Unknown Bolzano update status must fail closed")


def test_bolzano_cfg_digest_drift_fails_closed(tmp_path: Path):
    path = tmp_path / "applicants.docx"
    _applicant_fixture(
        path,
        [["RICHIEDENTE", "Bolzano", "", "01234567890", "Sezione 1", "01/09/2026", ""]],
    )
    try:
        bz.parse_bolzano_applicants(
            path, _cfg("bolzano-bozen-applicants", "applicant", "f" * 64)
        )
    except RuntimeError as exc:
        assert "approved-byte digest drift" in str(exc)
    else:
        raise AssertionError("Bolzano parser must not accept an unapproved source digest")
