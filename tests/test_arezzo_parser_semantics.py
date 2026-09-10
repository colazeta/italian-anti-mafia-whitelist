from datetime import date
from pathlib import Path

import openpyxl
from docx import Document

from white_list_archive.parsers.arezzo_openxml import parse_arezzo_applicants, parse_arezzo_listed


def _cfg(parser: str, source_key: str, population_scope: str) -> dict[str, str]:
    return {
        "parser": parser,
        "source_key": source_key,
        "authority_key": "arezzo",
        "authority_name": "Prefettura di Arezzo",
        "register_key": "arezzo-ordinary",
        "register_name": "White List ordinaria",
        "population_scope": population_scope,
        "reference_date": "2026-09-10",
        "source_page_url": "https://example.invalid/arezzo",
        "resource_url": "https://example.invalid/source",
        "sha256": "0" * 64,
    }


def _listed_sheet(workbook, title: str, section: str, rows: list[list[object]]):
    worksheet = workbook.create_sheet(title)
    worksheet.append(["PREFETTURA DI AREZZO"])
    worksheet.append([])
    worksheet.append([])
    worksheet.append([section])
    worksheet.append(
        [
            "Denominazione",
            "Sede legale",
            "Sede secondaria",
            "Codice fiscale/Partita I.V.A.",
            "Data iscrizione",
            "Data scadenza",
            "Aggiornamento in corso",
        ]
    )
    for row in rows:
        worksheet.append(row)


def test_arezzo_listed_groups_only_exact_sector_repetitions_and_preserves_conflicts(tmp_path: Path):
    path = tmp_path / "listed.xlsx"
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    _listed_sheet(
        workbook,
        "SEZ. I",
        "Sezione I – Estrazione, fornitura e trasporto di terra e materiali inerti",
        [
            ["IMPRESA UNO SRL", "Arezzo", "", "01234567890", date(2026, 1, 2), date(2027, 1, 1), ""],
            ["IMPRESA CONFLITTO SRL", "Arezzo", "", "12345678901", date(2025, 5, 1), date(2027, 5, 1), ""],
            ["IMPRESA MALFORMATA SRL", "Bucine", "", "02731820797", "0702/2025", date(2027, 3, 7), ""],
        ],
    )
    _listed_sheet(
        workbook,
        "SEZ. II",
        "Sezione II – Confezionamento, fornitura e trasporto di calcestruzzo e di bitume",
        [
            ["IMPRESA UNO SRL", "Arezzo", "", "01234567890", date(2026, 1, 2), date(2027, 1, 1), ""],
            ["IMPRESA CONFLITTO SRL", "Arezzo", "", "12345678901", date(2025, 6, 1), date(2027, 5, 1), ""],
            ["IMPRESA UPDATE SRL", "Arezzo", "", "SHORT-ID", date(2025, 7, 1), date(2026, 7, 1), "L’impresa ha comunicato il proprio interesse a permanere nell’elenco"],
        ],
    )
    workbook.save(path)

    batch = parse_arezzo_listed(path, _cfg("arezzo_listed", "arezzo-listed", "listed"))
    assert batch.diagnostics["sector_rows"] == 6
    assert batch.diagnostics["public_records"] == 5
    assert batch.diagnostics["dropped_date_rows"] == 0
    assert batch.diagnostics["malformed_date_rows_preserved"] == 1
    assert batch.diagnostics["date_conflict_identity_groups"] == 1

    one = [record for record in batch.records if record["name"] == "IMPRESA UNO SRL"]
    assert len(one) == 1
    assert len(one[0]["requested_activities"]) == 2

    conflict = [record for record in batch.records if record["name"] == "IMPRESA CONFLITTO SRL"]
    assert len(conflict) == 2
    assert all("listing_date" in record["source_fields"]["date_conflict_fields"] for record in conflict)

    malformed = next(record for record in batch.records if record["name"] == "IMPRESA MALFORMATA SRL")
    assert malformed["observed_listing_date"] == ""
    assert malformed["observed_expiry_date"] == "2027-03-07"
    assert malformed["source_fields"]["malformed_date_pairs"] == ["0702/2025 | 2027-03-07 00:00:00"]

    update = next(record for record in batch.records if record["name"] == "IMPRESA UPDATE SRL")
    assert update["source_status"] == "renewal_update_in_progress"
    assert update["identifier_field_raw"] == "SHORT-ID"
    assert update["identifiers"] == []


def test_arezzo_applicants_normalises_separator_typography_without_changing_digits(tmp_path: Path):
    path = tmp_path / "applicants.docx"
    document = Document()
    table = document.add_table(rows=1, cols=6)
    for cell, value in zip(
        table.rows[0].cells,
        ["Ragione Sociale", "Sede legale", "Sede secondaria", "Codice fiscale/Partita I.V.A.", "Attività", "Data di presentazione"],
    ):
        cell.text = value
    row = table.add_row().cells
    values = [
        "RICHIEDENTE SRL",
        "Arezzo",
        "",
        "1234",
        "Noli a freddo di macchinari Noli a caldo",
        "03//06/2026",
    ]
    for cell, value in zip(row, values):
        cell.text = value
    document.save(path)

    batch = parse_arezzo_applicants(path, _cfg("arezzo_applicants", "arezzo-applicants", "applicant"))
    assert batch.diagnostics["public_records"] == 1
    assert batch.diagnostics["dropped_date_rows"] == 0
    assert batch.diagnostics["separator_typography_normalised"] == 1
    record = batch.records[0]
    assert record["application_date"] == "2026-06-03"
    assert record["primary_date"] == "2026-06-03"
    assert record["source_status"] == "pending"
    assert record["outcome_raw"] == ""
    assert record["identifier_field_raw"] == "1234"
    assert record["identifiers"] == []
    assert record["requested_activities"] == ["Noli a freddo di macchinari", "Noli a caldo"]
    assert record["source_fields"]["malformed_date_pairs"] == ["application_date=03//06/2026"]


def test_arezzo_applicant_activity_vocabulary_fails_closed(tmp_path: Path):
    path = tmp_path / "unknown.docx"
    document = Document()
    table = document.add_table(rows=1, cols=6)
    for cell, value in zip(
        table.rows[0].cells,
        ["Ragione Sociale", "Sede legale", "Sede secondaria", "Codice fiscale", "Attività", "Data di presentazione"],
    ):
        cell.text = value
    row = table.add_row().cells
    for cell, value in zip(row, ["X SRL", "Arezzo", "", "01234567890", "Nuova attività non auditata", "01/09/2026"]):
        cell.text = value
    document.save(path)

    try:
        parse_arezzo_applicants(path, _cfg("arezzo_applicants", "arezzo-applicants", "applicant"))
    except ValueError as exc:
        assert "not in the audited source vocabulary" in str(exc)
    else:
        raise AssertionError("Unknown source activity must fail closed")
