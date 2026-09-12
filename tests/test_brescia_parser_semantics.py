from datetime import date
from pathlib import Path

import openpyxl

import white_list_archive.parsers.brescia_openxml as bs


def _cfg(parser: str, source_key: str, scope: str) -> dict[str, str]:
    return {
        "parser": parser,
        "source_key": source_key,
        "authority_key": "brescia",
        "authority_name": "Prefettura di Brescia",
        "register_key": "brescia-ordinary",
        "register_name": "White List ordinaria",
        "population_scope": scope,
        "reference_date": "2026-09-10",
        "source_page_url": "https://example.invalid/brescia",
        "resource_url": "https://example.invalid/source.xlsx",
        "sha256": "0" * 64,
    }


def _workbook(*, applicant: bool = False) -> tuple[openpyxl.Workbook, openpyxl.worksheet.worksheet.Worksheet]:
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Foglio1"
    if applicant:
        workbook.create_sheet("Foglio2")
        workbook.create_sheet("Foglio3")
    else:
        workbook.create_sheet("Foglio3")
        workbook.create_sheet("Foglio2")
    return workbook, ws


def _listed_header() -> list[object]:
    return [
        None,
        None,
        "Ragione sociale",
        "Sede legale",
        "Codice fiscale Partita iva",
        "Data iscrizione",
        "Data scadenza",
        "Aggiornamento in corso",
    ]


def _append_section(ws, roman: str, rows: list[list[object]]) -> None:
    ws.append([None, f"Sezione{roman}" if roman == "III" else f"Sezione {roman}"])
    ws.append([None, bs._EXPECTED_SECTIONS[roman]])
    ws.append(_listed_header())
    for row in rows:
        ws.append([None, None] + row)


def test_brescia_listed_groups_exact_identity_across_office_variants_and_preserves_source_offices(tmp_path: Path, monkeypatch):
    path = tmp_path / "listed.xlsx"
    workbook, ws = _workbook()
    by_section: dict[str, list[list[object]]] = {}
    for index, roman in enumerate(bs._EXPECTED_SECTIONS, start=1):
        by_section[roman] = [[
            f"UNICA {roman} SRL",
            f"Brescia via {roman}",
            f"000000000{index:02d}",
            date(2026, 1, 1),
            date(2027, 1, 1),
            "",
        ]]

    repeated_a = ["RIPETUTA SRL", "Brescia via Uno", "01234567890", date(2026, 2, 1), date(2027, 2, 1), ""]
    repeated_b = [
        "RIPETUTA SRL",
        "Brescia, Via Uno n. 1",
        "01234567890",
        date(2026, 2, 1),
        date(2027, 2, 1),
        bs._STANDARD_UPDATE,
    ]
    by_section["I"].append(repeated_a)
    by_section["II"].append(repeated_b)

    # The production source has one reviewed missing-identifier column shift.
    by_section["IV"].append(["ADMG SRL", "BRESCIA VIA CEFALONIA N70", "17/09/2025", "2026-09-17", "", ""])

    for roman, rows in by_section.items():
        _append_section(ws, roman, rows)
    # Reproduce one audited source layout anomaly: the company name moves
    # from labelled column C to column A while the other semantic columns stay put.
    ws.cell(row=4, column=1).value = "UNICA I SRL"
    ws.cell(row=4, column=3).value = None
    workbook.save(path)

    section_counts = {roman: len(rows) for roman, rows in by_section.items()}
    monkeypatch.setattr(bs, "_EXPECTED_SECTION_ROW_COUNTS", section_counts)
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_SECTOR_ROWS", sum(section_counts.values()))
    expected_shift = [[4, "I", "UNICA I SRL", "Brescia via I", "00000000001", "2026-01-01 00:00:00", "2027-01-01 00:00:00", ""]]
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_NAME_SHIFT_ROWS", 1)
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_NAME_SHIFT_SHA256", bs._layout_sha256(expected_shift))
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_NON_COMPANY_NOISE_ROWS", 0)
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_NON_COMPANY_NOISE_SHA256", bs._layout_sha256([]))
    monkeypatch.setattr(bs, "_EXPECTED_PEER_RESOLVED_DATE_ROWS", 0)
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_RECORDS", 12)
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_STATUS_COUNTS", {"listed": 11, "renewal_update_in_progress": 1})

    batch = bs.parse_brescia_listed(path, _cfg("brescia_listed", "brescia-listed", "listed"))
    assert batch.diagnostics["public_records"] == 12
    assert batch.diagnostics["status_counts"] == {"listed": 11, "renewal_update_in_progress": 1}

    repeated = [record for record in batch.records if record["name"] == "RIPETUTA SRL"]
    assert len(repeated) == 1
    assert repeated[0]["source_status"] == "renewal_update_in_progress"
    assert repeated[0]["source_fields"]["registered_office_variants"] == ["Brescia via Uno", "Brescia, Via Uno n. 1"]
    assert len(repeated[0]["requested_activities"]) == 2

    admg = next(record for record in batch.records if record["name"] == "ADMG SRL")
    assert admg["identifier_field_raw"] == ""
    assert admg["identifiers"] == []
    assert admg["observed_listing_date"] == "2025-09-17"
    assert admg["observed_expiry_date"] == "2026-09-17"
    assert admg["source_fields"]["malformed_date_pairs"]


def test_brescia_applicant_exact_duplicate_collapses_but_blank_dates_remain_pending(tmp_path: Path, monkeypatch):
    path = tmp_path / "applicants.xlsx"
    workbook, ws = _workbook(applicant=True)
    ws.append([])
    ws.append([
        None,
        "RAGIONE SOCIALE",
        "SEDE",
        "CODICE FISCALE/PARTITA IVA",
        "ATTIVITA' PER CUI E' RICHIESTA L'ISCRIZIONE",
        "DATA DI PRESENTAZIONE DELL'ISTANZA",
    ])
    duplicated = [None, "DUPLICATA SRL", "Iseo", "03059170179", "Noli a caldo", None]
    ws.append(duplicated)
    ws.append(duplicated)
    ws.append([None, "DATATA SRL", "Brescia", "01234567890", "Guardiania dei cantieri", date(2026, 8, 1)])
    workbook.save(path)

    monkeypatch.setattr(bs, "_EXPECTED_APPLICANT_ROWS", 3)
    monkeypatch.setattr(bs, "_EXPECTED_APPLICANT_RECORDS", 2)
    monkeypatch.setattr(bs, "_EXPECTED_APPLICANT_STATUS_COUNTS", {"pending": 2})
    monkeypatch.setattr(bs, "_EXPECTED_APPLICANT_MISSING_DATES", 2)
    monkeypatch.setattr(bs, "_EXPECTED_APPLICANT_MALFORMED_DATES", 0)
    monkeypatch.setattr(bs, "_EXPECTED_APPLICANT_DUPLICATE_GROUPS", 1)

    batch = bs.parse_brescia_applicants(path, _cfg("brescia_applicants", "brescia-applicants", "applicant"))
    assert batch.diagnostics["source_rows"] == 3
    assert batch.diagnostics["public_records"] == 2
    assert all(record["source_status"] == "pending" for record in batch.records)
    blank = next(record for record in batch.records if record["name"] == "DUPLICATA SRL")
    assert blank["application_date"] == ""


def test_brescia_unknown_date_typography_and_update_fail_closed(tmp_path: Path, monkeypatch):
    try:
        bs._source_date("2026??09??10", reviewed_malformed=bs._REVIEWED_LISTED_MALFORMED_DATES, allow_blank=False)
    except RuntimeError as exc:
        assert "unreviewed date typography" in str(exc)
    else:
        raise AssertionError("Unknown Brescia date typography must fail closed")

    path = tmp_path / "listed.xlsx"
    workbook, ws = _workbook()
    by_section = {}
    for index, roman in enumerate(bs._EXPECTED_SECTIONS, start=1):
        update = "STATUS NON AUDITATO" if roman == "I" else ""
        by_section[roman] = [[f"IMPRESA {roman}", "Brescia", f"000000000{index:02d}", date(2026, 1, 1), date(2027, 1, 1), update]]
    # Retain the one reviewed structural-shift row expected by the parser contract.
    by_section["IV"].append(["ADMG SRL", "BRESCIA VIA CEFALONIA N70", "17/09/2025", "2026-09-17", "", ""])
    for roman, rows in by_section.items():
        _append_section(ws, roman, rows)
    workbook.save(path)
    monkeypatch.setattr(bs, "_EXPECTED_SECTION_ROW_COUNTS", {roman: len(rows) for roman, rows in by_section.items()})
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_SECTOR_ROWS", sum(len(rows) for rows in by_section.values()))
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_NAME_SHIFT_ROWS", 0)
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_NAME_SHIFT_SHA256", bs._layout_sha256([]))
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_NON_COMPANY_NOISE_ROWS", 0)
    monkeypatch.setattr(bs, "_EXPECTED_LISTED_NON_COMPANY_NOISE_SHA256", bs._layout_sha256([]))

    try:
        bs.parse_brescia_listed(path, _cfg("brescia_listed", "brescia-listed", "listed"))
    except RuntimeError as exc:
        assert "unreviewed update status" in str(exc)
    else:
        raise AssertionError("Unknown Brescia update marker must fail closed")
