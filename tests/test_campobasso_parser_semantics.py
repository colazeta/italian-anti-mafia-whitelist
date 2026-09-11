from datetime import date
from pathlib import Path

import openpyxl
from docx import Document

import white_list_archive.parsers.campobasso_openxml as cb


def _cfg(parser: str, source_key: str, scope: str) -> dict[str, str]:
    return {
        "parser": parser,
        "source_key": source_key,
        "authority_key": "campobasso",
        "authority_name": "Prefettura di Campobasso",
        "register_key": "campobasso-ordinary",
        "register_name": "White List ordinaria",
        "population_scope": scope,
        "reference_date": "2026-08-11",
        "source_page_url": "https://example.invalid/campobasso",
        "resource_url": "https://example.invalid/source",
        "sha256": "0" * 64,
    }


def _append_section(ws, roman: str, activity: str, rows: list[list[object]]) -> None:
    # The real Section X heading is longer than its compact catalogue label and
    # explicitly refers to waste-management and remediation activities. Keep the
    # synthetic fixture faithful to that audited structure rather than weakening
    # the production parser's Section X guard.
    if roman == "X" and activity == "Servizi ambientali":
        activity = "Servizi ambientali - gestione rifiuti e bonifica"
    ws.append([])
    ws.append([None, f"SEZIONE {roman}"])
    ws.append([None, activity])
    ws.append(
        [
            None,
            "RAGIONE SOCIALE",
            "SEDE LEGALE",
            "PROV.",
            "SEDE SECONDARIA CON RAPPRESENTANZA STABILE IN ITALIA",
            "CODICE FISCALE/PARTITA IVA",
            "DATA DI ISCRIZIONE",
            "DATA SCADENZA ISCRIZIONE",
            "AGGIORNAMENTO IN CORSO",
        ]
    )
    for row in rows:
        ws.append([None] + row)


def test_campobasso_listed_groups_province_placeholders_but_not_office_conflicts(tmp_path: Path, monkeypatch):
    path = tmp_path / "listed.xlsx"
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Foglio1"

    section_rows: dict[str, list[list[object]]] = {}
    for roman, activity in cb._EXPECTED_SECTIONS.items():
        section_rows[roman] = [
            [
                f"UNICA {roman} SRL",
                "Campobasso",
                "CB",
                "",
                f"000000000{int({'I':1,'II':2,'III':3,'IV':4,'V':5,'VI':6,'VII':7,'VIII':8,'IX':9,'X':10}[roman]):02d}",
                date(2026, 1, 1),
                date(2027, 1, 1),
                "",
            ]
        ]

    # Exact identity/date/address repeat across two sections: province-only source
    # variation is not allowed to manufacture a duplicate observation.
    repeated = ["RIPETUTA SRL", "Bojano", "CB", "", "01234567890", date(2026, 2, 1), date(2027, 2, 1), "SI"]
    repeated_scoped = repeated.copy()
    repeated_scoped[2] = "=="
    repeated_scoped[-1] = "SI (non richiesto per questa sezione)"
    section_rows["I"].append(repeated)
    section_rows["II"].append(repeated_scoped)

    # Same name/id/dates but a genuinely different registered office remains a
    # separate observation: no fuzzy identity collapse.
    section_rows["III"].append(["SEDE VARIANTE SRL", "Bojano", "CB", "", "12345678901", date(2026, 3, 1), date(2027, 3, 1), ""])
    section_rows["IV"].append(["SEDE VARIANTE SRL", "Campochiaro", "CB", "", "12345678901", date(2026, 3, 1), date(2027, 3, 1), ""])

    for roman, activity in cb._EXPECTED_SECTIONS.items():
        _append_section(ws, roman, activity, section_rows[roman])
    workbook.save(path)

    expected_rows = {roman: len(rows) for roman, rows in section_rows.items()}
    expected_total = sum(expected_rows.values())
    # Ten unique section fixtures + one grouped repeat + two office variants.
    monkeypatch.setattr(cb, "_EXPECTED_SECTION_ROW_COUNTS", expected_rows)
    monkeypatch.setattr(cb, "_EXPECTED_SECTION_ROWS", expected_total)
    monkeypatch.setattr(cb, "_EXPECTED_LISTED_RECORDS", 13)
    monkeypatch.setattr(cb, "_EXPECTED_LISTED_STATUS_COUNTS", {"listed": 12, "renewal_update_in_progress": 1})

    batch = cb.parse_campobasso_listed(path, _cfg("campobasso_listed", "campobasso-listed", "listed"))
    assert batch.diagnostics["sector_rows"] == expected_total
    assert batch.diagnostics["public_records"] == 13
    assert batch.diagnostics["status_counts"] == {"listed": 12, "renewal_update_in_progress": 1}

    repeat = [record for record in batch.records if record["name"] == "RIPETUTA SRL"]
    assert len(repeat) == 1
    assert repeat[0]["source_status"] == "renewal_update_in_progress"
    assert len(repeat[0]["requested_activities"]) == 2

    office_variants = [record for record in batch.records if record["name"] == "SEDE VARIANTE SRL"]
    assert len(office_variants) == 2
    assert {record["registered_office"] for record in office_variants} == {"Bojano", "Campochiaro"}


def test_campobasso_malformed_identifier_stays_raw(tmp_path: Path, monkeypatch):
    path = tmp_path / "listed.xlsx"
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Foglio1"
    rows_by_section = {}
    for roman, activity in cb._EXPECTED_SECTIONS.items():
        identifier = "880470703" if roman == "I" else f"000000000{int({'I':1,'II':2,'III':3,'IV':4,'V':5,'VI':6,'VII':7,'VIII':8,'IX':9,'X':10}[roman]):02d}"
        rows_by_section[roman] = [[f"IMPRESA {roman}", "Campobasso", "CB", "", identifier, date(2026, 1, 1), date(2027, 1, 1), ""]]
        _append_section(ws, roman, activity, rows_by_section[roman])
    workbook.save(path)
    monkeypatch.setattr(cb, "_EXPECTED_SECTION_ROW_COUNTS", {roman: 1 for roman in cb._EXPECTED_SECTIONS})
    monkeypatch.setattr(cb, "_EXPECTED_SECTION_ROWS", 10)
    monkeypatch.setattr(cb, "_EXPECTED_LISTED_RECORDS", 10)
    monkeypatch.setattr(cb, "_EXPECTED_LISTED_STATUS_COUNTS", {"listed": 10})

    batch = cb.parse_campobasso_listed(path, _cfg("campobasso_listed", "campobasso-listed", "listed"))
    malformed = next(record for record in batch.records if record["identifier_field_raw"] == "880470703")
    assert malformed["identifiers"] == []


def test_campobasso_applicant_repeated_header_and_pending_status(tmp_path: Path, monkeypatch):
    path = tmp_path / "applicants.docx"
    document = Document()
    table = document.add_table(rows=1, cols=6)
    header = ["Ragione sociale", "Sede legale", "Sede secondaria con rappresentanza stabile in Italia", "Codice fiscale/ Partita IVA", "Attività per cui è richiesta l’iscrizione", "Data di presentazione dell’istanza"]
    for cell, value in zip(table.rows[0].cells, header):
        cell.text = value
    for values in (
        ["RICHIEDENTE UNO SRL", "Campobasso", "", "01234567890", "Noli a caldo", "14 novembre 2025"],
        header,
        ["RICHIEDENTE DUE SRL", "Bojano", "", "880470703", "Guardiania dei cantieri", "07 agosto 2026"],
    ):
        row = table.add_row().cells
        for cell, value in zip(row, values):
            cell.text = value
    document.save(path)
    monkeypatch.setattr(cb, "_EXPECTED_APPLICANT_RECORDS", 2)
    monkeypatch.setattr(cb, "_EXPECTED_APPLICANT_STATUS_COUNTS", {"pending": 2})

    batch = cb.parse_campobasso_applicants(path, _cfg("campobasso_applicants", "campobasso-applicants", "applicant"))
    assert batch.diagnostics["public_records"] == 2
    assert batch.diagnostics["repeated_headers"] == 2
    assert all(record["source_status"] == "pending" for record in batch.records)
    malformed = next(record for record in batch.records if record["identifier_field_raw"] == "880470703")
    assert malformed["identifiers"] == []


def test_campobasso_unknown_update_status_fails_closed(tmp_path: Path, monkeypatch):
    path = tmp_path / "listed.xlsx"
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Foglio1"
    for roman, activity in cb._EXPECTED_SECTIONS.items():
        update = "STATUS NON AUDITATO" if roman == "I" else ""
        _append_section(ws, roman, activity, [[f"IMPRESA {roman}", "Campobasso", "CB", "", "01234567890", date(2026, 1, 1), date(2027, 1, 1), update]])
    workbook.save(path)
    monkeypatch.setattr(cb, "_EXPECTED_SECTION_ROW_COUNTS", {roman: 1 for roman in cb._EXPECTED_SECTIONS})
    monkeypatch.setattr(cb, "_EXPECTED_SECTION_ROWS", 10)

    try:
        cb.parse_campobasso_listed(path, _cfg("campobasso_listed", "campobasso-listed", "listed"))
    except RuntimeError as exc:
        assert "unreviewed update status" in str(exc)
    else:
        raise AssertionError("Unknown Campobasso update status must fail closed")
