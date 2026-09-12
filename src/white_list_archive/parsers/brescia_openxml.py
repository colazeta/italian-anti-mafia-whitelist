from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record

PARSER_VERSION = "1"
_REFERENCE_DATE = "2026-09-10"
_EXPECTED_LISTED_SHEETS = ["Foglio1", "Foglio3", "Foglio2"]
_EXPECTED_APPLICANT_SHEETS = ["Foglio1", "Foglio2", "Foglio3"]
_EXPECTED_SECTIONS = {
    "I": "Estrazione, fornitura e trasporto di terra e materiali inerti",
    "II": "Confezionamento, fornitura e trasporto di calcestruzzo e bitume",
    "III": "Noli a freddo di macchinari",
    "IV": "Fornitura di ferro lavorato",
    "V": "Noli a caldo",
    "VI": "Autotrasporto per conto di terzi",
    "VII": "Guardiania dei cantieri",
    "VIII": "Servizi Funerari e Cimiteriali",
    "IX": "Ristorazione, gestione delle mense e catering",
    "X": (
        "Servizi ambientali, comprese le attività di raccolta, di trasporto nazionale e transfrontaliero, "
        "anche per conto di terzi, di trattamento e di smaltimento dei rifiuti, nonché le attività di "
        "risanamento e di bonifica e gli altri servizi connessi alla gestione dei rifiuti"
    ),
}
_EXPECTED_SECTION_ROW_COUNTS = {
    "I": 416,
    "II": 218,
    "III": 567,
    "IV": 436,
    "V": 582,
    "VI": 457,
    "VII": 68,
    "VIII": 35,
    "IX": 79,
    "X": 450,
}
_EXPECTED_LISTED_SECTOR_ROWS = 3308
# Two malformed source date tokens have an exact clean repetition for the same
# name + raw identifier + opposite date in another statutory sector. Those two
# rows are associated with the unique clean date value from that same source.
# All other malformed dates remain raw and unnormalised.
_EXPECTED_PEER_RESOLVED_DATE_ROWS = 2
_EXPECTED_LISTED_RECORDS = 2032
_EXPECTED_LISTED_STATUS_COUNTS = {"listed": 1820, "renewal_update_in_progress": 212}
_EXPECTED_APPLICANT_ROWS = 1264
_EXPECTED_APPLICANT_RECORDS = 1263
_EXPECTED_APPLICANT_STATUS_COUNTS = {"pending": 1263}
_EXPECTED_APPLICANT_MISSING_DATES = 841
_EXPECTED_APPLICANT_MALFORMED_DATES = 1
_EXPECTED_APPLICANT_DUPLICATE_GROUPS = 1

_SECTION_RE = re.compile(r"^\s*sezione\s*(i{1,3}|iv|v|vi{0,3}|ix|x)\s*$", re.I)
_DATE_DMY = re.compile(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{4})$")
_STANDARD_UPDATE = "in fase di aggiornamento l'iscrizione resta valida anche oltre la scadenza, fino all'esito definitivo"
_REVIEWED_NEUTRAL_UPDATE = "`"
_REVIEWED_LISTED_MALFORMED_DATES = frozenset(
    {
        "30/010/2026",
        "29/08/203",
        "05.082027",
        "1706.2027",
        "14/07/205",
        "01.09,2027",
        "18.09.208",
        "14.05,2027",
        "03112025",
        "16/16/2024",
        "27/09/20222",
        "32.06.2026",
    }
)
_REVIEWED_APPLICANT_MALFORMED_DATES = frozenset({"25/092025"})


def _validate_cfg(cfg: dict[str, Any], source_key: str) -> None:
    if cfg.get("source_key") != source_key:
        raise RuntimeError(f"Brescia parser/source mismatch: {cfg.get('source_key')!r} != {source_key!r}")
    if cfg.get("authority_key") != "brescia":
        raise RuntimeError("Brescia parser bound to a non-Brescia authority")
    if cfg.get("reference_date") != _REFERENCE_DATE:
        raise RuntimeError(f"Brescia reference date drift: {cfg.get('reference_date')!r}")


def _source_date(value: Any, *, reviewed_malformed: frozenset[str], allow_blank: bool) -> tuple[str, str, bool]:
    if isinstance(value, datetime):
        raw = value.date().isoformat()
        return raw, raw, False
    if isinstance(value, date):
        raw = value.isoformat()
        return raw, raw, False
    raw = _clean(value)
    if not raw:
        if allow_blank:
            return "", "", False
        raise RuntimeError("Brescia required source date is blank")
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:T00:00:00| 00:00:00)?", raw):
        try:
            return date.fromisoformat(raw[:10]).isoformat(), raw, False
        except ValueError as exc:
            if raw in reviewed_malformed:
                return "", raw, True
            raise RuntimeError(f"Brescia invalid calendar date: {raw!r}") from exc
    match = _DATE_DMY.fullmatch(raw)
    if match:
        day, month, year = map(int, match.groups())
        try:
            return date(year, month, day).isoformat(), raw, False
        except ValueError as exc:
            if raw in reviewed_malformed:
                return "", raw, True
            raise RuntimeError(f"Brescia invalid calendar date: {raw!r}") from exc
    if raw in reviewed_malformed:
        return "", raw, True
    raise RuntimeError(f"Brescia unreviewed date typography: {raw!r}")


def _strict_identifiers(value: str) -> list[str]:
    """Extract only source-explicit exact identifiers; never pad, trim or infer digits."""
    raw = _clean(value).upper()
    found: list[str] = []

    def add(token: str) -> None:
        if ((token.isdigit() and len(token) == 11) or (len(token) == 16 and token.isalnum())) and token not in found:
            found.append(token)

    for token in re.split(r"[^A-Z0-9]+", raw):
        if token:
            add(token)
    # Some source cells attach an explicit CF/PI label directly to the digits.
    for match in re.finditer(r"(?:^|[^A-Z0-9])(?:CF|P\.?\s*I\.?|PI)\s*[:=-]?\s*(\d{11})(?!\d)", raw):
        add(match.group(1))
    return found


def _header_indexes(values: list[str], *, listed: bool) -> tuple[int, ...] | None:
    folded = [value.casefold() for value in values]
    if not any("ragione sociale" in value for value in folded):
        return None

    def first(predicate) -> int | None:
        return next((index for index, value in enumerate(folded) if predicate(value)), None)

    name = first(lambda value: "ragione sociale" in value)
    office = first(lambda value: "sede legale" in value) if listed else first(lambda value: value.strip() == "sede" or "sede legale" in value)
    identifier = first(lambda value: "codice fiscale" in value or "partita iva" in value)
    if listed:
        listing = first(lambda value: "data iscrizione" in value or "data di iscrizione" in value)
        expiry = first(lambda value: "data scadenza" in value or "data di scadenza" in value)
        update = first(lambda value: "aggiornamento" in value)
        indexes = (name, office, identifier, listing, expiry, update)
    else:
        activity = first(lambda value: "attivit" in value or "settore" in value)
        application = first(lambda value: "data" in value and any(token in value for token in ("istanza", "presentazione", "richiesta")))
        indexes = (name, office, identifier, activity, application)
    if any(index is None for index in indexes):
        raise RuntimeError(f"Brescia unreviewed header layout: {values!r}")
    return tuple(int(index) for index in indexes if index is not None)


def _section_from_row(values: list[str]) -> str | None:
    hits = []
    for value in values:
        match = _SECTION_RE.fullmatch(value)
        if match:
            hits.append(match.group(1).upper())
    if not hits:
        return None
    if len(set(hits)) != 1 or hits[0] not in _EXPECTED_SECTIONS:
        raise RuntimeError(f"Brescia ambiguous section marker: {hits!r}")
    return hits[0]


def _activity_matches(section: str, value: str) -> bool:
    return _clean(value).casefold() == _EXPECTED_SECTIONS[section].casefold()


def _listed_rows(path: Path) -> tuple[list[dict[str, Any]], Counter[str], Counter[str], int]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if workbook.sheetnames != _EXPECTED_LISTED_SHEETS:
        raise RuntimeError(f"Brescia listed worksheet drift: {workbook.sheetnames!r}")
    for sheet_name in _EXPECTED_LISTED_SHEETS[1:]:
        sheet = workbook[sheet_name]
        if any(_clean(cell.value) for row in sheet.iter_rows() for cell in row):
            raise RuntimeError(f"Brescia unexpected data in auxiliary listed worksheet {sheet_name!r}")

    worksheet = workbook[_EXPECTED_LISTED_SHEETS[0]]
    section: str | None = None
    activity = ""
    columns: tuple[int, ...] | None = None
    section_headers: Counter[str] = Counter()
    section_counts: Counter[str] = Counter()
    rows: list[dict[str, Any]] = []
    structural_shift_rows = 0

    for source_row, raw_row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        raw_values = list(raw_row)
        values = [_clean(value) for value in raw_values]
        current = _section_from_row(values)
        if current:
            section = current
            activity = ""
            columns = None
            continue

        header = _header_indexes(values, listed=True)
        if header:
            if not section:
                raise RuntimeError(f"Brescia listed header before section at row {source_row}")
            if not activity or not _activity_matches(section, activity):
                raise RuntimeError(f"Brescia section {section} activity heading changed before row {source_row}: {activity!r}")
            columns = header
            section_headers[section] += 1
            continue

        if section and columns is None:
            nonempty = [value for value in values if value]
            for candidate in sorted(nonempty, key=len, reverse=True):
                if _activity_matches(section, candidate):
                    activity = candidate
                    break
            continue
        if not section or columns is None or not any(values):
            continue

        name_i, office_i, id_i, listing_i, expiry_i, update_i = columns
        name = values[name_i] if name_i < len(values) else ""
        office = values[office_i] if office_i < len(values) else ""
        identifier_raw = values[id_i] if id_i < len(values) else ""
        update_raw = values[update_i] if update_i < len(values) else ""
        listing_value: Any = raw_values[listing_i] if listing_i < len(raw_values) else ""
        expiry_value: Any = raw_values[expiry_i] if expiry_i < len(raw_values) else ""

        if not name:
            raise RuntimeError(f"Brescia nonempty listed row without company name at row {source_row}: {values!r}")

        structural_shift = False
        # The audited 10 September source has one row (ADMG SRL) where the missing
        # identifier causes the two date cells to be shifted one column left. Keep
        # the identifier absent; move only the source-explicit date values back to
        # their labelled semantic fields. Do not import the identifier from another
        # population or external source.
        if (
            name == "ADMG SRL"
            and identifier_raw == "17/09/2025"
            and _clean(listing_value) == "2026-09-17"
            and not _clean(expiry_value)
        ):
            structural_shift = True
            structural_shift_rows += 1
            original_identifier_cell = identifier_raw
            identifier_raw = ""
            expiry_value = listing_value
            listing_value = original_identifier_cell

        if not _clean(listing_value) or not _clean(expiry_value):
            raise RuntimeError(f"Brescia unreviewed incomplete listed row {source_row}: {values!r}")

        listing_date, listing_raw, listing_malformed = _source_date(
            listing_value, reviewed_malformed=_REVIEWED_LISTED_MALFORMED_DATES, allow_blank=False
        )
        expiry_date, expiry_raw, expiry_malformed = _source_date(
            expiry_value, reviewed_malformed=_REVIEWED_LISTED_MALFORMED_DATES, allow_blank=False
        )
        update_folded = update_raw.casefold()
        if update_folded not in ("", _STANDARD_UPDATE.casefold(), _REVIEWED_NEUTRAL_UPDATE):
            raise RuntimeError(f"Brescia unreviewed update status at row {source_row}: {update_raw!r}")
        rows.append(
            {
                "source_row": source_row,
                "section": section,
                "activity": activity,
                "name": name,
                "office": office,
                "identifier_raw": identifier_raw,
                "listing_date": listing_date,
                "listing_raw": listing_raw,
                "listing_malformed": listing_malformed,
                "expiry_date": expiry_date,
                "expiry_raw": expiry_raw,
                "expiry_malformed": expiry_malformed,
                "update_raw": update_raw,
                "structural_shift": structural_shift,
            }
        )
        section_counts[section] += 1

    if section_headers != Counter({section: 1 for section in _EXPECTED_SECTIONS}):
        raise RuntimeError(f"Brescia listed header structure drift: {dict(section_headers)!r}")
    if dict(section_counts) != _EXPECTED_SECTION_ROW_COUNTS:
        raise RuntimeError(f"Brescia listed section-row drift: {dict(section_counts)!r}")
    if len(rows) != _EXPECTED_LISTED_SECTOR_ROWS:
        raise RuntimeError(f"Brescia listed source-row drift: {len(rows)} != {_EXPECTED_LISTED_SECTOR_ROWS}")
    if structural_shift_rows != 1:
        raise RuntimeError(f"Brescia reviewed structural-shift row drift: {structural_shift_rows} != 1")
    return rows, section_counts, section_headers, structural_shift_rows


def _resolve_peer_dates(rows: list[dict[str, Any]]) -> int:
    """Resolve only a malformed date with one unique clean same-source peer value."""
    by_identity: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        if row["identifier_raw"]:
            by_identity[(row["name"], row["identifier_raw"])].append(row)

    resolved = 0
    for row in rows:
        peers = by_identity.get((row["name"], row["identifier_raw"]), []) if row["identifier_raw"] else []
        if row["listing_malformed"] and row["expiry_date"]:
            candidates = {
                peer["listing_date"]
                for peer in peers
                if peer is not row and peer["expiry_date"] == row["expiry_date"] and peer["listing_date"]
            }
            if len(candidates) > 1:
                raise RuntimeError(f"Brescia ambiguous peer listing date for source row {row['source_row']}")
            if len(candidates) == 1:
                row["listing_date"] = next(iter(candidates))
                resolved += 1
        if row["expiry_malformed"] and row["listing_date"]:
            candidates = {
                peer["expiry_date"]
                for peer in peers
                if peer is not row and peer["listing_date"] == row["listing_date"] and peer["expiry_date"]
            }
            if len(candidates) > 1:
                raise RuntimeError(f"Brescia ambiguous peer expiry date for source row {row['source_row']}")
            if len(candidates) == 1:
                row["expiry_date"] = next(iter(candidates))
                resolved += 1
    if resolved != _EXPECTED_PEER_RESOLVED_DATE_ROWS:
        raise RuntimeError(
            f"Brescia peer-resolved date-row drift: {resolved} != {_EXPECTED_PEER_RESOLVED_DATE_ROWS}"
        )
    return resolved


def parse_brescia_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, "brescia-listed")
    rows, section_counts, _section_headers, structural_shift_rows = _listed_rows(path)
    peer_resolved = _resolve_peer_dates(rows)

    grouped: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in rows:
        # Cross-sector grouping is anchored on exact company name, raw identifier
        # and the reviewed date pair. Registered-office text is intentionally not
        # normalised or used to split otherwise exact repeated identity/date rows;
        # every distinct source office remains visible in provenance.
        key = (
            row["name"],
            row["identifier_raw"],
            row["listing_date"] or row["listing_raw"],
            row["expiry_date"] or row["expiry_raw"],
        )
        group = grouped.setdefault(
            key,
            {
                "row": row,
                "sections": [],
                "activities": [],
                "offices": [],
                "updates": [],
                "listing_raw": [],
                "expiry_raw": [],
                "normalised_listing": [],
                "normalised_expiry": [],
                "malformed_pairs": [],
            },
        )
        if row["section"] in group["sections"]:
            raise RuntimeError(f"Brescia duplicate semantic observation inside section {row['section']}: {key!r}")
        for field, value in (
            ("sections", row["section"]),
            ("activities", row["activity"]),
            ("offices", row["office"]),
            ("updates", row["update_raw"]),
            ("listing_raw", row["listing_raw"]),
            ("expiry_raw", row["expiry_raw"]),
            ("normalised_listing", row["listing_date"]),
            ("normalised_expiry", row["expiry_date"]),
        ):
            if value and value not in group[field]:
                group[field].append(value)
        if row["listing_malformed"] or row["expiry_malformed"] or row["structural_shift"]:
            pair = f"listing={row['listing_raw']} | expiry={row['expiry_raw']}"
            if row["structural_shift"]:
                pair += " | reviewed_source_column_shift=missing_identifier"
            if pair not in group["malformed_pairs"]:
                group["malformed_pairs"].append(pair)

    if len(grouped) != _EXPECTED_LISTED_RECORDS:
        raise RuntimeError(f"Brescia listed semantic-group drift: {len(grouped)} != {_EXPECTED_LISTED_RECORDS}")

    records: list[dict[str, Any]] = []
    for group in grouped.values():
        row = group["row"]
        has_positive_update = any(value.casefold() == _STANDARD_UPDATE.casefold() for value in group["updates"])
        status = "renewal_update_in_progress" if has_positive_update else "listed"
        record = _record(
            cfg,
            len(records) + 1,
            name=row["name"],
            office=row["office"],
            identifier_raw=row["identifier_raw"],
            activities=group["activities"],
            status=status,
            outcome_raw=" | ".join(group["updates"]),
            listing_date=row["listing_date"],
            expiry_date=row["expiry_date"],
            primary_date_label="Data iscrizione",
            source_fields={
                "sections": group["sections"],
                "registered_office_variants": group["offices"],
                "listing_date_raw_variants": group["listing_raw"],
                "expiry_date_raw_variants": group["expiry_raw"],
                "normalised_listing_date_variants": group["normalised_listing"],
                "normalised_expiry_date_variants": group["normalised_expiry"],
                "malformed_date_pairs": group["malformed_pairs"],
                "requested_activities_source": " | ".join(group["activities"]),
            },
        )
        record["identifiers"] = _strict_identifiers(row["identifier_raw"])
        records.append(record)

    status_counts = dict(Counter(record["source_status"] for record in records))
    if status_counts != _EXPECTED_LISTED_STATUS_COUNTS:
        raise RuntimeError(f"Brescia listed status drift: {status_counts!r}")

    return ParsedBatch(
        records=records,
        diagnostics={
            "parser": "brescia_listed",
            "sector_rows": len(rows),
            "public_records": len(records),
            "status_counts": status_counts,
            "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
            "section_row_counts": dict(section_counts),
            "peer_resolved_date_rows": peer_resolved,
            "reviewed_structural_shift_rows": structural_shift_rows,
            "malformed_source_date_rows": sum(row["listing_malformed"] or row["expiry_malformed"] for row in rows),
        },
    )


def parse_brescia_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, "brescia-applicants")
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if workbook.sheetnames != _EXPECTED_APPLICANT_SHEETS:
        raise RuntimeError(f"Brescia applicant worksheet drift: {workbook.sheetnames!r}")
    for sheet_name in _EXPECTED_APPLICANT_SHEETS[1:]:
        sheet = workbook[sheet_name]
        if any(_clean(cell.value) for row in sheet.iter_rows() for cell in row):
            raise RuntimeError(f"Brescia unexpected data in auxiliary applicant worksheet {sheet_name!r}")

    worksheet = workbook[_EXPECTED_APPLICANT_SHEETS[0]]
    columns: tuple[int, ...] | None = None
    header_count = 0
    rows: list[dict[str, Any]] = []
    for source_row, raw_row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        raw_values = list(raw_row)
        values = [_clean(value) for value in raw_values]
        header = _header_indexes(values, listed=False)
        if header:
            columns = header
            header_count += 1
            continue
        if columns is None or not any(values):
            continue
        name_i, office_i, id_i, activity_i, application_i = columns
        name = values[name_i] if name_i < len(values) else ""
        if not name:
            raise RuntimeError(f"Brescia nonempty applicant row without company name at row {source_row}: {values!r}")
        office = values[office_i] if office_i < len(values) else ""
        identifier_raw = values[id_i] if id_i < len(values) else ""
        activity = values[activity_i] if activity_i < len(values) else ""
        if not identifier_raw or not activity:
            raise RuntimeError(f"Brescia incomplete applicant identity/activity at row {source_row}: {values!r}")
        application_date, application_raw, malformed = _source_date(
            raw_values[application_i] if application_i < len(raw_values) else "",
            reviewed_malformed=_REVIEWED_APPLICANT_MALFORMED_DATES,
            allow_blank=True,
        )
        rows.append(
            {
                "source_row": source_row,
                "name": name,
                "office": office,
                "identifier_raw": identifier_raw,
                "activity": activity,
                "application_date": application_date,
                "application_raw": application_raw,
                "malformed": malformed,
            }
        )

    if header_count != 1:
        raise RuntimeError(f"Brescia applicant header drift: {header_count} != 1")
    if len(rows) != _EXPECTED_APPLICANT_ROWS:
        raise RuntimeError(f"Brescia applicant source-row drift: {len(rows)} != {_EXPECTED_APPLICANT_ROWS}")
    missing_dates = sum(not row["application_raw"] for row in rows)
    malformed_dates = sum(row["malformed"] for row in rows)
    if missing_dates != _EXPECTED_APPLICANT_MISSING_DATES:
        raise RuntimeError(f"Brescia applicant missing-date drift: {missing_dates} != {_EXPECTED_APPLICANT_MISSING_DATES}")
    if malformed_dates != _EXPECTED_APPLICANT_MALFORMED_DATES:
        raise RuntimeError(f"Brescia applicant malformed-date drift: {malformed_dates} != {_EXPECTED_APPLICANT_MALFORMED_DATES}")

    grouped: dict[tuple[str, ...], dict[str, Any]] = {}
    duplicate_groups = 0
    for row in rows:
        key = (
            row["name"], row["office"], row["identifier_raw"], row["activity"],
            row["application_date"] or row["application_raw"],
        )
        if key in grouped:
            duplicate_groups += 1
            group = grouped[key]
        else:
            group = grouped.setdefault(key, {"row": row, "application_raw": [], "malformed_pairs": []})
        if row["application_raw"] and row["application_raw"] not in group["application_raw"]:
            group["application_raw"].append(row["application_raw"])
        if row["malformed"]:
            item = f"application={row['application_raw']}"
            if item not in group["malformed_pairs"]:
                group["malformed_pairs"].append(item)

    if duplicate_groups != _EXPECTED_APPLICANT_DUPLICATE_GROUPS:
        raise RuntimeError(
            f"Brescia applicant duplicate-group drift: {duplicate_groups} != {_EXPECTED_APPLICANT_DUPLICATE_GROUPS}"
        )
    if len(grouped) != _EXPECTED_APPLICANT_RECORDS:
        raise RuntimeError(f"Brescia applicant semantic-group drift: {len(grouped)} != {_EXPECTED_APPLICANT_RECORDS}")

    records: list[dict[str, Any]] = []
    for group in grouped.values():
        row = group["row"]
        record = _record(
            cfg,
            len(records) + 1,
            name=row["name"],
            office=row["office"],
            identifier_raw=row["identifier_raw"],
            activities=[row["activity"]],
            status="pending",
            application_date=row["application_date"],
            primary_date_label="Data presentazione istanza" if row["application_date"] else "",
            source_fields={
                "application_date_raw_variants": group["application_raw"],
                "malformed_date_pairs": group["malformed_pairs"],
                "requested_activities_source": row["activity"],
            },
        )
        record["identifiers"] = _strict_identifiers(row["identifier_raw"])
        records.append(record)

    status_counts = dict(Counter(record["source_status"] for record in records))
    if status_counts != _EXPECTED_APPLICANT_STATUS_COUNTS:
        raise RuntimeError(f"Brescia applicant status drift: {status_counts!r}")
    return ParsedBatch(
        records=records,
        diagnostics={
            "parser": "brescia_applicants",
            "source_rows": len(rows),
            "public_records": len(records),
            "status_counts": status_counts,
            "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
            "missing_application_dates": missing_dates,
            "malformed_application_dates": malformed_dates,
            "duplicate_exact_groups": duplicate_groups,
        },
    )


PARSERS = {
    "brescia_listed": parse_brescia_listed,
    "brescia_applicants": parse_brescia_applicants,
}
