from __future__ import annotations

import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record

PARSER_VERSION = "1"
_REFERENCE_DATE = "2026-09-11"
_SHEET = "Foglio1"

_LISTED_FIRST_ROW = 3
_LISTED_LAST_ROW = 1634
_LISTED_RECORDS = 1632
_APPLICANT_FIRST_ROW = 5
_APPLICANT_LAST_ROW = 328
_APPLICANT_RECORDS = 324
_FOOTER_ROWS = 11

_EXPECTED_LISTED_STATUS_COUNTS = {"listed": 1281, "renewal_update_in_progress": 351}
_EXPECTED_APPLICANT_STATUS_COUNTS = {"pending": 324}
_EXPECTED_LISTED_IDENTIFIER_COVERAGE = 1588
_EXPECTED_APPLICANT_IDENTIFIER_COVERAGE = 322
_EXPECTED_LISTED_FORMULAS = 1491
_EXPECTED_LISTED_STANDARD_FORMULAS = 1490

_REVIEWED_LISTED_SECTOR_ANOMALIES = {
    350: "1, 5, 6 10",
    758: "1,2,34,,5,6,10",
    1083: "1,2,5,610,",
    1163: "1,2,3,4,5,69,10",
}
_REVIEWED_LISTED_MALFORMED_LISTING_DATES = {
    158: "16/072026",
    274: "09/062026",
}
_REVIEWED_LISTED_BLANK_EXPIRIES = {366, 400, 822}
_REVIEWED_LISTED_BLANK_OFFICES = {777}
_REVIEWED_LISTED_ABNORMAL_FORMULA = {
    1064: "=A1487=DATE(YEAR(F1064)+1,MONTH(F1064),DAY(F1064))"
}
_REVIEWED_APPLICANT_NONSTANDARD_IDENTIFIERS = {
    41: "061411110871",
    67: "CCOSVT731C351M",
}

_DATE_DMY = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")
_DATE_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?:[ T]00:00:00)?$")
_STANDARD_EXPIRY_FORMULA = re.compile(
    r"=DATE\(YEAR\(F(?P<row>\d+)\)\+1,MONTH\(F(?P=row)\),DAY\(F(?P=row)\)\)",
    re.I,
)


def _validate_cfg(cfg: dict[str, Any], expected_source_key: str) -> None:
    if cfg.get("source_key") != expected_source_key:
        raise RuntimeError(
            f"Catania parser/source mismatch: {cfg.get('source_key')!r} != {expected_source_key!r}"
        )
    if cfg.get("authority_key") != "catania":
        raise RuntimeError("Catania parser bound to a non-Catania authority")
    if cfg.get("reference_date") != _REFERENCE_DATE:
        raise RuntimeError(f"Catania reference date drift: {cfg.get('reference_date')!r}")


def _strict_date(value: Any, *, context: str) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _clean(value)
    match = _DATE_DMY.fullmatch(raw)
    if match:
        day, month, year = map(int, match.groups())
        try:
            return date(year, month, day).isoformat()
        except ValueError as exc:
            raise RuntimeError(f"Catania invalid calendar date in {context}: {raw!r}") from exc
    match = _DATE_ISO.fullmatch(raw)
    if match:
        year, month, day = map(int, match.groups())
        try:
            return date(year, month, day).isoformat()
        except ValueError as exc:
            raise RuntimeError(f"Catania invalid calendar date in {context}: {raw!r}") from exc
    raise RuntimeError(f"Catania unreviewed date typography in {context}: {raw!r}")


def _strict_identifiers(value: Any) -> list[str]:
    raw = _clean(value).replace(" ", "").upper()
    if (raw.isdigit() and len(raw) == 11) or (len(raw) == 16 and raw.isalnum()):
        return [raw]
    return []


def _sections(value: Any, *, source_row: int, population: str) -> list[str]:
    raw = _clean(value)
    if not raw:
        raise RuntimeError(f"Catania {population} blank sections at source row {source_row}")

    if population == "listed" and source_row in _REVIEWED_LISTED_SECTOR_ANOMALIES:
        expected = _REVIEWED_LISTED_SECTOR_ANOMALIES[source_row]
        if raw != expected:
            raise RuntimeError(
                f"Catania listed reviewed section anomaly drift at row {source_row}: {raw!r} != {expected!r}"
            )
        tokens = [token.strip() for token in re.split(r"[,.;]+", raw) if token.strip()]
        # Retain only source-explicit valid section tokens. Fused or otherwise
        # invalid tokens are never decomposed or repaired.
        valid = [token for token in tokens if token in {str(i) for i in range(1, 11)}]
    else:
        tokens = [token.strip() for token in re.split(r"[,.;]+", raw) if token.strip()]
        if not tokens or any(token not in {str(i) for i in range(1, 11)} for token in tokens):
            raise RuntimeError(
                f"Catania {population} unreviewed section value at source row {source_row}: {raw!r}"
            )
        valid = tokens

    if len(valid) != len(set(valid)):
        raise RuntimeError(f"Catania {population} duplicate section at source row {source_row}: {raw!r}")
    return [f"Sezione {token}" for token in valid]


def _footer_signature(sheet: Any, *, applicant: bool) -> list[tuple[str, str]]:
    start = sheet.max_row - (_FOOTER_ROWS - 1)
    result: list[tuple[str, str]] = []
    for row_number in range(start, sheet.max_row + 1):
        values = [_clean(cell.value) for cell in sheet[row_number]]
        if applicant:
            result.append((values[1] if len(values) > 1 else "", values[2] if len(values) > 2 else ""))
        else:
            result.append((values[0] if values else "", values[1] if len(values) > 1 else ""))
    return result


_EXPECTED_FOOTER_LABELS = [
    "SEZIONI WHITE LIST",
    "SEZ. I",
    "SEZ. II",
    "SEZ III",
    "SEZ IV",
    "SEZ IX",
    "SEZ V",
    "SEZ VI",
    "SEZ VII",
    "SEZ VIII",
    "SEZ X",
]


def _validate_structure(sheet: Any, *, population: str) -> None:
    applicant = population == "applicants"
    expected_max_row = 1806 if applicant else 3380
    expected_max_col = 9 if applicant else 8
    first = _APPLICANT_FIRST_ROW if applicant else _LISTED_FIRST_ROW
    last = _APPLICANT_LAST_ROW if applicant else _LISTED_LAST_ROW
    header = 4 if applicant else 2
    if sheet.max_row != expected_max_row or sheet.max_column != expected_max_col:
        raise RuntimeError(
            f"Catania {population} worksheet dimensions drift: "
            f"{sheet.max_row}x{sheet.max_column} != {expected_max_row}x{expected_max_col}"
        )

    header_values = [_clean(cell.value) for cell in sheet[header]]
    if applicant:
        required = {
            0: "RAGIONE SOCIALE",
            1: "SEDE LEGALE",
            2: "SEDE SECONDARIA CON RAPPRESENTANZA IN ITALIA",
            4: "CODICE FISCALE/PARTITA IVA",
            5: "DATA PRESENTAZIONE ISTANZA",
            6: "ESITO",
        }
    else:
        required = {
            0: "RAGIONE SOCIALE",
            1: "SEDE LEGALE",
            2: "SEDE SECONDARIA CON RAPPRESENTANZA IN ITALIA",
            4: "CODICE FISCALE/PARTITA IVA",
            5: "DATA DI ISCRIZIONE",
            6: "SCADENZA ISCRIZIONE",
            7: "AGGIORNAMENTO IN CORSO",
        }
    for index, expected in required.items():
        if header_values[index] != expected:
            raise RuntimeError(
                f"Catania {population} header drift at column {index + 1}: {header_values[index]!r} != {expected!r}"
            )
    if applicant and "ATTIVITA'" not in header_values[3]:
        raise RuntimeError(f"Catania applicant activity header drift: {header_values[3]!r}")

    for row_number in range(last + 1, sheet.max_row - _FOOTER_ROWS + 1):
        if any(_clean(cell.value) for cell in sheet[row_number]):
            raise RuntimeError(f"Catania {population} unexpected data after company block at row {row_number}")

    footer = _footer_signature(sheet, applicant=applicant)
    labels = [item[0] for item in footer]
    if labels != _EXPECTED_FOOTER_LABELS:
        raise RuntimeError(f"Catania {population} footer-label drift: {labels!r}")
    if first != header + 1:
        raise AssertionError("Internal Catania row-boundary definition error")


def _listed_status(raw: Any) -> str:
    # Presence in the official listed-company population is itself positive
    # evidence of listed status. Only source-explicit renewal wording elevates
    # the observation to renewal/update in progress; cryptic or misplaced notes
    # are preserved raw and never interpreted as renewal.
    return "renewal_update_in_progress" if "rinnovo" in _clean(raw).casefold() else "listed"


def parse_catania_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, "catania-listed")
    formulas_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    values_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if formulas_book.sheetnames != [_SHEET] or values_book.sheetnames != [_SHEET]:
        raise RuntimeError(f"Catania listed worksheet drift: {formulas_book.sheetnames!r}")
    formula_sheet = formulas_book[_SHEET]
    value_sheet = values_book[_SHEET]
    _validate_structure(formula_sheet, population="listed")

    status_values: Counter[str] = Counter()
    formula_count = 0
    standard_formula_count = 0
    blank_expiries: set[int] = set()
    blank_offices: set[int] = set()
    reviewed_malformed_dates: dict[int, str] = {}
    identifier_coverage = 0
    records: list[dict[str, Any]] = []

    for source_row in range(_LISTED_FIRST_ROW, _LISTED_LAST_ROW + 1):
        raw = [cell.value for cell in formula_sheet[source_row]][:8]
        cached = [cell.value for cell in value_sheet[source_row]][:8]
        if len(raw) != 8 or not any(_clean(value) for value in raw):
            raise RuntimeError(f"Catania listed company-row drift at source row {source_row}")
        name, office, secondary, sections_raw, identifier_raw, listing_raw, expiry_raw, update_raw = raw
        if not _clean(name):
            raise RuntimeError(f"Catania listed blank company name at source row {source_row}")
        if not _clean(office):
            blank_offices.add(source_row)
        sections = _sections(sections_raw, source_row=source_row, population="listed")
        identifiers = _strict_identifiers(identifier_raw)
        identifier_coverage += bool(identifiers)

        listing_source = _clean(listing_raw)
        if source_row in _REVIEWED_LISTED_MALFORMED_LISTING_DATES:
            expected = _REVIEWED_LISTED_MALFORMED_LISTING_DATES[source_row]
            if listing_source != expected:
                raise RuntimeError(f"Catania reviewed listed-date drift at row {source_row}: {listing_source!r}")
            listing_date = ""
            reviewed_malformed_dates[source_row] = listing_source
        else:
            listing_date = _strict_date(listing_raw, context=f"listed row {source_row} listing")

        expiry_source = _clean(expiry_raw)
        if not expiry_source:
            blank_expiries.add(source_row)
            expiry_date = ""
        elif isinstance(expiry_raw, str) and expiry_raw.startswith("="):
            formula_count += 1
            if source_row in _REVIEWED_LISTED_ABNORMAL_FORMULA:
                expected = _REVIEWED_LISTED_ABNORMAL_FORMULA[source_row]
                if expiry_raw != expected or cached[6] is not False:
                    raise RuntimeError(f"Catania reviewed abnormal expiry-formula drift at row {source_row}")
                expiry_date = ""
            else:
                match = _STANDARD_EXPIRY_FORMULA.fullmatch(expiry_raw)
                if not match or int(match.group("row")) != source_row:
                    raise RuntimeError(f"Catania unreviewed expiry formula at row {source_row}: {expiry_raw!r}")
                standard_formula_count += 1
                expiry_date = _strict_date(cached[6], context=f"listed row {source_row} cached expiry")
        else:
            expiry_date = _strict_date(expiry_raw, context=f"listed row {source_row} expiry")

        status_values[_clean(update_raw)] += 1
        status = _listed_status(update_raw)
        record = _record(
            cfg,
            len(records) + 1,
            name=_clean(name),
            office=_clean(office),
            secondary=_clean(secondary),
            identifier_raw=_clean(identifier_raw),
            activities=sections,
            status=status,
            outcome_raw=_clean(update_raw),
            listing_date=listing_date,
            expiry_date=expiry_date,
            primary_date_label="Data iscrizione",
            source_fields={
                "sections": sections,
                "requested_activities_source": _clean(sections_raw),
                "listing_date_raw_variants": [listing_source] if listing_source else [],
                "expiry_date_raw_variants": [expiry_source] if expiry_source else [],
                "in_aggiornamento": _clean(update_raw) if status == "renewal_update_in_progress" else "",
            },
        )
        record["identifiers"] = identifiers
        records.append(record)

    status_counts = dict(Counter(record["source_status"] for record in records))
    if len(records) != _LISTED_RECORDS:
        raise RuntimeError(f"Catania listed row drift: {len(records)} != {_LISTED_RECORDS}")
    if status_counts != _EXPECTED_LISTED_STATUS_COUNTS:
        raise RuntimeError(f"Catania listed status drift: {status_counts!r}")
    if identifier_coverage != _EXPECTED_LISTED_IDENTIFIER_COVERAGE:
        raise RuntimeError(
            f"Catania listed identifier coverage drift: {identifier_coverage} != {_EXPECTED_LISTED_IDENTIFIER_COVERAGE}"
        )
    if formula_count != _EXPECTED_LISTED_FORMULAS or standard_formula_count != _EXPECTED_LISTED_STANDARD_FORMULAS:
        raise RuntimeError(
            f"Catania listed formula drift: total={formula_count}, standard={standard_formula_count}"
        )
    if blank_expiries != _REVIEWED_LISTED_BLANK_EXPIRIES:
        raise RuntimeError(f"Catania listed blank-expiry drift: {sorted(blank_expiries)!r}")
    if blank_offices != _REVIEWED_LISTED_BLANK_OFFICES:
        raise RuntimeError(f"Catania listed blank-office drift: {sorted(blank_offices)!r}")
    if reviewed_malformed_dates != _REVIEWED_LISTED_MALFORMED_LISTING_DATES:
        raise RuntimeError(f"Catania listed malformed-date class drift: {reviewed_malformed_dates!r}")

    return ParsedBatch(
        records=records,
        diagnostics={
            "parser": "catania_listed",
            "parser_version": PARSER_VERSION,
            "public_records": len(records),
            "status_counts": status_counts,
            "identifier_coverage": identifier_coverage,
            "expiry_formula_rows": formula_count,
            "standard_expiry_formula_rows": standard_formula_count,
            "blank_expiry_rows": len(blank_expiries),
            "reviewed_malformed_listing_dates": len(reviewed_malformed_dates),
        },
    )


def parse_catania_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, "catania-applicants")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != [_SHEET]:
        raise RuntimeError(f"Catania applicant worksheet drift: {workbook.sheetnames!r}")
    sheet = workbook[_SHEET]
    _validate_structure(sheet, population="applicants")

    identifier_coverage = 0
    nonstandard_identifiers: dict[int, str] = {}
    outcome_values: Counter[str] = Counter()
    records: list[dict[str, Any]] = []

    for source_row in range(_APPLICANT_FIRST_ROW, _APPLICANT_LAST_ROW + 1):
        raw = [cell.value for cell in sheet[source_row]][:7]
        if len(raw) != 7 or not any(_clean(value) for value in raw):
            raise RuntimeError(f"Catania applicant company-row drift at source row {source_row}")
        name, office, secondary, sections_raw, identifier_raw, application_raw, outcome_raw = raw
        if not all(_clean(value) for value in (name, office, sections_raw, identifier_raw, application_raw)):
            raise RuntimeError(f"Catania applicant incomplete company row {source_row}: {[_clean(x) for x in raw]!r}")
        if _clean(outcome_raw):
            raise RuntimeError(f"Catania applicant outcome became nonblank at row {source_row}: {_clean(outcome_raw)!r}")

        sections = _sections(sections_raw, source_row=source_row, population="applicants")
        identifiers = _strict_identifiers(identifier_raw)
        identifier_coverage += bool(identifiers)
        if not identifiers:
            nonstandard_identifiers[source_row] = _clean(identifier_raw)
        application_source = _clean(application_raw)
        application_date = _strict_date(application_raw, context=f"applicant row {source_row} application")
        outcome_values[_clean(outcome_raw)] += 1

        record = _record(
            cfg,
            len(records) + 1,
            name=_clean(name),
            office=_clean(office),
            secondary=_clean(secondary),
            identifier_raw=_clean(identifier_raw),
            activities=sections,
            status="pending",
            outcome_raw=_clean(outcome_raw),
            application_date=application_date,
            primary_date_label="Data presentazione istanza",
            source_fields={
                "sections": sections,
                "requested_activities_source": _clean(sections_raw),
                "application_date_raw_variants": [application_source],
            },
        )
        record["identifiers"] = identifiers
        records.append(record)

    status_counts = dict(Counter(record["source_status"] for record in records))
    if len(records) != _APPLICANT_RECORDS:
        raise RuntimeError(f"Catania applicant row drift: {len(records)} != {_APPLICANT_RECORDS}")
    if status_counts != _EXPECTED_APPLICANT_STATUS_COUNTS:
        raise RuntimeError(f"Catania applicant status drift: {status_counts!r}")
    if identifier_coverage != _EXPECTED_APPLICANT_IDENTIFIER_COVERAGE:
        raise RuntimeError(
            f"Catania applicant identifier coverage drift: {identifier_coverage} != {_EXPECTED_APPLICANT_IDENTIFIER_COVERAGE}"
        )
    if nonstandard_identifiers != _REVIEWED_APPLICANT_NONSTANDARD_IDENTIFIERS:
        raise RuntimeError(f"Catania applicant nonstandard-identifier drift: {nonstandard_identifiers!r}")
    if dict(outcome_values) != {"": _APPLICANT_RECORDS}:
        raise RuntimeError(f"Catania applicant outcome drift: {dict(outcome_values)!r}")

    return ParsedBatch(
        records=records,
        diagnostics={
            "parser": "catania_applicants",
            "parser_version": PARSER_VERSION,
            "public_records": len(records),
            "status_counts": status_counts,
            "identifier_coverage": identifier_coverage,
            "nonstandard_identifier_rows": len(nonstandard_identifiers),
        },
    )


PARSERS = {
    "catania_listed": parse_catania_listed,
    "catania_applicants": parse_catania_applicants,
}
