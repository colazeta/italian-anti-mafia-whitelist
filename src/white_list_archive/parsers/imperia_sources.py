from __future__ import annotations

import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record

PARSER_VERSION = "1"
_REFERENCE_DATE = "2026-09-18"

_LISTED_SHEETS = (
    "legenda",
    "SEZIONE I",
    "SEZIONE II",
    "SEZIONE III",
    "SEZIONE IV",
    "SEZIONE V",
    "SEZIONE VI",
    "SEZIONE VII",
    "SEZIONE VIII",
    "SEZIONE IX",
    "SEZIONE X",
    "Foglio1",
)
_ROMAN = ("I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X")
_SECTION_LABELS = {
    1: "Estrazione, fornitura e trasporto di terra e materiali inerti",
    2: "Confezionamento, fornitura e trasporto di calcestruzzo e bitume",
    3: "Nolo a freddo di macchinari",
    4: "Fornitura di ferro lavorato",
    5: "Noli a caldo",
    6: "Autotrasporto per conto di terzi",
    7: "Guardiania dei cantieri",
    8: "Servizi Funerari e cimiteriali",
    9: "Ristorazione, gestione delle mense e catering",
    10: "Servizi ambientali, comprese le attività di raccolta, di trasporto nazionale e tranfrontaliero, anche per conto terzi, di trattamento e smaltimento dei rifiuti, nonché le attività di risanamento e di bonifica e gli altri servizi connessi alla gestione dei rifiuti",
}
_EXPECTED_LISTED_HEADER = (
    "Ragione Sociale",
    "Sede legale",
    "Sede secondaria con rappresentanza in Italia",
    "Codice fiscale/Partita IVA",
    "Data di iscrizione",
    "Data di scadenza iscrizione",
    "Aggiornamento in corso",
)
_EXPECTED_LISTED_SECTOR_COUNTS = {1: 53, 2: 11, 3: 42, 4: 17, 5: 29, 6: 47, 7: 2, 8: 1, 9: 6, 10: 38}
_EXPECTED_LISTED_SECTOR_ROWS = 246
_EXPECTED_LISTED_RECORDS = 142
_EXPECTED_LISTED_STATUS_COUNTS = Counter({"listed": 105, "renewal_update_in_progress": 37})
_EXPECTED_LISTED_UPDATE_VALUES = Counter({"": 105, "IN CORSO - PER RINNOVO": 36, "IN CORSO - PER MODIFICHE SOCIETARIE": 1})
_EXPECTED_LISTED_IDENTIFIER_COVERAGE = 131
_EXPECTED_LISTED_MALFORMED_SOURCE_DATES = Counter({"06/07/206": 4, "07/07/206": 2})

_EXPECTED_APPLICANT_TABLE_ROWS = 337
_EXPECTED_APPLICANT_HEADER_ROWS = 68
_EXPECTED_APPLICANT_BLANK_ROWS = 13
_EXPECTED_APPLICANT_CONTINUATION_ROWS = 125
_EXPECTED_APPLICANT_RECORDS = 131
_EXPECTED_APPLICANT_STATUS_COUNTS = Counter({
    "listed": 117,
    "pending": 6,
    "cancellation_related": 5,
    "other_or_unknown": 3,
})
_EXPECTED_APPLICANT_IDENTIFIER_COVERAGE = 126
_EXPECTED_APPLICANT_DATE_DIAGNOSTICS = Counter({"valid": 122, "blank": 7, "malformed": 2})
_EXPECTED_APPLICANT_MALFORMED_DATES = Counter({"25/0720522": 1, "10/1\\0/2024": 1})

_STRICT_VAT = re.compile(r"(?<!\d)(\d{11})(?!\d)")
_STRICT_CF = re.compile(r"(?<![A-Z0-9])([A-Z]{6}\d{2}[A-Z]\d{2}[A-Z]\d{3}[A-Z])(?![A-Z0-9])", re.I)
_MONTHS = {
    "gennaio": 1,
    "febbraio": 2,
    "marzo": 3,
    "aprile": 4,
    "maggio": 5,
    "giugno": 6,
    "luglio": 7,
    "agosto": 8,
    "settembre": 9,
    "ottobre": 10,
    "novembre": 11,
    "dicembre": 12,
}


def _source_identifiers(raw: str) -> list[str]:
    values: list[str] = []
    for pattern in (_STRICT_VAT, _STRICT_CF):
        for match in pattern.finditer(_clean(raw)):
            value = match.group(1).upper()
            if value not in values:
                values.append(value)
    return values


def _raw_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _clean(value)


def _source_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _clean(value)
    if not raw:
        return ""
    folded = raw.casefold().replace("º", "°")
    match = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", folded)
    if match is not None:
        day, month, year = map(int, match.groups())
    else:
        short = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2})", folded)
        if short is not None:
            day, month, year = map(int, short.groups())
            year += 2000
        else:
            named = re.fullmatch(r"(\d{1,2})°?\s+([a-zàèéìòù]+)\s+(\d{4})", folded)
            if named is None or named.group(2) not in _MONTHS:
                return ""
            day = int(named.group(1))
            month = _MONTHS[named.group(2)]
            year = int(named.group(3))
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return ""


def _listed_status(update: str) -> str:
    value = _clean(update)
    if not value:
        return "listed"
    if value in {"IN CORSO - PER RINNOVO", "IN CORSO - PER MODIFICHE SOCIETARIE"}:
        return "renewal_update_in_progress"
    raise RuntimeError(f"Imperia listed update/status vocabulary changed: {update!r}")


def _applicant_status(name: str, outcome: str) -> str:
    folded_name = _clean(name).casefold()
    folded = _clean(outcome).casefold()
    if "cancellat" in folded or "cancellat" in folded_name:
        return "cancellation_related"
    if folded.startswith("iscritta "):
        return "listed"
    if folded == "in corso":
        return "pending"
    if not folded or folded == "06/02/2025":
        return "other_or_unknown"
    raise RuntimeError(f"Imperia applicant outcome vocabulary changed for {name!r}: {outcome!r}")


def _outcome_listing_date(outcome: str) -> str:
    value = _clean(outcome).replace("’", "'")
    if not value.casefold().startswith("iscritta "):
        return ""
    raw = re.sub(r"^iscritta\s+(?:dall['/]|dal|del)\s*", "", value, flags=re.I)
    parsed = _source_date(raw)
    if not parsed:
        raise RuntimeError(f"Imperia applicant explicit enrolment date became unparseable: {outcome!r}")
    return parsed


def _antiword_docbook(path: Path) -> str:
    executable = shutil.which("antiword")
    if executable is None:
        raise RuntimeError("antiword is required to parse the reviewed Imperia applicant legacy Word source")
    process = subprocess.run(
        [executable, "-x", "db", str(path)],
        check=False,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if process.returncode:
        raise RuntimeError(
            f"antiword failed for Imperia source {path.name}: "
            f"{process.stderr.decode('utf-8', 'replace').strip()}"
        )
    return process.stdout.decode("utf-8", "replace")


def _validate_cfg(cfg: dict[str, Any], source_key: str) -> None:
    if cfg.get("source_key") != source_key or cfg.get("authority_key") != "imperia":
        raise RuntimeError(f"Imperia parser/source binding changed: {cfg.get('source_key')!r}")
    if cfg.get("reference_date") != _REFERENCE_DATE:
        raise RuntimeError(f"{source_key}: unexpected reference date {cfg.get('reference_date')!r}")


def _listed_sector_rows(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if tuple(workbook.sheetnames) != _LISTED_SHEETS:
        raise RuntimeError(f"Imperia listed workbook sheet set changed: {workbook.sheetnames!r}")
    foglio = workbook["Foglio1"]
    if any(_clean(value) for row in foglio.iter_rows(values_only=True) for value in row):
        raise RuntimeError("Imperia listed Foglio1 unexpectedly contains source values")

    source_rows: list[dict[str, Any]] = []
    sector_counts: Counter[int] = Counter()
    malformed_dates: Counter[str] = Counter()
    for section, roman in enumerate(_ROMAN, 1):
        worksheet = workbook[f"SEZIONE {roman}"]
        section_marker = _clean(worksheet.cell(15, 1).value)
        if section_marker.strip() != f"Sezione {roman}":
            raise RuntimeError(f"Imperia listed section marker changed on {roman}: {section_marker!r}")
        title = _clean(worksheet.cell(17, 1).value)
        if title != _SECTION_LABELS[section]:
            raise RuntimeError(f"Imperia listed section title changed on {roman}: {title!r}")
        header_row = 21 if section == 10 else 19
        observed_header = tuple(_clean(worksheet.cell(header_row, column).value) for column in range(1, 8))
        if observed_header != _EXPECTED_LISTED_HEADER:
            raise RuntimeError(
                f"Imperia listed header changed on section {roman}: {observed_header!r}"
            )
        for row_number in range(header_row + 1, worksheet.max_row + 1):
            values = [worksheet.cell(row_number, column).value for column in range(1, 9)]
            if not any(_clean(value) for value in values):
                continue
            if _clean(values[7]):
                raise RuntimeError(f"Imperia listed unexpected eighth-column content at {roman}:{row_number}")
            name, office, secondary, identifier_raw, listing_value, expiry_value, update_value = values[:7]
            name = _clean(name)
            if not name:
                raise RuntimeError(f"Imperia listed nonblank row lost company identity at {roman}:{row_number}")
            listing_raw = _raw_date(listing_value)
            expiry_raw = _raw_date(expiry_value)
            listing_date = _source_date(listing_value)
            expiry_date = _source_date(expiry_value)
            if listing_raw and not listing_date:
                malformed_dates[listing_raw] += 1
            if not expiry_raw or not expiry_date:
                raise RuntimeError(f"Imperia listed expiry date changed at {roman}:{row_number}: {expiry_raw!r}")
            update = _clean(update_value)
            source_rows.append({
                "section": section,
                "row": row_number,
                "name": name,
                "office": _clean(office),
                "secondary": _clean(secondary),
                "identifier_raw": _clean(identifier_raw),
                "listing_raw": listing_raw,
                "listing_date": listing_date,
                "expiry_raw": expiry_raw,
                "expiry_date": expiry_date,
                "update": update,
                "status": _listed_status(update),
            })
            sector_counts[section] += 1

    if dict(sector_counts) != _EXPECTED_LISTED_SECTOR_COUNTS:
        raise RuntimeError(f"Imperia listed sector-row counts changed: {dict(sector_counts)!r}")
    if len(source_rows) != _EXPECTED_LISTED_SECTOR_ROWS:
        raise RuntimeError(
            f"Imperia listed source-row denominator changed: {len(source_rows)} != {_EXPECTED_LISTED_SECTOR_ROWS}"
        )
    if malformed_dates != _EXPECTED_LISTED_MALFORMED_SOURCE_DATES:
        raise RuntimeError(f"Imperia listed reviewed malformed-date boundary changed: {dict(malformed_dates)!r}")
    return source_rows


def parse_imperia_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, "imperia-listed")
    sector_rows = _listed_sector_rows(path)
    grouped: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in sector_rows:
        key = (
            row["name"].casefold(),
            row["office"].casefold(),
            row["secondary"].casefold(),
            row["identifier_raw"].casefold(),
            row["listing_date"],
            row["expiry_date"],
            row["status"],
            row["update"].casefold(),
        )
        group = grouped.setdefault(key, {"row": row, "sections": [], "source_rows": []})
        if row["section"] in group["sections"]:
            raise RuntimeError(
                f"Imperia listed duplicate same-section source row in one observation: {row['name']!r}, section {row['section']}"
            )
        group["sections"].append(row["section"])
        group["source_rows"].append(row["row"])

    if len(grouped) != _EXPECTED_LISTED_RECORDS:
        raise RuntimeError(
            f"Imperia listed grouped-observation denominator changed: {len(grouped)} != {_EXPECTED_LISTED_RECORDS}"
        )

    records: list[dict[str, Any]] = []
    for group in grouped.values():
        row = group["row"]
        sections = sorted(group["sections"])
        record = _record(
            cfg,
            len(records) + 1,
            name=row["name"],
            office=row["office"],
            secondary=row["secondary"],
            identifier_raw=row["identifier_raw"],
            activities=[_SECTION_LABELS[section] for section in sections],
            status=row["status"],
            outcome_raw=row["update"],
            listing_date=row["listing_date"],
            expiry_date=row["expiry_date"],
            primary_date_label="Data iscrizione",
            source_fields={
                "sections": [f"Sezione {_ROMAN[section - 1]}" for section in sections],
                "physical_locators": [
                    f"SEZIONE {_ROMAN[section - 1]}:r{source_row}"
                    for section, source_row in zip(sections, group["source_rows"], strict=True)
                ],
                "listing_date_raw_variants": [row["listing_raw"]] if row["listing_raw"] else [],
                "expiry_date_raw_variants": [row["expiry_raw"]] if row["expiry_raw"] else [],
                "in_aggiornamento": row["update"] if row["status"] == "renewal_update_in_progress" else "",
            },
        )
        record["identifiers"] = _source_identifiers(row["identifier_raw"])
        records.append(record)

    status_counts = Counter(record["source_status"] for record in records)
    if status_counts != _EXPECTED_LISTED_STATUS_COUNTS:
        raise RuntimeError(f"Imperia listed status drift: {dict(status_counts)!r}")
    update_counts = Counter(record["outcome_raw"] for record in records)
    if update_counts != _EXPECTED_LISTED_UPDATE_VALUES:
        raise RuntimeError(f"Imperia listed update-text drift: {dict(update_counts)!r}")
    coverage = sum(bool(record["identifiers"]) for record in records)
    if coverage != _EXPECTED_LISTED_IDENTIFIER_COVERAGE:
        raise RuntimeError(f"Imperia listed identifier-coverage drift: {coverage}")
    malformed_logical = Counter(
        record["source_fields"]["listing_date_raw_variants"][0]
        for record in records
        if record["source_fields"]["listing_date_raw_variants"] and not record["observed_listing_date"]
    )
    if malformed_logical != Counter({"07/07/206": 1, "06/07/206": 1}):
        raise RuntimeError(f"Imperia listed logical malformed-date boundary changed: {dict(malformed_logical)!r}")
    return ParsedBatch(records, {
        "parser": "imperia_listed",
        "sector_rows": len(sector_rows),
        "public_records": len(records),
        "status_counts": dict(status_counts),
        "identifier_coverage": coverage,
        "raw_only_identifier_records": len(records) - coverage,
        "malformed_listing_date_records": sum(not record["observed_listing_date"] for record in records),
    })


def _applicant_rows(docbook: str) -> list[dict[str, Any]]:
    xml = re.sub(r"<!DOCTYPE book PUBLIC.*?docbookx\.dtd\">\s*", "", docbook, flags=re.S)
    root = ET.fromstring(xml)
    tables = root.findall(".//informaltable")
    if len(tables) != 1:
        raise RuntimeError(f"Imperia applicants: expected one table, got {len(tables)}")
    groups = tables[0].findall(".//tgroup")
    if len(groups) != 1 or groups[0].get("cols") != "7":
        raise RuntimeError("Imperia applicants: table column-count drift")

    raw_rows: list[list[str]] = []
    for row in tables[0].findall(".//row"):
        cells = [_clean("".join(entry.itertext())) for entry in row.findall("./entry")]
        if len(cells) != 7:
            raise RuntimeError(f"Imperia applicants: row column-count drift: {len(cells)}")
        raw_rows.append(cells)
    if len(raw_rows) != _EXPECTED_APPLICANT_TABLE_ROWS:
        raise RuntimeError(
            f"Imperia applicants: table-row denominator changed: {len(raw_rows)} != {_EXPECTED_APPLICANT_TABLE_ROWS}"
        )

    header_rows = 0
    blank_rows = 0
    continuation_rows = 0
    logical: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for row_number, cells in enumerate(raw_rows, 1):
        if cells[0].casefold() == "ragione sociale" and cells[-1].casefold() == "esito":
            expected = (
                "Ragione sociale",
                "Sede legale",
                cells[2],
                "Codice fiscale / Partita IVA",
                "Attività per cui è richiesta l’iscrizione",
                "Data di presentazione dell’istanza",
                "Esito",
            )
            if tuple(cells) != expected or cells[2] not in {"", "Sede secondaria con rappresentanza stabile in Italia"}:
                raise RuntimeError(f"Imperia applicants: repeated header changed at row {row_number}: {cells!r}")
            header_rows += 1
            continue
        if not any(cells):
            blank_rows += 1
            continue
        if cells[0]:
            current = {
                "source_row": row_number,
                "name": cells[0],
                "office": cells[1],
                "secondary": cells[2],
                "identifier_raw": cells[3],
                "activities": [cells[4]] if cells[4] else [],
                "application_raw": cells[5],
                "outcome": cells[6],
                "continuation_rows": [],
            }
            logical.append(current)
            continue
        if current is None:
            raise RuntimeError(f"Imperia applicants: orphan continuation row {row_number}")
        if any(cells[index] for index in (1, 2, 3, 5, 6)) or not cells[4]:
            raise RuntimeError(f"Imperia applicants: continuation-row shape changed at row {row_number}: {cells!r}")
        current["activities"].append(cells[4])
        current["continuation_rows"].append(row_number)
        continuation_rows += 1

    if header_rows != _EXPECTED_APPLICANT_HEADER_ROWS:
        raise RuntimeError(f"Imperia applicants: repeated-header denominator changed: {header_rows}")
    if blank_rows != _EXPECTED_APPLICANT_BLANK_ROWS:
        raise RuntimeError(f"Imperia applicants: blank-row denominator changed: {blank_rows}")
    if continuation_rows != _EXPECTED_APPLICANT_CONTINUATION_ROWS:
        raise RuntimeError(f"Imperia applicants: continuation-row denominator changed: {continuation_rows}")
    if len(logical) != _EXPECTED_APPLICANT_RECORDS:
        raise RuntimeError(f"Imperia applicants: observation denominator changed: {len(logical)}")
    return logical


def parse_imperia_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, "imperia-applicants")
    rows = _applicant_rows(_antiword_docbook(path))
    records: list[dict[str, Any]] = []
    date_diagnostics: Counter[str] = Counter()
    malformed_dates: Counter[str] = Counter()
    for row in rows:
        status = _applicant_status(row["name"], row["outcome"])
        application_date = _source_date(row["application_raw"])
        if application_date:
            date_diagnostics["valid"] += 1
        elif row["application_raw"]:
            date_diagnostics["malformed"] += 1
            malformed_dates[row["application_raw"]] += 1
        else:
            date_diagnostics["blank"] += 1
        listing_date = _outcome_listing_date(row["outcome"]) if status == "listed" else ""
        physical_rows = [row["source_row"], *row["continuation_rows"]]
        record = _record(
            cfg,
            len(records) + 1,
            name=row["name"],
            office=row["office"],
            secondary=row["secondary"],
            identifier_raw=row["identifier_raw"],
            activities=list(row["activities"]),
            status=status,
            outcome_raw=row["outcome"],
            application_date=application_date,
            listing_date=listing_date,
            primary_date_label="Data presentazione istanza",
            source_fields={
                "physical_locators": [f"table:r{number}" for number in physical_rows],
                "requested_activities_source": " · ".join(row["activities"]),
                "application_date_raw_variants": [row["application_raw"]] if row["application_raw"] else [],
                "listing_date_raw_variants": [row["outcome"]] if status == "listed" else [],
                "notes": [row["outcome"]] if row["outcome"] else [],
            },
        )
        record["identifiers"] = _source_identifiers(row["identifier_raw"])
        records.append(record)

    status_counts = Counter(record["source_status"] for record in records)
    if status_counts != _EXPECTED_APPLICANT_STATUS_COUNTS:
        raise RuntimeError(f"Imperia applicant status drift: {dict(status_counts)!r}")
    coverage = sum(bool(record["identifiers"]) for record in records)
    if coverage != _EXPECTED_APPLICANT_IDENTIFIER_COVERAGE:
        raise RuntimeError(f"Imperia applicant identifier-coverage drift: {coverage}")
    if date_diagnostics != _EXPECTED_APPLICANT_DATE_DIAGNOSTICS:
        raise RuntimeError(f"Imperia applicant application-date boundary changed: {dict(date_diagnostics)!r}")
    if malformed_dates != _EXPECTED_APPLICANT_MALFORMED_DATES:
        raise RuntimeError(f"Imperia applicant reviewed malformed-date tokens changed: {dict(malformed_dates)!r}")
    listed_dates = sum(bool(record["observed_listing_date"]) for record in records if record["source_status"] == "listed")
    if listed_dates != 117:
        raise RuntimeError(f"Imperia applicant explicit subsequent-enrolment date coverage changed: {listed_dates}")
    return ParsedBatch(records, {
        "parser": "imperia_applicants",
        "table_rows": _EXPECTED_APPLICANT_TABLE_ROWS,
        "continuation_rows": _EXPECTED_APPLICANT_CONTINUATION_ROWS,
        "public_records": len(records),
        "status_counts": dict(status_counts),
        "identifier_coverage": coverage,
        "raw_only_identifier_records": len(records) - coverage,
        "application_date_diagnostics": dict(date_diagnostics),
        "explicit_subsequent_enrolment_dates": listed_dates,
    })


PARSERS = {
    "imperia_listed": parse_imperia_listed,
    "imperia_applicants": parse_imperia_applicants,
}
