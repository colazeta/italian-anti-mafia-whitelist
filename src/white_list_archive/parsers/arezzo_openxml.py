from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

import openpyxl
from docx import Document

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record


_SECTION_PREFIX = re.compile(r"^SEZIONE\s+[IVX]+\s*[–—-]\s*", re.I)
_DMY_REPEATED_SLASH = re.compile(r"^(\d{1,2})/+(\d{1,2})/+(\d{4})$")

# These are the ten statutory/source section labels observed in the current Arezzo
# lists.  They are used only to recover explicit labels from the DOCX applicant
# cell, where Word concatenates multiple requested activities without separators.
# Unknown residue fails closed instead of being silently reinterpreted.
_ACTIVITY_LABELS = (
    "Servizi ambientali, comprese le attività di raccolta, di trasporto nazionale e transfrontaliero, anche per conto di terzi, di trattamento e di smaltimento dei rifiuti (ex sez. I e sez. II), nonché le attività di risanamento e di bonifica e gli altri servizi connessi alla gestione dei rifiuti",
    "Confezionamento, fornitura e trasporto di calcestruzzo e di bitume",
    "Estrazione, fornitura e trasporto di terra e materiali inerti",
    "Ristorazione, gestione delle mense e catering",
    "Autotrasporti per conto di terzi",
    "Autotrasporti per conto terzi",
    "Servizi funerari e cimiteriali",
    "Noli a freddo di macchinari",
    "Fornitura di ferro lavorato",
    "Guardiania ai cantieri",
    "Noli a caldo",
    # The current Arezzo document abbreviates section X to this label.
    "Servizi ambientali",
)
_ACTIVITY_PATTERN = re.compile(
    "|".join(re.escape(label) for label in sorted(_ACTIVITY_LABELS, key=len, reverse=True)),
    re.I,
)


def _source_date(value: Any) -> str:
    """Normalise only unambiguous date representation, never missing source digits."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _clean(value)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?: 00:00:00)?", raw):
        try:
            return date.fromisoformat(raw[:10]).isoformat()
        except ValueError:
            return ""
    match = _DMY_REPEATED_SLASH.fullmatch(raw)
    if not match:
        return ""
    day, month, year = map(int, match.groups())
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return ""


def _section_activity(value: str) -> str:
    source = _clean(value)
    return _SECTION_PREFIX.sub("", source).strip() or source


def _applicant_activities(value: str) -> list[str]:
    source = _clean(value)
    if not source:
        return []
    matches = list(_ACTIVITY_PATTERN.finditer(source))
    if not matches:
        raise ValueError(f"Arezzo applicant activity is not in the audited source vocabulary: {source!r}")
    residue = list(source)
    for match in matches:
        for index in range(match.start(), match.end()):
            residue[index] = " "
    if _clean("".join(residue)):
        raise ValueError(f"Unparsed Arezzo applicant activity residue: {source!r}")
    return [_clean(match.group(0)) for match in matches]


def _validate_listed_header(rows: list[tuple[Any, ...]], sheet: str) -> None:
    if len(rows) < 5:
        raise ValueError(f"Arezzo listed sheet {sheet!r} is shorter than the audited layout")
    header = [_clean(value).casefold() for value in rows[4]]
    required = ("denominazione", "sede legale", "codice fiscale", "data iscrizione", "data scadenza")
    joined = " | ".join(header)
    if any(token not in joined for token in required):
        raise ValueError(f"Arezzo listed sheet {sheet!r} header changed: {header!r}")


def parse_arezzo_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sector_rows: list[dict[str, str]] = []
    malformed_date_rows = 0
    dropped = 0

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        _validate_listed_header(rows, worksheet.title)
        section = _clean(rows[3][0]) if len(rows[3]) else worksheet.title
        if not section.casefold().startswith("sezione"):
            raise ValueError(f"Arezzo listed section heading changed in {worksheet.title!r}: {section!r}")
        for source_row_number, raw_row in enumerate(rows[5:], start=6):
            values = list(raw_row) + [None] * max(0, 7 - len(raw_row))
            cleaned = [_clean(value) for value in values[:7]]
            if not any(cleaned):
                continue
            name, office, secondary, identifier_raw, listing_raw, expiry_raw, update = cleaned
            if not name or not identifier_raw or not listing_raw or not expiry_raw:
                dropped += 1
                continue
            listing = _source_date(values[4])
            expiry = _source_date(values[5])
            # A malformed listing date remains a valid source observation when the
            # row identity and expiry date are explicit.  We retain its raw form and
            # leave the canonical listing date blank rather than guessing punctuation.
            if not expiry:
                dropped += 1
                continue
            if not listing:
                malformed_date_rows += 1
            sector_rows.append(
                {
                    "section": section,
                    "sheet": worksheet.title,
                    "source_row": str(source_row_number),
                    "name": name,
                    "office": office,
                    "secondary": secondary,
                    "identifier_raw": identifier_raw,
                    "listing_raw": listing_raw,
                    "expiry_raw": expiry_raw,
                    "listing": listing,
                    "expiry": expiry,
                    "update": update,
                }
            )

    grouped: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in sector_rows:
        # Preserve source conflicts.  The raw malformed listing value participates
        # in the key so an unknown date can never collapse into another observation.
        listing_key = row["listing"] or f"RAW:{row['listing_raw']}"
        key = (
            row["name"], row["office"], row["secondary"], row["identifier_raw"],
            listing_key, row["expiry"], row["update"],
        )
        group = grouped.setdefault(
            key,
            {
                "row": row,
                "sections": [],
                "offices": [],
                "secondary_offices": [],
                "listing_raw": [],
                "expiry_raw": [],
                "malformed": [],
            },
        )
        for field, value in (
            ("sections", row["section"]),
            ("offices", row["office"]),
            ("secondary_offices", row["secondary"]),
            ("listing_raw", row["listing_raw"]),
            ("expiry_raw", row["expiry_raw"]),
        ):
            if value and value not in group[field]:
                group[field].append(value)
        if not row["listing"]:
            pair = f"{row['listing_raw']} | {row['expiry_raw']}"
            if pair not in group["malformed"]:
                group["malformed"].append(pair)

    identity_variants: dict[tuple[str, str], set[tuple[str, str, str]]] = defaultdict(set)
    for group in grouped.values():
        row = group["row"]
        identity_variants[(row["name"], row["identifier_raw"])].add(
            (row["listing"] or f"RAW:{row['listing_raw']}", row["expiry"], row["update"])
        )

    records: list[dict[str, Any]] = []
    conflict_groups = 0
    for group in grouped.values():
        row = group["row"]
        variants = identity_variants[(row["name"], row["identifier_raw"])]
        conflict_fields: list[str] = []
        if len(variants) > 1:
            listings = {variant[0] for variant in variants}
            expiries = {variant[1] for variant in variants}
            updates = {variant[2] for variant in variants}
            if len(listings) > 1:
                conflict_fields.append("listing_date")
            if len(expiries) > 1:
                conflict_fields.append("expiry_date")
            if len(updates) > 1:
                conflict_fields.append("update_status")
        records.append(
            _record(
                cfg,
                len(records) + 1,
                name=row["name"],
                office=row["office"],
                secondary=row["secondary"],
                identifier_raw=row["identifier_raw"],
                activities=[_section_activity(section) for section in group["sections"]],
                status="renewal_update_in_progress" if row["update"] else "listed",
                outcome_raw=row["update"],
                listing_date=row["listing"],
                expiry_date=row["expiry"],
                primary_date_label="Data iscrizione",
                source_fields={
                    "sections": group["sections"],
                    "registered_office_variants": group["offices"],
                    "secondary_office_variants": group["secondary_offices"],
                    "listing_date_raw_variants": group["listing_raw"],
                    "expiry_date_raw_variants": group["expiry_raw"],
                    "normalised_listing_date_variants": [row["listing"]] if row["listing"] else [],
                    "normalised_expiry_date_variants": [row["expiry"]],
                    "date_conflict_fields": conflict_fields,
                    "malformed_date_pairs": group["malformed"],
                },
            )
        )

    conflict_groups = sum(1 for variants in identity_variants.values() if len(variants) > 1)
    return ParsedBatch(
        records,
        {
            "parser": "arezzo_listed",
            "sector_rows": len(sector_rows),
            "date_rows": len(sector_rows),
            "public_records": len(records),
            "dropped_date_rows": dropped,
            "malformed_date_rows_preserved": malformed_date_rows,
            "date_conflict_identity_groups": conflict_groups,
            "status_counts": dict(Counter(record["source_status"] for record in records)),
            "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
            "raw_identifier_only": sum(bool(record["identifier_field_raw"]) and not record["identifiers"] for record in records),
        },
    )


def parse_arezzo_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    document = Document(path)
    if len(document.tables) != 1:
        raise ValueError(f"Arezzo applicant document changed: expected one table, got {len(document.tables)}")
    table = document.tables[0]
    if not table.rows:
        raise ValueError("Arezzo applicant document has no table rows")
    header = [_clean(cell.text).casefold() for cell in table.rows[0].cells]
    joined = " | ".join(header)
    for token in ("ragione sociale", "sede legale", "codice fiscale", "attività", "data di presentazione"):
        if token not in joined:
            raise ValueError(f"Arezzo applicant header changed: {header!r}")

    records: list[dict[str, Any]] = []
    dropped = 0
    typography_normalised = 0
    for source_row_number, source_row in enumerate(table.rows[1:], start=2):
        values = [_clean(cell.text) for cell in source_row.cells]
        values += [""] * max(0, 6 - len(values))
        name, office, secondary, identifier_raw, activity_raw, application_raw = values[:6]
        if not any(values[:6]):
            continue
        application_date = _source_date(application_raw)
        if not name or not identifier_raw or not application_date:
            dropped += 1
            continue
        activities = _applicant_activities(activity_raw)
        source_fields: dict[str, Any] = {"requested_activities_source": activity_raw}
        if "//" in application_raw:
            source_fields["malformed_date_pairs"] = [f"application_date={application_raw}"]
            typography_normalised += 1
        records.append(
            _record(
                cfg,
                len(records) + 1,
                name=name,
                office=office,
                secondary=secondary,
                identifier_raw=identifier_raw,
                activities=activities,
                status="pending",
                outcome_raw="",
                application_date=application_date,
                primary_date_label="Data presentazione istanza",
                source_fields=source_fields,
            )
        )

    return ParsedBatch(
        records,
        {
            "parser": "arezzo_applicants",
            "date_rows": len(records) + dropped,
            "public_records": len(records),
            "dropped_date_rows": dropped,
            "separator_typography_normalised": typography_normalised,
            "status_counts": dict(Counter(record["source_status"] for record in records)),
            "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
            "raw_identifier_only": sum(bool(record["identifier_field_raw"]) and not record["identifiers"] for record in records),
        },
    )


PARSERS: dict[str, Callable[[Path, dict[str, Any]], ParsedBatch]] = {
    "arezzo_listed": parse_arezzo_listed,
    "arezzo_applicants": parse_arezzo_applicants,
}
