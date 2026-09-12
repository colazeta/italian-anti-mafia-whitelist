from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record

PARSER_VERSION = "1"
_REFERENCE_DATE = "2026-09-11"
_LISTED_SOURCE_KEY = "bolzano-bozen-listed"
_APPLICANT_SOURCE_KEY = "bolzano-bozen-applicants"
_LISTED_SHA256 = "96992db4caac16200fbebfa573bb602e1e396961ff854bbf006cb42e011c1e04"
_APPLICANT_SHA256 = "8b7321745edb19db688a001b9a1a706e84bee95cb457f9fdca1af144521ef3ae"
_EXPECTED_SECTION_ROWS = {
    1: 307,
    2: 144,
    3: 219,
    4: 158,
    5: 269,
    6: 226,
    7: 21,
    8: 10,
    9: 165,
    10: 187,
}
_EXPECTED_LISTED_SECTOR_ROWS = 1706
_EXPECTED_LISTED_RECORDS = 871
_EXPECTED_LISTED_STATUS_COUNTS = {"listed": 620, "renewal_update_in_progress": 251}
_EXPECTED_LISTED_UPDATE_VALUES = {
    "SI -JA": 4,
    "": 1142,
    "SI-JA": 446,
    "SI- Ja": 2,
    "Si-JA": 23,
    "SI - JA": 6,
    "Si - JA": 1,
    "Si-ja": 27,
    "Si- ja": 19,
    "SI- JA": 5,
    "Si- Ja": 9,
    "Si–JA": 1,
    "SI-Ja": 3,
    "Si - ja": 1,
    "Si -JA": 4,
    "SI/JA": 4,
    "Si.JA": 1,
    "SI- ja": 1,
    "Si_JA": 1,
    "Si- JA": 1,
    "SI-ja": 1,
    "SI- jA": 1,
    "SI-- Ja": 1,
    "Si-Ja": 1,
    "SI - Ja": 1,
}
_REVIEWED_LISTED_DATE_TYPOGRAPHY = frozenset(
    {
        "02/09/20259",
        "027/07/2026",
        "03/0/2025",
        "06/05/206",
        "1 9 /01/2026",
        "12/06/02025",
        "23/06/25025",
        "27/06/204",
        "29/6/2026",
    }
)
_REVIEWED_EXPIRY_DATE_TYPOGRAPHY = frozenset(
    {
        "1 7 /1 2 /202 6",
        "12/1 2/2026",
        "15/012026",
        "16/19/2026",
        "2/8/01/2027",
        "22/ 01/2027",
        "22/ 02487690212 04/2027",
        "22/0 4 /202 7",
        "25/008/2026",
        "27/08/202 6",
        "29/0/2026",
    }
)
_EXPECTED_APPLICANT_ROWS = 351
_EXPECTED_APPLICANT_RECORDS = 351
_EXPECTED_APPLICANT_STATUS_COUNTS = {"pending": 351}
_REVIEWED_APPLICANT_DATE_TYPOGRAPHY = frozenset({"14/032025"})
_DATE_DMY = re.compile(r"^(\d{2})/(\d{2})/(\d{4})$")


def _physical_cells(row: Any) -> list[str]:
    values: list[str] = []
    for tc in row._tr.tc_lst:
        text = _clean(" ".join((node.text or "") for node in tc.iter(qn("w:t"))))
        values.append(text)
    return values


def _validate_cfg(cfg: dict[str, Any], *, source_key: str, population_scope: str, sha256: str) -> None:
    if cfg.get("source_key") != source_key:
        raise RuntimeError(f"Bolzano parser/source mismatch: {cfg.get('source_key')!r} != {source_key!r}")
    if cfg.get("authority_key") != "bolzano-bozen":
        raise RuntimeError("Bolzano parser bound to a non-Bolzano authority")
    if cfg.get("population_scope") != population_scope:
        raise RuntimeError(f"Bolzano population-scope drift: {cfg.get('population_scope')!r}")
    if cfg.get("reference_date") != _REFERENCE_DATE:
        raise RuntimeError(f"Bolzano reference date drift: {cfg.get('reference_date')!r}")
    if cfg.get("sha256") != sha256:
        raise RuntimeError(f"Bolzano approved-byte digest drift: {cfg.get('sha256')!r}")


def _strict_source_date(raw: str, *, reviewed_malformed: frozenset[str], label: str) -> str:
    raw = _clean(raw)
    if raw in reviewed_malformed:
        return ""
    match = _DATE_DMY.fullmatch(raw)
    if match:
        day, month, year = map(int, match.groups())
        try:
            return date(year, month, day).isoformat()
        except ValueError as exc:
            raise RuntimeError(f"Bolzano unreviewed invalid {label} calendar date: {raw!r}") from exc
    raise RuntimeError(f"Bolzano unreviewed {label} date typography: {raw!r}")


def _strict_identifiers(raw: str) -> list[str]:
    value = _clean(raw).upper()
    if value.isdigit() and len(value) == 11:
        return [value]
    if len(value) == 16 and value.isalnum():
        return [value]
    return []


def _is_affirmative_update(raw: str) -> bool:
    if not raw:
        return False
    folded = re.sub(r"[\s_./–—-]+", "", raw.casefold())
    if folded not in {"sija", "sìja"}:
        raise RuntimeError(f"Bolzano unreviewed update status: {raw!r}")
    return True


def _validate_listed_header(values: list[str], section: int) -> None:
    if (
        len(values) != 7
        or values[0].casefold() != "ragione sociale"
        or "sede legale" not in values[1].casefold()
        or "sede secondaria" not in values[2].casefold()
        or "codice fiscale" not in values[3].casefold()
        or "data di iscrizione" not in values[4].casefold()
        or "data validità iscrizione" not in values[5].casefold()
        or "aggiornamento in corso" not in values[6].casefold()
    ):
        raise RuntimeError(f"Bolzano listed header drift in section {section}: {values!r}")


def _validate_applicant_header(values: list[str]) -> None:
    if (
        len(values) != 7
        or values[0].casefold() != "ragione sociale"
        or "sede legale" not in values[1].casefold()
        or "sede secondaria" not in values[2].casefold()
        or "codice fiscale" not in values[3].casefold()
        or "attività per cui" not in values[4].casefold()
        or "data di presentazione" not in values[5].casefold()
        or values[6].casefold() != "esito"
    ):
        raise RuntimeError(f"Bolzano applicant header drift: {values!r}")


def _one_update(values: list[str], *, section: int, source_row: int) -> str:
    present: list[str] = []
    for value in values:
        if value and value not in present:
            present.append(value)
    if len(present) > 1:
        raise RuntimeError(
            f"Bolzano ambiguous merged update cells in section {section}, row {source_row}: {present!r}"
        )
    return present[0] if present else ""


def parse_bolzano_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, source_key=_LISTED_SOURCE_KEY, population_scope="listed", sha256=_LISTED_SHA256)
    document = Document(path)
    if len(document.tables) != 10:
        raise RuntimeError(f"Bolzano listed table-count drift: {len(document.tables)}")

    sector_rows: list[dict[str, Any]] = []
    section_counts: Counter[int] = Counter()
    update_values: Counter[str] = Counter()
    seen_in_section: set[tuple[int, str, str, str, str, str, str]] = set()

    for section, table in enumerate(document.tables, start=1):
        if len(table.rows) < 3:
            raise RuntimeError(f"Bolzano listed section {section} unexpectedly short")
        _validate_listed_header(_physical_cells(table.rows[0]), section)
        for source_row, row in enumerate(table.rows[2:], start=3):
            values = _physical_cells(row)
            if not any(values):
                continue
            if len(values) < 7:
                raise RuntimeError(f"Bolzano short listed row in section {section}, row {source_row}: {values!r}")
            name, office, secondary, identifier_raw, listing_raw, expiry_raw = values[:6]
            update_raw = _one_update(values[6:], section=section, source_row=source_row)
            if not name or not listing_raw or not expiry_raw:
                raise RuntimeError(
                    f"Bolzano incomplete listed row in section {section}, row {source_row}: {values!r}"
                )
            listing_date = _strict_source_date(
                listing_raw, reviewed_malformed=_REVIEWED_LISTED_DATE_TYPOGRAPHY, label="listing"
            )
            expiry_date = _strict_source_date(
                expiry_raw, reviewed_malformed=_REVIEWED_EXPIRY_DATE_TYPOGRAPHY, label="expiry"
            )
            _is_affirmative_update(update_raw)
            key_in_section = (section, name, office, secondary, identifier_raw, listing_raw, expiry_raw)
            if key_in_section in seen_in_section:
                raise RuntimeError(f"Bolzano exact listed duplicate inside section {section}: {key_in_section!r}")
            seen_in_section.add(key_in_section)
            section_counts[section] += 1
            update_values[update_raw] += 1
            sector_rows.append(
                {
                    "section": section,
                    "source_row": source_row,
                    "name": name,
                    "office": office,
                    "secondary": secondary,
                    "identifier_raw": identifier_raw,
                    "listing_raw": listing_raw,
                    "listing_date": listing_date,
                    "expiry_raw": expiry_raw,
                    "expiry_date": expiry_date,
                    "update_raw": update_raw,
                }
            )

    if dict(section_counts) != _EXPECTED_SECTION_ROWS:
        raise RuntimeError(f"Bolzano listed section-row drift: {dict(section_counts)!r}")
    if len(sector_rows) != _EXPECTED_LISTED_SECTOR_ROWS:
        raise RuntimeError(f"Bolzano listed source-row drift: {len(sector_rows)}")
    if dict(update_values) != _EXPECTED_LISTED_UPDATE_VALUES:
        raise RuntimeError(f"Bolzano listed update-lexeme drift: {dict(update_values)!r}")

    grouped: dict[tuple[str, str, str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in sector_rows:
        grouped[
            (
                row["name"],
                row["office"],
                row["secondary"],
                row["identifier_raw"],
                row["listing_raw"],
                row["expiry_raw"],
            )
        ].append(row)
    if len(grouped) != _EXPECTED_LISTED_RECORDS:
        raise RuntimeError(f"Bolzano listed semantic-group drift: {len(grouped)}")

    records: list[dict[str, Any]] = []
    status_counts: Counter[str] = Counter()
    for ordinal, members in enumerate(grouped.values(), start=1):
        sections = [member["section"] for member in members]
        if len(sections) != len(set(sections)):
            raise RuntimeError(f"Bolzano listed same-section repeat escaped guard: {sections!r}")
        raw_updates = []
        for member in members:
            if member["update_raw"] and member["update_raw"] not in raw_updates:
                raw_updates.append(member["update_raw"])
        status = "renewal_update_in_progress" if raw_updates else "listed"
        status_counts[status] += 1
        first = members[0]
        record = _record(
            cfg,
            ordinal,
            name=first["name"],
            office=first["office"],
            secondary=first["secondary"],
            identifier_raw=first["identifier_raw"],
            activities=[f"Sezione {section}" for section in sections],
            status=status,
            listing_date=first["listing_date"],
            expiry_date=first["expiry_date"],
            primary_date_label="Data iscrizione",
            source_fields={
                "listed_sections": sections,
                "listed_source_rows": [member["source_row"] for member in members],
                "listing_date_raw": first["listing_raw"],
                "expiry_date_raw": first["expiry_raw"],
                "update_raw_values": raw_updates,
            },
        )
        record["identifiers"] = _strict_identifiers(first["identifier_raw"])
        records.append(record)

    if dict(status_counts) != _EXPECTED_LISTED_STATUS_COUNTS:
        raise RuntimeError(f"Bolzano listed status drift: {dict(status_counts)!r}")
    return ParsedBatch(
        records=records,
        diagnostics={
            "sector_rows": len(sector_rows),
            "section_rows": dict(section_counts),
            "public_records": len(records),
            "status_counts": dict(status_counts),
            "update_values": dict(update_values),
            "malformed_listing_values": sorted(
                {row["listing_raw"] for row in sector_rows if not row["listing_date"]}
            ),
            "malformed_expiry_values": sorted(
                {row["expiry_raw"] for row in sector_rows if not row["expiry_date"]}
            ),
        },
    )


def parse_bolzano_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, source_key=_APPLICANT_SOURCE_KEY, population_scope="applicant", sha256=_APPLICANT_SHA256)
    document = Document(path)
    if len(document.tables) != 1:
        raise RuntimeError(f"Bolzano applicant table-count drift: {len(document.tables)}")
    table = document.tables[0]
    if len(table.rows) < 3:
        raise RuntimeError("Bolzano applicant table unexpectedly short")
    _validate_applicant_header(_physical_cells(table.rows[0]))

    rows: list[dict[str, Any]] = []
    exact_seen: set[tuple[str, str, str, str, str, str, str]] = set()
    for source_row, row in enumerate(table.rows[2:], start=3):
        values = _physical_cells(row)
        if not any(values):
            continue
        if len(values) != 7:
            raise RuntimeError(f"Bolzano applicant row-shape drift at row {source_row}: {values!r}")
        name, office, secondary, identifier_raw, activities_raw, application_raw, outcome_raw = values
        if not name or not application_raw:
            raise RuntimeError(f"Bolzano incomplete applicant row {source_row}: {values!r}")
        if outcome_raw:
            raise RuntimeError(f"Bolzano applicant outcome unexpectedly nonblank at row {source_row}: {outcome_raw!r}")
        application_date = _strict_source_date(
            application_raw,
            reviewed_malformed=_REVIEWED_APPLICANT_DATE_TYPOGRAPHY,
            label="application",
        )
        exact_key = (name, office, secondary, identifier_raw, activities_raw, application_raw, outcome_raw)
        if exact_key in exact_seen:
            raise RuntimeError(f"Bolzano unexpected exact applicant duplicate: {exact_key!r}")
        exact_seen.add(exact_key)
        rows.append(
            {
                "source_row": source_row,
                "name": name,
                "office": office,
                "secondary": secondary,
                "identifier_raw": identifier_raw,
                "activities_raw": activities_raw,
                "application_raw": application_raw,
                "application_date": application_date,
                "outcome_raw": outcome_raw,
            }
        )

    if len(rows) != _EXPECTED_APPLICANT_ROWS or len(exact_seen) != _EXPECTED_APPLICANT_RECORDS:
        raise RuntimeError(
            f"Bolzano applicant population drift: rows={len(rows)}, exact={len(exact_seen)}"
        )

    records: list[dict[str, Any]] = []
    for ordinal, row in enumerate(rows, start=1):
        activities = [row["activities_raw"]] if row["activities_raw"] else []
        record = _record(
            cfg,
            ordinal,
            name=row["name"],
            office=row["office"],
            secondary=row["secondary"],
            identifier_raw=row["identifier_raw"],
            activities=activities,
            status="pending",
            outcome_raw=row["outcome_raw"],
            application_date=row["application_date"],
            primary_date_label="Data presentazione istanza",
            source_fields={
                "source_row": row["source_row"],
                "activities_raw": row["activities_raw"],
                "application_date_raw": row["application_raw"],
                "outcome_raw": row["outcome_raw"],
            },
        )
        record["identifiers"] = _strict_identifiers(row["identifier_raw"])
        records.append(record)

    status_counts = {"pending": len(records)}
    if status_counts != _EXPECTED_APPLICANT_STATUS_COUNTS:
        raise RuntimeError(f"Bolzano applicant status drift: {status_counts!r}")
    return ParsedBatch(
        records=records,
        diagnostics={
            "source_rows": len(rows),
            "public_records": len(records),
            "status_counts": status_counts,
            "malformed_application_values": sorted(
                {row["application_raw"] for row in rows if not row["application_date"]}
            ),
        },
    )
