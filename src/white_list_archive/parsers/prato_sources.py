from __future__ import annotations

import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

import openpyxl

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record


_LISTED_REFERENCE_DATE = "2026-09-21"
_APPLICANT_REFERENCE_DATE = "2026-09-21"
_ROMAN = ("I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X")
_LISTED_SHEETS = tuple(f"SEZ. {roman}" for roman in _ROMAN)
_LISTED_HEADER = (
    "Denominazione / Ragione Sociale / Ditta",
    "Sede Legale",
    "Sede Secondaria con rappresentanza stabile in italia",
    "Codice fiscale / Partita IVA",
    "Data iscrizione",
    "Data scadenza",
    "Aggiornamento in corso",
)
_SECTION_COUNTS = {1: 32, 2: 18, 3: 31, 4: 22, 5: 33, 6: 39, 7: 1, 8: 2, 9: 11, 10: 47}
_UPDATE_RENEWAL = "IN CORSO ISTRUTTORIA PER RINNOVO ISCRIZIONE"
_APPLICANT_HEADER = (
    "Ragione Sociale",
    "Sede legale",
    "Sede secondaria con rappresentanza stabile in Italia",
    "Codice fiscale / Partita IVA",
    "Attività per cui è richiesta l’iscrizione",
    "Data presentazione istanza",
    "Esito",
)
_DMY = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")


def _require_cfg(
    cfg: dict[str, Any],
    *,
    source_key: str,
    population_scope: str,
    reference_date: str,
) -> None:
    if cfg.get("reference_date") != reference_date:
        raise RuntimeError(
            f"Prato reference-date drift: expected {reference_date}, got {cfg.get('reference_date')!r}"
        )
    if cfg.get("source_key") != source_key:
        raise RuntimeError(f"Prato source-key drift: expected {source_key!r}, got {cfg.get('source_key')!r}")
    if cfg.get("population_scope") != population_scope:
        raise RuntimeError(
            f"Prato population-scope drift: expected {population_scope!r}, got {cfg.get('population_scope')!r}"
        )


def _source_date(value: Any) -> tuple[str, str]:
    if isinstance(value, datetime):
        parsed = value.date()
        return parsed.isoformat(), parsed.isoformat()
    if isinstance(value, date):
        return value.isoformat(), value.isoformat()
    raw = _clean(value)
    match = _DMY.fullmatch(raw)
    if match is None:
        return "", raw
    day, month, year = map(int, match.groups())
    try:
        parsed = date(year, month, day)
    except ValueError:
        return "", raw
    return parsed.isoformat(), raw


def _section_activity(raw: str, roman: str) -> str:
    value = _clean(raw)
    match = re.fullmatch(rf"SEZIONE\s+{re.escape(roman)}\s*[–-]\s*(.+)", value, flags=re.I)
    if match is None:
        raise ValueError(f"Prato section-heading drift for {roman}: {value!r}")
    activity = _clean(match.group(1))
    if not activity:
        raise ValueError(f"Prato section heading lost its activity for {roman}")
    return activity


def _applicant_activities(raw: str) -> list[str]:
    source = _clean(raw)
    if not source:
        raise ValueError("Prato applicant activity field is unexpectedly blank")
    body = re.sub(r"^[-–]\s*", "", source).strip()
    values = [_clean(value) for value in re.split(r"\s+[-–]\s+", body) if _clean(value)]
    if not values:
        raise ValueError(f"Prato applicant activity field could not be tokenised: {source!r}")
    return values


def _assert_no_extra_values(ws: Any, *, logical_columns: int = 7) -> None:
    for row in range(1, ws.max_row + 1):
        for col in range(logical_columns + 1, ws.max_column + 1):
            value = _clean(ws.cell(row, col).value)
            if value:
                raise ValueError(
                    f"Prato listed unexpected value outside audited seven-column table "
                    f"at {ws.title}!R{row}C{col}: {value!r}"
                )


def parse_prato_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _require_cfg(
        cfg,
        source_key="prato-listed",
        population_scope="listed",
        reference_date=_LISTED_REFERENCE_DATE,
    )
    book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if tuple(book.sheetnames) != _LISTED_SHEETS:
        raise ValueError(f"Prato listed sheet-set drift: expected {_LISTED_SHEETS!r}, got {tuple(book.sheetnames)!r}")

    physical_rows: list[dict[str, Any]] = []
    observed_section_counts: Counter[int] = Counter()
    for section, roman in enumerate(_ROMAN, 1):
        ws = book[f"SEZ. {roman}"]
        _assert_no_extra_values(ws)
        header = tuple(_clean(ws.cell(5, col).value) for col in range(1, 8))
        if header != _LISTED_HEADER:
            raise ValueError(f"Prato listed header drift in {ws.title}: {header!r}")
        section_heading = _clean(ws.cell(4, 1).value)
        section_activity = _section_activity(section_heading, roman)
        if any(_clean(ws.cell(4, col).value) for col in range(2, 8)):
            raise ValueError(f"Prato section-heading row gained unexpected values in {ws.title}")

        for row_number in range(6, ws.max_row + 1):
            cells = [ws.cell(row_number, col).value for col in range(1, 8)]
            cleaned = [_clean(value) for value in cells]
            if not any(cleaned):
                continue
            name, office, secondary, identifier_raw, _listing_raw, _expiry_raw, update = cleaned
            if not name or not office or not identifier_raw:
                raise ValueError(f"Prato listed identity/layout drift at {ws.title}!R{row_number}: {cleaned!r}")
            if update not in {"", _UPDATE_RENEWAL}:
                raise ValueError(f"Prato listed update vocabulary drift at {ws.title}!R{row_number}: {update!r}")
            listing_date, listing_raw = _source_date(cells[4])
            expiry_date, expiry_raw = _source_date(cells[5])
            if not listing_date or not expiry_date:
                raise ValueError(
                    f"Prato listed date drift at {ws.title}!R{row_number}: "
                    f"listing={listing_raw!r}, expiry={expiry_raw!r}"
                )
            physical_rows.append(
                {
                    "section": section,
                    "section_label": f"Sezione {roman}",
                    "section_heading": section_heading,
                    "section_activity": section_activity,
                    "physical_locator": f"{ws.title}:R{row_number}",
                    "name": name,
                    "office": office,
                    "secondary": secondary,
                    "identifier_raw": identifier_raw,
                    "listing_date": listing_date,
                    "listing_raw": listing_raw,
                    "expiry_date": expiry_date,
                    "expiry_raw": expiry_raw,
                    "update": update,
                }
            )
            observed_section_counts[section] += 1

    if dict(observed_section_counts) != _SECTION_COUNTS:
        raise ValueError(
            f"Prato listed section-count drift: {dict(observed_section_counts)!r} != {_SECTION_COUNTS!r}"
        )
    if len(physical_rows) != 236:
        raise ValueError(f"Prato listed physical-row drift: {len(physical_rows)} != 236")

    grouped: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in physical_rows:
        key = tuple(
            str(row[field]).casefold()
            for field in (
                "name",
                "office",
                "secondary",
                "identifier_raw",
                "listing_date",
                "expiry_date",
                "update",
            )
        )
        group = grouped.setdefault(
            key,
            {
                "row": row,
                "sections": [],
                "section_headings": [],
                "activities": [],
                "physical_locators": [],
            },
        )
        for field, value in (
            ("sections", row["section_label"]),
            ("section_headings", row["section_heading"]),
            ("activities", row["section_activity"]),
            ("physical_locators", row["physical_locator"]),
        ):
            if value not in group[field]:
                group[field].append(value)

    records: list[dict[str, Any]] = []
    for group in grouped.values():
        row = group["row"]
        records.append(
            _record(
                cfg,
                len(records) + 1,
                name=row["name"],
                office=row["office"],
                secondary=row["secondary"],
                identifier_raw=row["identifier_raw"],
                activities=group["activities"],
                status="renewal_update_in_progress" if row["update"] else "listed",
                outcome_raw=row["update"],
                listing_date=row["listing_date"],
                expiry_date=row["expiry_date"],
                primary_date_label="Data iscrizione",
                source_fields={
                    "sections": group["sections"],
                    "physical_locators": group["physical_locators"],
                    "listing_date_raw_variants": [row["listing_raw"]],
                    "expiry_date_raw_variants": [row["expiry_raw"]],
                    "in_aggiornamento": row["update"],
                },
            )
        )

    diagnostics = {
        "parser": "prato_openxml_listed",
        "physical_sector_rows": len(physical_rows),
        "public_records": len(records),
        "section_counts": dict(observed_section_counts),
        "section_memberships": sum(len(record["source_fields"]["sections"]) for record in records),
        "status_counts": dict(Counter(record["source_status"] for record in records)),
        "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
        "raw_identifier_only": sum(
            bool(record["identifier_field_raw"]) and not record["identifiers"] for record in records
        ),
    }
    expected = {
        "physical_sector_rows": 236,
        "public_records": 137,
        "section_memberships": 236,
        "status_counts": {"listed": 123, "renewal_update_in_progress": 14},
        "identifier_coverage": 137,
        "raw_identifier_only": 0,
    }
    for key, value in expected.items():
        if diagnostics[key] != value:
            raise ValueError(f"Prato listed audited-boundary drift for {key}: {diagnostics[key]!r} != {value!r}")
    return ParsedBatch(records, diagnostics)


def _antiword_paragraphs(path: Path) -> list[str]:
    executable = shutil.which("antiword")
    if executable is None:
        raise RuntimeError("Prato applicant parsing requires antiword")
    process = subprocess.run(
        [executable, "-x", "db", str(path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
    )
    if process.returncode:
        raise RuntimeError(f"antiword failed for Prato applicant source: {process.stderr.decode('utf-8', 'replace').strip()}")
    xml = re.sub(rb'<!DOCTYPE book PUBLIC.*?docbookx\.dtd">\s*', b"", process.stdout, flags=re.S)
    root = ET.fromstring(xml)
    paragraphs: list[str] = []
    for node in root.iter():
        if node.tag.rsplit("}", 1)[-1].lower() != "para":
            continue
        text = _clean(" ".join(node.itertext()))
        if text:
            paragraphs.append(text)
    return paragraphs


def _applicant_rows(path: Path) -> list[list[str]]:
    paragraphs = _antiword_paragraphs(path)
    if len(paragraphs) != 21:
        raise ValueError(f"Prato applicant paragraph-count drift: {len(paragraphs)} != 21")
    try:
        header_index = next(
            index for index, text in enumerate(paragraphs) if "ragione sociale" in text.casefold()
        )
    except StopIteration as exc:
        raise ValueError("Prato applicant header paragraph not found") from exc
    if header_index != 5:
        raise ValueError(f"Prato applicant header-paragraph drift: {header_index} != 5")
    anchored = " ".join(paragraphs[header_index:])
    position = anchored.casefold().find("ragione sociale")
    if position < 0:
        raise ValueError("Prato applicant header anchor disappeared")
    tokens = [_clean(value) for value in anchored[position:].split("|")]
    if len(tokens) != 193:
        raise ValueError(f"Prato applicant pipe-token drift: {len(tokens)} != 193")
    header = tuple(tokens[:7])
    if header != _APPLICANT_HEADER:
        raise ValueError(f"Prato applicant header drift: {header!r}")

    cursor = 7
    while cursor < len(tokens) and not tokens[cursor]:
        cursor += 1
    records: list[list[str]] = []
    separators: Counter[int] = Counter()
    while cursor < len(tokens):
        if not tokens[cursor]:
            cursor += 1
            continue
        if cursor + 6 >= len(tokens):
            tail = tokens[cursor:]
            if all(not value for value in tail):
                break
            raise ValueError(f"Prato applicant truncated token group at {cursor}: {tail!r}")
        row = tokens[cursor : cursor + 7]
        records.append(row)
        cursor += 7
        blanks = 0
        while cursor < len(tokens) and not tokens[cursor]:
            blanks += 1
            cursor += 1
        separators[blanks] += 1
    if len(records) != 23:
        raise ValueError(f"Prato applicant record-count drift: {len(records)} != 23")
    if dict(separators) != {1: 22, 2: 1}:
        raise ValueError(f"Prato applicant row-separator drift: {dict(separators)!r}")
    return records


def parse_prato_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _require_cfg(
        cfg,
        source_key="prato-applicants",
        population_scope="applicant",
        reference_date=_APPLICANT_REFERENCE_DATE,
    )
    rows = _applicant_rows(path)
    records: list[dict[str, Any]] = []
    for row_number, values in enumerate(rows, 1):
        name, office, secondary, identifier_raw, activity_raw, application_raw, outcome = values
        if not name or not office or not identifier_raw or not activity_raw or not application_raw:
            raise ValueError(f"Prato applicant identity/layout drift at logical row {row_number}: {values!r}")
        if outcome:
            raise ValueError(f"Prato applicant outcome vocabulary drift at logical row {row_number}: {outcome!r}")
        application_date, raw = _source_date(application_raw)
        if not application_date:
            raise ValueError(f"Prato applicant date drift at logical row {row_number}: {raw!r}")
        records.append(
            _record(
                cfg,
                row_number,
                name=name,
                office=office,
                secondary=secondary,
                identifier_raw=identifier_raw,
                activities=_applicant_activities(activity_raw),
                status="pending",
                application_date=application_date,
                primary_date_label="Data presentazione istanza",
                source_fields={
                    "requested_activities_source": activity_raw,
                    "application_date_raw": raw,
                },
            )
        )

    diagnostics = {
        "parser": "prato_legacy_applicants",
        "public_records": len(records),
        "status_counts": dict(Counter(record["source_status"] for record in records)),
        "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
        "raw_identifier_only": sum(
            bool(record["identifier_field_raw"]) and not record["identifiers"] for record in records
        ),
    }
    expected = {
        "public_records": 23,
        "status_counts": {"pending": 23},
        "identifier_coverage": 23,
        "raw_identifier_only": 0,
    }
    for key, value in expected.items():
        if diagnostics[key] != value:
            raise ValueError(f"Prato applicant audited-boundary drift for {key}: {diagnostics[key]!r} != {value!r}")
    return ParsedBatch(records, diagnostics)


PARSERS: dict[str, Callable[[Path, dict[str, Any]], ParsedBatch]] = {
    "prato_openxml_listed": parse_prato_listed,
    "prato_legacy_applicants": parse_prato_applicants,
}
