from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from docx import Document

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record

PARSER_VERSION = "1"
_REFERENCE_DATE = "2026-08-11"
_EXPECTED_SHEET = "Foglio1"
_EXPECTED_SECTION_ROWS = 882
_EXPECTED_LISTED_RECORDS = 323
_EXPECTED_LISTED_STATUS_COUNTS = {"listed": 276, "renewal_update_in_progress": 47}
_EXPECTED_APPLICANT_RECORDS = 19
_EXPECTED_APPLICANT_STATUS_COUNTS = {"pending": 19}
_EXPECTED_SECTIONS = {
    "I": "Estrazione, fornitura e trasporto di terra e materiali inerti",
    "II": "Confezionamento, fornitura e trasporto di calcestruzzo e di bitume",
    "III": "Noli a freddo di macchinari",
    "IV": "Fornitura di ferro lavorato",
    "V": "Noli a caldo",
    "VI": "Autotrasporto per conto di terzi",
    "VII": "Guardiania dei cantieri",
    "VIII": "Servizi funerari e cimiteriali",
    "IX": "Ristorazione, gestione delle mense e catering",
    "X": "Servizi ambientali",
}
_EXPECTED_SECTION_ROW_COUNTS = {
    "I": 134,
    "II": 60,
    "III": 199,
    "IV": 51,
    "V": 179,
    "VI": 107,
    "VII": 31,
    "VIII": 19,
    "IX": 19,
    "X": 83,
}
_SECTION_RE = re.compile(r"\bSEZIONE\s*(10|[1-9]|[IVX]+)\s*[°º^]?\b", re.I)
_DATE_DMY = re.compile(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{4})$")
_STRICT_IDENTIFIER = re.compile(r"^(?:\d{11}|[A-Za-z0-9]{16})$")
_SCOPED_UPDATE = "si (non richiesto per questa sezione)"
_ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V", 6: "VI", 7: "VII", 8: "VIII", 9: "IX", 10: "X"}


def _strict_identifiers(value: str) -> list[str]:
    raw = _clean(value).upper()
    return [raw] if _STRICT_IDENTIFIER.fullmatch(raw) else []


def _source_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _clean(value)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?: 00:00:00)?", raw):
        try:
            return date.fromisoformat(raw[:10]).isoformat()
        except ValueError as exc:
            raise RuntimeError(f"Campobasso invalid calendar date: {raw!r}") from exc
    match = _DATE_DMY.fullmatch(raw)
    if not match:
        raise RuntimeError(f"Campobasso unreviewed date typography: {raw!r}")
    day, month, year = map(int, match.groups())
    try:
        return date(year, month, day).isoformat()
    except ValueError as exc:
        raise RuntimeError(f"Campobasso invalid calendar date: {raw!r}") from exc


def _normalise_section(value: str) -> str | None:
    match = _SECTION_RE.search(_clean(value))
    if not match:
        return None
    token = match.group(1).upper().rstrip("°º^")
    if token.isdigit():
        token = _ROMAN.get(int(token), "")
    return token if token in _EXPECTED_SECTIONS else None


def _validate_cfg(cfg: dict[str, Any], expected_source_key: str) -> None:
    if cfg.get("source_key") != expected_source_key:
        raise RuntimeError(
            f"Campobasso parser/source mismatch: {cfg.get('source_key')!r} != {expected_source_key!r}"
        )
    if cfg.get("authority_key") != "campobasso":
        raise RuntimeError("Campobasso parser bound to a non-Campobasso authority")
    if cfg.get("reference_date") != _REFERENCE_DATE:
        raise RuntimeError(f"Campobasso reference date drift: {cfg.get('reference_date')!r}")


def _activity_matches(section: str, value: str) -> bool:
    folded = _clean(value).casefold()
    expected = _EXPECTED_SECTIONS[section].casefold()
    if section != "X":
        return expected in folded
    required = ("servizi ambientali", "rifiuti", "bonifica")
    return all(token in folded for token in required)


def parse_campobasso_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, "campobasso-listed")
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if workbook.sheetnames != [_EXPECTED_SHEET]:
        raise RuntimeError(f"Campobasso listed worksheet drift: {workbook.sheetnames!r}")
    worksheet = workbook[_EXPECTED_SHEET]

    current_section: str | None = None
    current_activity = ""
    in_table = False
    rows: list[dict[str, Any]] = []
    section_counts: Counter[str] = Counter()
    section_headers: Counter[str] = Counter()

    for source_row_number, raw_row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        values = list(raw_row)
        cleaned = [_clean(value) for value in values]
        section_tokens = [section for value in cleaned if (section := _normalise_section(value))]
        if section_tokens:
            if len(set(section_tokens)) != 1:
                raise RuntimeError(f"Campobasso ambiguous section marker at row {source_row_number}: {section_tokens!r}")
            current_section = section_tokens[0]
            current_activity = ""
            in_table = False
            continue

        joined = " | ".join(cleaned).casefold()
        if current_section and not in_table and cleaned:
            nonempty = [value for value in cleaned if value]
            candidates = [
                value for value in nonempty
                if "prefettura" not in value.casefold()
                and "elenco dei fornitori" not in value.casefold()
                and "ministero" not in value.casefold()
                and "ragione sociale" not in value.casefold()
            ]
            for candidate in candidates:
                if _activity_matches(current_section, candidate):
                    current_activity = candidate
                    break

        if all(token in joined for token in ("ragione sociale", "codice fiscale", "data di iscrizione", "aggiornamento in corso")):
            if not current_section:
                raise RuntimeError(f"Campobasso listed header before section at row {source_row_number}")
            if not current_activity or not _activity_matches(current_section, current_activity):
                raise RuntimeError(f"Campobasso section {current_section} activity heading changed")
            section_headers[current_section] += 1
            in_table = True
            continue

        if not in_table or not current_section:
            continue
        # Current official XLSX has an empty leading column A and data in B:I.
        padded = (cleaned[1:9] + [""] * 8)[:8]
        name, office, province, secondary, identifier_raw, listing_raw, expiry_raw, update_raw = padded
        if not any(padded):
            continue
        # Section break/title rows are processed above. Other incomplete rows are not
        # silently coerced into company observations.
        if not name or not listing_raw or not expiry_raw:
            continue
        listing_date = _source_date(values[6] if len(values) > 6 else listing_raw)
        expiry_date = _source_date(values[7] if len(values) > 7 else expiry_raw)
        update_folded = update_raw.casefold()
        if update_folded not in ("", "si", _SCOPED_UPDATE):
            raise RuntimeError(f"Campobasso unreviewed update status at row {source_row_number}: {update_raw!r}")
        rows.append(
            {
                "section": current_section,
                "activity": current_activity,
                "source_row": source_row_number,
                "name": name,
                "office": office,
                "province": province,
                "secondary": secondary,
                "identifier_raw": identifier_raw,
                "listing_raw": listing_raw,
                "expiry_raw": expiry_raw,
                "listing_date": listing_date,
                "expiry_date": expiry_date,
                "update_raw": update_raw,
            }
        )
        section_counts[current_section] += 1

    if section_headers != Counter({section: 1 for section in _EXPECTED_SECTIONS}):
        raise RuntimeError(f"Campobasso listed header structure drift: {dict(section_headers)!r}")
    if dict(section_counts) != _EXPECTED_SECTION_ROW_COUNTS:
        raise RuntimeError(f"Campobasso section-row drift: {dict(section_counts)!r}")
    if len(rows) != _EXPECTED_SECTION_ROWS:
        raise RuntimeError(f"Campobasso listed source-row drift: {len(rows)} != {_EXPECTED_SECTION_ROWS}")

    grouped: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in rows:
        # Province is intentionally absent from this key. The audited source uses
        # literal placeholders such as '==' and 'CB' inconsistently for otherwise
        # exact repetitions. Registered-office differences remain identity-relevant
        # and are never fuzzy-merged.
        key = (
            row["name"], row["office"], row["secondary"], row["identifier_raw"],
            row["listing_date"], row["expiry_date"],
        )
        group = grouped.setdefault(
            key,
            {"row": row, "sections": [], "activities": [], "provinces": [], "updates": [], "listing_raw": [], "expiry_raw": []},
        )
        if row["section"] in group["sections"]:
            raise RuntimeError(f"Campobasso duplicate observation inside section {row['section']}: {key!r}")
        for field, value in (
            ("sections", row["section"]), ("activities", row["activity"]),
            ("provinces", row["province"]), ("updates", row["update_raw"]),
            ("listing_raw", row["listing_raw"]), ("expiry_raw", row["expiry_raw"]),
        ):
            if value and value not in group[field]:
                group[field].append(value)

    if len(grouped) != _EXPECTED_LISTED_RECORDS:
        raise RuntimeError(f"Campobasso listed semantic-group drift: {len(grouped)} != {_EXPECTED_LISTED_RECORDS}")

    records: list[dict[str, Any]] = []
    for group in grouped.values():
        row = group["row"]
        updates = [value.casefold() for value in group["updates"]]
        has_plain_si = "si" in updates
        scoped_only = bool(updates) and not has_plain_si and all(value == _SCOPED_UPDATE for value in updates)
        if scoped_only:
            raise RuntimeError(f"Campobasso scoped update note without source-explicit plain SI: {row['name']!r}")
        status = "renewal_update_in_progress" if has_plain_si else "listed"
        record = _record(
            cfg,
            len(records) + 1,
            name=row["name"],
            office=row["office"],
            secondary=row["secondary"],
            identifier_raw=row["identifier_raw"],
            activities=group["activities"],
            status=status,
            outcome_raw=" | ".join(group["updates"]),
            listing_date=row["listing_date"],
            expiry_date=row["expiry_date"],
            primary_date_label="Data iscrizione",
            source_fields={
                "sections": group["sections"],
                "registered_office_variants": [row["office"]] if row["office"] else [],
                "secondary_office_variants": [row["secondary"]] if row["secondary"] else [],
                "listing_date_raw_variants": group["listing_raw"],
                "expiry_date_raw_variants": group["expiry_raw"],
                "requested_activities_source": " | ".join(group["activities"]),
            },
        )
        record["identifiers"] = _strict_identifiers(row["identifier_raw"])
        records.append(record)

    status_counts = dict(Counter(record["source_status"] for record in records))
    if status_counts != _EXPECTED_LISTED_STATUS_COUNTS:
        raise RuntimeError(f"Campobasso listed status drift: {status_counts!r}")

    return ParsedBatch(
        records=records,
        diagnostics={
            "parser": "campobasso_listed",
            "parser_version": PARSER_VERSION,
            "sector_rows": len(rows),
            "public_records": len(records),
            "status_counts": status_counts,
            "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
            "raw_identifier_only": sum(bool(record["identifier_field_raw"]) and not record["identifiers"] for record in records),
            "province_variant_groups": sum(len(group["provinces"]) > 1 for group in grouped.values()),
        },
    )


def parse_campobasso_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, "campobasso-applicants")
    document = Document(path)
    if len(document.tables) != 1:
        raise RuntimeError(f"Campobasso applicant table-count drift: {len(document.tables)} != 1")
    table = document.tables[0]
    records: list[dict[str, Any]] = []
    header_count = 0
    seen: set[tuple[str, ...]] = set()

    for source_row_number, source_row in enumerate(table.rows, start=1):
        values = [_clean(cell.text) for cell in source_row.cells]
        values += [""] * max(0, 6 - len(values))
        joined = " | ".join(values[:6]).casefold()
        if "ragione sociale" in joined and "data di presentazione" in joined:
            header_count += 1
            continue
        name, office, secondary, identifier_raw, activities_raw, application_raw = values[:6]
        if not any(values[:6]):
            continue
        if not name or not application_raw:
            raise RuntimeError(f"Campobasso incomplete applicant row {source_row_number}: {values[:6]!r}")
        application_date = _source_date(application_raw)
        key = (name, office, secondary, identifier_raw, activities_raw, application_date)
        if key in seen:
            raise RuntimeError(f"Campobasso duplicate applicant row: {key!r}")
        seen.add(key)
        record = _record(
            cfg,
            len(records) + 1,
            name=name,
            office=office,
            secondary=secondary,
            identifier_raw=identifier_raw,
            activities=[activities_raw] if activities_raw else [],
            status="pending",
            outcome_raw="",
            application_date=application_date,
            primary_date_label="Data presentazione istanza",
            source_fields={
                "application_date_raw_variants": [application_raw],
                "requested_activities_source": activities_raw,
            },
        )
        record["identifiers"] = _strict_identifiers(identifier_raw)
        records.append(record)

    if header_count != 2:
        raise RuntimeError(f"Campobasso applicant repeated-header drift: {header_count} != 2")
    if len(records) != _EXPECTED_APPLICANT_RECORDS:
        raise RuntimeError(f"Campobasso applicant row drift: {len(records)} != {_EXPECTED_APPLICANT_RECORDS}")
    status_counts = dict(Counter(record["source_status"] for record in records))
    if status_counts != _EXPECTED_APPLICANT_STATUS_COUNTS:
        raise RuntimeError(f"Campobasso applicant status drift: {status_counts!r}")

    return ParsedBatch(
        records=records,
        diagnostics={
            "parser": "campobasso_applicants",
            "parser_version": PARSER_VERSION,
            "public_records": len(records),
            "repeated_headers": header_count,
            "status_counts": status_counts,
            "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
            "raw_identifier_only": sum(bool(record["identifier_field_raw"]) and not record["identifiers"] for record in records),
        },
    )


PARSERS = {
    "campobasso_listed": parse_campobasso_listed,
    "campobasso_applicants": parse_campobasso_applicants,
}
