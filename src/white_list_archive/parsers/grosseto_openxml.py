from __future__ import annotations

import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record

LISTED_PARSER_NAME = "grosseto_listed"
APPLICANT_PARSER_NAME = "grosseto_applicants"
PARSER_VERSION = "1"
_REFERENCE_DATE = "2026-09-19"

_EXPECTED_LISTED_SHAPE = (800, 19)
_EXPECTED_APPLICANT_SHAPE = (72, 11)
_EXPECTED_LISTED_HEADER_ROWS = {1: 6, 2: 132, 3: 213, 4: 342, 5: 394, 6: 531, 7: 618, 8: 656, 9: 668, 10: 695}
_EXPECTED_LISTED_END_ROWS = {1: 129, 2: 210, 3: 339, 4: 391, 5: 528, 6: 615, 7: 653, 8: 665, 9: 692, 10: 800}
_EXPECTED_APPLICANT_HEADER_ROWS = {1: 5, 2: 12, 3: 19, 4: 30, 5: 37, 6: 46, 7: 52, 8: 56, 9: 59, 10: 65}
_EXPECTED_APPLICANT_END_ROWS = {1: 9, 2: 16, 3: 27, 4: 34, 5: 43, 6: 49, 7: 53, 8: 56, 9: 62, 10: 72}
_EXPECTED_LISTED_MARKERS = ((3, 1), (4, 1), (130, 2), (211, 3), (340, 4), (392, 5), (529, 6), (616, 7), (654, 8), (666, 9), (693, 10))
_EXPECTED_APPLICANT_MARKERS = ((3, 1), (10, 2), (17, 3), (28, 4), (35, 5), (44, 6), (50, 7), (54, 8), (57, 9), (63, 10))
_EXPECTED_LISTED_SECTOR_COUNTS = {1: 123, 2: 75, 3: 125, 4: 49, 5: 133, 6: 84, 7: 35, 8: 9, 9: 24, 10: 105}
_EXPECTED_APPLICANT_SECTOR_COUNTS = {1: 4, 2: 4, 3: 7, 4: 4, 5: 5, 6: 2, 7: 1, 9: 2, 10: 6}
_EXPECTED_LISTED_SOURCE_ROWS = 762
_EXPECTED_APPLICANT_SOURCE_ROWS = 35
_EXPECTED_LISTED_RECORDS = 406
_EXPECTED_APPLICANT_RECORDS = 12
_EXPECTED_LISTED_STATUS_COUNTS = {"listed": 328, "renewal_update_in_progress": 78}
_EXPECTED_APPLICANT_STATUS_COUNTS = {"pending": 12}
_EXPECTED_LISTED_IDENTIFIER_COVERAGE = 371
_EXPECTED_APPLICANT_IDENTIFIER_COVERAGE = 12
_EXPECTED_LISTED_BLANK_LISTING_DATES = 8
_EXPECTED_LISTED_BLANK_EXPIRY_DATES = 8

_SECTION_LABELS = {
    1: "Estrazione, fornitura e trasporto di terra e materiali inerti",
    2: "Confezionamento, fornitura e trasporto di calcestruzzo e di bitume",
    3: "Noli a freddo di macchinari",
    4: "Fornitura di ferro lavorato",
    5: "Noli a caldo",
    6: "Autotrasporti per conto di terzi",
    7: "Guardiania dei Cantieri",
    8: "Servizi funerari e cimiteriali",
    9: "Ristorazione, gestione delle mense e catering",
    10: "Servizi ambientali",
}
_ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10}
_STRICT_IDENTIFIER = re.compile(r"^(?:\d{11}|[A-Z]{6}[0-9A-Z]{10})$")
_LISTED_PLAIN_NOTES = {"", "iscrizione", "iscritta", "iscizione", "iscrzione"}


def _source_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _clean(value)
    if not raw:
        return ""
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?: 00:00:00)?", raw):
        try:
            return date.fromisoformat(raw[:10]).isoformat()
        except ValueError:
            return ""
    match = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", raw)
    if match:
        day, month, year = map(int, match.groups())
        try:
            return date(year, month, day).isoformat()
        except ValueError:
            return ""
    return ""


def _structured_identifier(raw: str) -> bool:
    return bool(_STRICT_IDENTIFIER.fullmatch(_clean(raw).upper().replace(" ", "")))


def _listed_status(note: str) -> str:
    folded = _clean(note).casefold()
    if "rinnovo" in folded:
        return "renewal_update_in_progress"
    if folded in _LISTED_PLAIN_NOTES:
        return "listed"
    raise RuntimeError(f"Grosseto listed note/status vocabulary changed: {note!r}")


def _section_marker(value: Any) -> int | None:
    match = re.fullmatch(r"Sezione\s+([IVX]+)", _clean(value), re.I)
    return _ROMAN.get(match.group(1).upper()) if match else None


def _raw_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _clean(value)


def _workbook(path: Path, *, population: str):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if workbook.sheetnames != ["Foglio1"]:
        raise RuntimeError(f"Grosseto {population} workbook sheet set changed: {workbook.sheetnames!r}")
    worksheet = workbook["Foglio1"]
    expected_shape = _EXPECTED_LISTED_SHAPE if population == "listed" else _EXPECTED_APPLICANT_SHAPE
    shape = (worksheet.max_row, worksheet.max_column)
    if shape != expected_shape:
        raise RuntimeError(f"Grosseto {population} workbook shape changed: expected {expected_shape!r}, got {shape!r}")
    rows = {number: list(values) for number, values in enumerate(worksheet.iter_rows(values_only=True), 1)}
    return rows, shape


def _validate_layout(rows: dict[int, list[Any]], *, population: str) -> None:
    if population == "listed":
        marker_expected = _EXPECTED_LISTED_MARKERS
        header_rows = _EXPECTED_LISTED_HEADER_ROWS
        expected_header = [
            "ragione sociale",
            "sede legale",
            "sede secondaria con rappresentanza stabile in italia",
            "codice fiscale partita iva",
            "data di iscrizione",
            "data scadenza iscrizione",
            "aggiornamento in corso",
        ]
    else:
        marker_expected = _EXPECTED_APPLICANT_MARKERS
        header_rows = _EXPECTED_APPLICANT_HEADER_ROWS
        expected_header = [
            "ragione sociale",
            "sede legale",
            "sede secondaria con rappresentanza stabile in italia",
            "codice fiscale partita iva",
            "data presentazione domanda",
        ]

    observed_markers = tuple(
        (row_number, section)
        for row_number, values in rows.items()
        if (section := _section_marker(values[0] if values else "")) is not None
    )
    if observed_markers != marker_expected:
        raise RuntimeError(f"Grosseto {population} section-marker layout changed: {observed_markers!r}")

    for section, header_row in header_rows.items():
        values = rows[header_row]
        header = [re.sub(r"\s+", " ", _clean(value).casefold()) for value in values[: len(expected_header)]]
        if header != expected_header:
            raise RuntimeError(
                f"Grosseto {population} section {section} header changed: expected {expected_header!r}, got {header!r}"
            )
        title = _clean(rows[header_row - 1][0])
        expected_title = _SECTION_LABELS[section]
        if section == 10:
            if not title.startswith(expected_title):
                raise RuntimeError(f"Grosseto {population} section {section} title changed: {title!r}")
        elif re.sub(r"\s+", "", title).casefold() != re.sub(r"\s+", "", expected_title).casefold():
            raise RuntimeError(f"Grosseto {population} section {section} title changed: {title!r}")


def _listed_sector_rows(rows: dict[int, list[Any]]) -> list[dict[str, Any]]:
    source_rows: list[dict[str, Any]] = []
    for section, header_row in _EXPECTED_LISTED_HEADER_ROWS.items():
        for row_number in range(header_row + 1, _EXPECTED_LISTED_END_ROWS[section] + 1):
            values = list(rows[row_number])
            if not any(_clean(value) for value in values):
                continue
            if any(_clean(value) for value in values[7:]):
                raise RuntimeError(f"Grosseto listed unexpected trailing content at row {row_number}")
            name, office, secondary, identifier_raw, listing_value, expiry_value, note_value = values[:7]
            name = _clean(name)
            if not name:
                raise RuntimeError(f"Grosseto listed identity unexpectedly blank at row {row_number}")
            listing_raw = _raw_date(listing_value)
            expiry_raw = _raw_date(expiry_value)
            listing_date = _source_date(listing_value)
            expiry_date = _source_date(expiry_value)
            if listing_raw and not listing_date:
                raise RuntimeError(f"Grosseto listed unreviewed listing-date typography at row {row_number}: {listing_raw!r}")
            if expiry_raw and not expiry_date:
                raise RuntimeError(f"Grosseto listed unreviewed expiry-date typography at row {row_number}: {expiry_raw!r}")
            note = _clean(note_value)
            source_rows.append(
                {
                    "section": section,
                    "row": row_number,
                    "name": name,
                    "office": _clean(office),
                    "secondary": _clean(secondary),
                    "identifier_raw": _clean(identifier_raw),
                    "listing_raw": listing_raw,
                    "listing_date": listing_date,
                    "expiry_raw": expiry_raw,
                    "expiry_date": expiry_date,
                    "note": note,
                    "status": _listed_status(note),
                }
            )
    return source_rows


def _applicant_sector_rows(rows: dict[int, list[Any]]) -> list[dict[str, Any]]:
    source_rows: list[dict[str, Any]] = []
    for section, header_row in _EXPECTED_APPLICANT_HEADER_ROWS.items():
        end_row = _EXPECTED_APPLICANT_END_ROWS[section]
        if end_row <= header_row:
            continue
        for row_number in range(header_row + 1, end_row + 1):
            values = list(rows[row_number])
            if not any(_clean(value) for value in values):
                continue
            if any(_clean(value) for value in values[5:]):
                raise RuntimeError(f"Grosseto applicant unexpected trailing content at row {row_number}")
            name, office, secondary, identifier_raw, application_value = values[:5]
            name = _clean(name)
            if not name:
                raise RuntimeError(f"Grosseto applicant identity unexpectedly blank at row {row_number}")
            application_raw = _raw_date(application_value)
            application_date = _source_date(application_value)
            if not application_raw or not application_date:
                raise RuntimeError(
                    f"Grosseto applicant application-date boundary changed at row {row_number}: {application_raw!r}"
                )
            source_rows.append(
                {
                    "section": section,
                    "row": row_number,
                    "name": name,
                    "office": _clean(office),
                    "secondary": _clean(secondary),
                    "identifier_raw": _clean(identifier_raw),
                    "application_raw": application_raw,
                    "application_date": application_date,
                }
            )
    return source_rows


def _validate_cfg(cfg: dict[str, Any], source_key: str) -> None:
    if cfg.get("source_key") != source_key or cfg.get("authority_key") != "grosseto":
        raise RuntimeError(f"Grosseto parser/source binding changed: {cfg.get('source_key')!r}")
    if cfg.get("reference_date") != _REFERENCE_DATE:
        raise RuntimeError(f"Grosseto reference boundary changed: {cfg.get('reference_date')!r}")


def parse_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, "grosseto-listed")
    rows, shape = _workbook(path, population="listed")
    _validate_layout(rows, population="listed")
    sector_rows = _listed_sector_rows(rows)

    if len(sector_rows) != _EXPECTED_LISTED_SOURCE_ROWS:
        raise RuntimeError(
            f"Grosseto listed source-row denominator changed: {len(sector_rows)} != {_EXPECTED_LISTED_SOURCE_ROWS}"
        )
    section_counts = dict(Counter(row["section"] for row in sector_rows))
    if section_counts != _EXPECTED_LISTED_SECTOR_COUNTS:
        raise RuntimeError(f"Grosseto listed section counts changed: {section_counts!r}")

    grouped: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in sector_rows:
        key = (
            row["name"].casefold(),
            row["identifier_raw"].casefold(),
            row["office"].casefold(),
            row["secondary"].casefold(),
            row["listing_date"],
            row["expiry_date"],
            row["status"],
            row["note"].casefold(),
        )
        group = grouped.setdefault(key, {"row": row, "sections": [], "source_rows": []})
        if row["section"] in group["sections"]:
            raise RuntimeError(
                f"Grosseto listed duplicate same-section source row in one observation: {row['name']!r}, section {row['section']}"
            )
        group["sections"].append(row["section"])
        group["source_rows"].append(row["row"])

    if len(grouped) != _EXPECTED_LISTED_RECORDS:
        raise RuntimeError(
            f"Grosseto listed grouped-observation denominator changed: {len(grouped)} != {_EXPECTED_LISTED_RECORDS}"
        )

    output: list[dict[str, Any]] = []
    for group in grouped.values():
        row = group["row"]
        sections = sorted(group["sections"])
        record = _record(
            cfg,
            len(output) + 1,
            name=row["name"],
            office=row["office"],
            secondary=row["secondary"],
            identifier_raw=row["identifier_raw"],
            activities=[f"Sezione {section}" for section in sections],
            status=row["status"],
            outcome_raw=row["note"],
            listing_date=row["listing_date"],
            expiry_date=row["expiry_date"],
            primary_date_label="Data iscrizione",
            source_fields={
                "sections": [f"Sezione {section}" for section in sections],
                "listing_date_raw_variants": [row["listing_raw"]] if row["listing_raw"] else [],
                "expiry_date_raw_variants": [row["expiry_raw"]] if row["expiry_raw"] else [],
                "in_aggiornamento": row["note"] if row["status"] == "renewal_update_in_progress" else "",
            },
        )
        record["identifiers"] = (
            [_clean(row["identifier_raw"]).upper().replace(" ", "")]
            if _structured_identifier(row["identifier_raw"])
            else []
        )
        output.append(record)

    statuses = dict(Counter(record["source_status"] for record in output))
    if statuses != _EXPECTED_LISTED_STATUS_COUNTS:
        raise RuntimeError(f"Grosseto listed grouped status counts changed: {statuses!r}")
    identifier_coverage = sum(bool(record["identifiers"]) for record in output)
    if identifier_coverage != _EXPECTED_LISTED_IDENTIFIER_COVERAGE:
        raise RuntimeError(f"Grosseto listed identifier coverage changed: {identifier_coverage}")
    blank_listing = sum(not record["observed_listing_date"] for record in output)
    blank_expiry = sum(not record["observed_expiry_date"] for record in output)
    if blank_listing != _EXPECTED_LISTED_BLANK_LISTING_DATES or blank_expiry != _EXPECTED_LISTED_BLANK_EXPIRY_DATES:
        raise RuntimeError(
            f"Grosseto listed reviewed blank-date boundary changed: listing={blank_listing}, expiry={blank_expiry}"
        )

    return ParsedBatch(
        output,
        {
            "parser": LISTED_PARSER_NAME,
            "parser_version": PARSER_VERSION,
            "workbook_shape": shape,
            "sector_rows": len(sector_rows),
            "section_counts": section_counts,
            "public_records": len(output),
            "status_counts": statuses,
            "identifier_coverage": identifier_coverage,
            "raw_identifier_only": len(output) - identifier_coverage,
            "blank_listing_dates": blank_listing,
            "blank_expiry_dates": blank_expiry,
        },
    )


def parse_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, "grosseto-applicants")
    rows, shape = _workbook(path, population="applicant")
    _validate_layout(rows, population="applicant")
    sector_rows = _applicant_sector_rows(rows)

    if len(sector_rows) != _EXPECTED_APPLICANT_SOURCE_ROWS:
        raise RuntimeError(
            f"Grosseto applicant source-row denominator changed: {len(sector_rows)} != {_EXPECTED_APPLICANT_SOURCE_ROWS}"
        )
    section_counts = dict(Counter(row["section"] for row in sector_rows))
    if section_counts != _EXPECTED_APPLICANT_SECTOR_COUNTS:
        raise RuntimeError(f"Grosseto applicant section counts changed: {section_counts!r}")

    grouped: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in sector_rows:
        key = (
            row["name"].casefold(),
            row["identifier_raw"].casefold(),
            row["office"].casefold(),
            row["secondary"].casefold(),
            row["application_date"],
        )
        group = grouped.setdefault(key, {"row": row, "sections": [], "source_rows": []})
        if row["section"] in group["sections"]:
            raise RuntimeError(
                f"Grosseto applicant duplicate same-section source row in one observation: {row['name']!r}, section {row['section']}"
            )
        group["sections"].append(row["section"])
        group["source_rows"].append(row["row"])

    if len(grouped) != _EXPECTED_APPLICANT_RECORDS:
        raise RuntimeError(
            f"Grosseto applicant grouped-observation denominator changed: {len(grouped)} != {_EXPECTED_APPLICANT_RECORDS}"
        )

    output: list[dict[str, Any]] = []
    for group in grouped.values():
        row = group["row"]
        sections = sorted(group["sections"])
        record = _record(
            cfg,
            len(output) + 1,
            name=row["name"],
            office=row["office"],
            secondary=row["secondary"],
            identifier_raw=row["identifier_raw"],
            activities=[f"Sezione {section}" for section in sections],
            status="pending",
            application_date=row["application_date"],
            primary_date_label="Data presentazione istanza",
            source_fields={
                "requested_activities_source": " · ".join(f"Sezione {section}" for section in sections),
                "application_date_raw_variants": [row["application_raw"]],
            },
        )
        record["identifiers"] = (
            [_clean(row["identifier_raw"]).upper().replace(" ", "")]
            if _structured_identifier(row["identifier_raw"])
            else []
        )
        output.append(record)

    statuses = dict(Counter(record["source_status"] for record in output))
    if statuses != _EXPECTED_APPLICANT_STATUS_COUNTS:
        raise RuntimeError(f"Grosseto applicant grouped status counts changed: {statuses!r}")
    identifier_coverage = sum(bool(record["identifiers"]) for record in output)
    if identifier_coverage != _EXPECTED_APPLICANT_IDENTIFIER_COVERAGE:
        raise RuntimeError(f"Grosseto applicant identifier coverage changed: {identifier_coverage}")

    return ParsedBatch(
        output,
        {
            "parser": APPLICANT_PARSER_NAME,
            "parser_version": PARSER_VERSION,
            "workbook_shape": shape,
            "sector_rows": len(sector_rows),
            "section_counts": section_counts,
            "public_records": len(output),
            "status_counts": statuses,
            "identifier_coverage": identifier_coverage,
        },
    )


PARSERS = {
    LISTED_PARSER_NAME: parse_listed,
    APPLICANT_PARSER_NAME: parse_applicants,
}
