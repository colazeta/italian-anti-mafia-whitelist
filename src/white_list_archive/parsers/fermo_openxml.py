from __future__ import annotations

import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record


LISTED_PARSER_NAME = "fermo_listed"
APPLICANT_PARSER_NAME = "fermo_applicants"
PARSER_VERSION = "1"

_EXPECTED_LISTED_ROWS = 174
_EXPECTED_APPLICANT_ROWS = 75
_EXPECTED_LISTED_ORDINALS = tuple(i for i in range(1, 177) if i not in {105, 121})
_EXPECTED_APPLICANT_ORDINALS = tuple(i for i in range(1, 77) if i != 38)
_EXPECTED_LISTED_STATUS_COUNTS = {"listed": 97, "renewal_update_in_progress": 77}
_EXPECTED_APPLICANT_STATUS_COUNTS = {"pending": 75}

_RENEWAL = "IN FASE DI RINNOVO (l'iscrizione resta valida anche oltre la scadenza, fino all'esito definitivo)"
_RENEWAL_CORPORATE = "IN FASE DI RINNOVO PER MODIFICA ASSETTO SOCIETARIO (l'iscrizione resta valida anche oltre la scadenza, fino all'esito definitivo)"
_EXPECTED_LISTED_NOTES = {"": 97, _RENEWAL: 76, _RENEWAL_CORPORATE: 1}

_EXPECTED_LISTED_RAW_ONLY = {
    ("A.M. EDILIZIA SNC", "016778200443"),
    ("AIRTECNOHOUSE DI SACHA CAPRIOTTI", "2542610445"),
    ("AUTOTRASPORTI MORETTI SRL", "0172190449"),
    ("AZIENDA AGRICOLA CURTI CLAUDIO", "023415440447"),
    ("COGNIGNI Bruno", "1062460447"),
    ("ECO EDILIZIA SRL", "0213601440"),
    ("FERMANA LOGISTICA SRL", "023974300444"),
    ("GEOSERVICE SRL", "0209650048"),
    ("PIERAGOSTINI EMANUELE", "023115360442"),
    ("RC SCAVI E COSTRUZIONI", "022386400440"),
    ("STEVI SRL", "0244732447"),
}
_EXPECTED_APPLICANT_RAW_ONLY = {
    ("BOSCO VITTORIO", "1579880442"),
    ("EDIL NORD DI MOHAMED BOUTLATA", "2592610444"),
    ("GARDEN TALAMONTI S.R.L.", "1310970429"),
    ("PAKO SRL", "2499190441"),
    ("SIEM SRL", "0184220446"),
}
_EXPECTED_MALFORMED_LISTING_DATES = {
    ("MOVIMENTO TERRA DI SQUARCIA FRANCESCO", "28/11/204"),
}
_EXPECTED_MALFORMED_APPLICATION_DATES = {
    ("EDIL NORD DI MOHAMED BOUTLATA", "46092"),
}

_STRICT_11 = re.compile(r"^\d{11}$")
_STRICT_16 = re.compile(r"^[A-Z]{6}[0-9A-Z]{10}$")
_ROMAN = re.compile(r"^(?:I|II|III|IV|V|VI|VII|VIII|IX|X)$")


def _source_date(value: Any) -> str:
    """Normalise only dates explicitly represented as dates or complete source strings."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _clean(value)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?: 00:00:00)?", raw):
        try:
            return date.fromisoformat(raw[:10]).isoformat()
        except ValueError:
            return ""
    match = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", raw)
    if not match:
        return ""
    day, month, year = map(int, match.groups())
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return ""


def _structured_identifier(raw: str) -> bool:
    token = _clean(raw).upper().replace(" ", "")
    return bool(_STRICT_11.fullmatch(token) or _STRICT_16.fullmatch(token))


def _sections(raw: str) -> list[str]:
    source = _clean(raw)
    if not source:
        raise RuntimeError("Fermo activity field unexpectedly blank")
    folded = re.sub(r"(?i)^sez\.?\s*", "", source).upper()
    parts = [item for item in re.split(r"\s*-\s*", folded) if item]
    if not parts or any(not _ROMAN.fullmatch(item) for item in parts):
        raise RuntimeError(f"Fermo activity vocabulary changed: {source!r}")
    return [f"Sez. {item}" for item in parts]


def _listed_status(note: str) -> str:
    note = _clean(note)
    if not note:
        return "listed"
    if note in {_RENEWAL, _RENEWAL_CORPORATE}:
        return "renewal_update_in_progress"
    raise RuntimeError(f"Fermo listed note/status vocabulary changed: {note!r}")


def _workbook_rows(path: Path, *, population: str) -> tuple[list[tuple[int, list[Any]]], tuple[int, int]]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if workbook.sheetnames != ["Foglio1"]:
        raise RuntimeError(f"Fermo {population} workbook sheet set changed: {workbook.sheetnames!r}")
    worksheet = workbook["Foglio1"]
    expected_shape = (400, 12) if population == "listed" else (267, 15)
    observed_shape = (worksheet.max_row, worksheet.max_column)
    if observed_shape != expected_shape:
        raise RuntimeError(
            f"Fermo {population} workbook shape changed: expected {expected_shape!r}, got {observed_shape!r}"
        )
    rows = [(row_number, list(values)) for row_number, values in enumerate(worksheet.iter_rows(values_only=True), 1)]
    return rows, observed_shape


def _validate_title_and_header(rows: list[tuple[int, list[Any]]], *, population: str) -> None:
    by_number = {number: values for number, values in rows}
    title = _clean((by_number.get(3) or [""])[0]).casefold()
    if population == "listed":
        required_title = "elenco delle imprese iscritte nell'elenco dei fornitori"
        expected_header = ["n.", "ragione sociale", "sede legale", "c.f./p.i.", "data iscrizione", "data scadenza", "settori di attivita'", "note"]
    else:
        required_title = "elenco delle imprese richiedenti l'iscrizione nell'elenco dei fornitori"
        expected_header = ["n.", "ragione sociale", "sede legale", "c.f./p.i.", "data presentazione", "settori di attivita'", "note"]
    if required_title not in title:
        raise RuntimeError(f"Fermo {population} title changed: {title!r}")
    header = [_clean(value).casefold() for value in (by_number.get(9) or [])[: len(expected_header)]]
    if header != expected_header:
        raise RuntimeError(f"Fermo {population} header changed: expected {expected_header!r}, got {header!r}")


def _extract(rows: list[tuple[int, list[Any]]], *, population: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    expected_width = 8 if population == "listed" else 7
    for source_row, values in rows:
        if source_row < 10:
            continue
        padded = list(values) + [None] * max(0, expected_width - len(values))
        ordinal_raw = _clean(padded[0])
        name = _clean(padded[1])
        if not ordinal_raw and not name:
            continue
        if not ordinal_raw.isdigit():
            if any(_clean(value) for value in padded[:expected_width]):
                raise RuntimeError(f"Fermo {population} unexpected non-data row {source_row}: {padded[:expected_width]!r}")
            continue
        # The source workbooks contain styled trailing ordinal-only rows. They are
        # not company observations. Keep the reviewed distinction explicit and
        # fail closed if any identity-bearing field unexpectedly appears there.
        if not name:
            if any(_clean(value) for value in padded[2:expected_width]):
                raise RuntimeError(f"Fermo {population} ordinal-only row gained content at {source_row}: {padded[:expected_width]!r}")
            continue
        ordinal = int(ordinal_raw)
        office = _clean(padded[2])
        identifier_raw = _clean(padded[3])
        if population == "listed":
            listing_raw = _clean(padded[4])
            expiry_raw = _clean(padded[5])
            activity_raw = _clean(padded[6])
            note = _clean(padded[7])
            records.append(
                {
                    "source_row": source_row,
                    "source_ordinal": ordinal,
                    "name": name,
                    "office": office,
                    "identifier_raw": identifier_raw,
                    "listing_raw": listing_raw,
                    "listing_date": _source_date(padded[4]),
                    "expiry_raw": expiry_raw,
                    "expiry_date": _source_date(padded[5]),
                    "activity_raw": activity_raw,
                    "activities": _sections(activity_raw),
                    "note": note,
                }
            )
        else:
            application_raw = _clean(padded[4])
            activity_raw = _clean(padded[5])
            note = _clean(padded[6])
            if note:
                raise RuntimeError(f"Fermo applicant note/outcome requires review: {note!r}")
            records.append(
                {
                    "source_row": source_row,
                    "source_ordinal": ordinal,
                    "name": name,
                    "office": office,
                    "identifier_raw": identifier_raw,
                    "application_raw": application_raw,
                    "application_date": _source_date(padded[4]),
                    "activity_raw": activity_raw,
                    "activities": _sections(activity_raw),
                    "note": note,
                }
            )
    return records


def parse_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    rows, shape = _workbook_rows(path, population="listed")
    _validate_title_and_header(rows, population="listed")
    extracted = _extract(rows, population="listed")

    ordinals = tuple(record["source_ordinal"] for record in extracted)
    if ordinals != _EXPECTED_LISTED_ORDINALS:
        raise RuntimeError(f"Fermo listed source-ordinal boundary changed: {ordinals!r}")
    if len(extracted) != _EXPECTED_LISTED_ROWS:
        raise RuntimeError(f"Fermo listed denominator changed: expected {_EXPECTED_LISTED_ROWS}, got {len(extracted)}")
    if any(not record["name"] or not record["identifier_raw"] or not record["expiry_date"] for record in extracted):
        raise RuntimeError("Fermo listed identity/identifier/expiry boundary contains an unexpected blank or malformed value")

    notes = Counter(record["note"] for record in extracted)
    if dict(notes) != _EXPECTED_LISTED_NOTES:
        raise RuntimeError(f"Fermo listed note vocabulary/counts changed: {dict(notes)!r}")

    raw_only = {
        (record["name"], record["identifier_raw"])
        for record in extracted
        if not _structured_identifier(record["identifier_raw"])
    }
    if raw_only != _EXPECTED_LISTED_RAW_ONLY:
        raise RuntimeError(f"Fermo listed reviewed raw-only identifiers changed: {sorted(raw_only)!r}")

    malformed_listing = {
        (record["name"], record["listing_raw"])
        for record in extracted
        if not record["listing_date"]
    }
    if malformed_listing != _EXPECTED_MALFORMED_LISTING_DATES:
        raise RuntimeError(f"Fermo listed malformed listing-date boundary changed: {sorted(malformed_listing)!r}")

    output: list[dict[str, Any]] = []
    for record in extracted:
        output.append(
            _record(
                cfg,
                len(output) + 1,
                name=record["name"],
                office=record["office"],
                identifier_raw=record["identifier_raw"],
                activities=record["activities"],
                status=_listed_status(record["note"]),
                outcome_raw=record["note"],
                listing_date=record["listing_date"],
                expiry_date=record["expiry_date"],
                primary_date_label="Data iscrizione",
                source_fields={
                    "sections": list(record["activities"]),
                    "listing_date_raw_variants": [record["listing_raw"]] if record["listing_raw"] else [],
                    "expiry_date_raw_variants": [record["expiry_raw"]] if record["expiry_raw"] else [],
                    "notes": [record["note"]] if record["note"] else [],
                    "malformed_date_pairs": [f"listing_date={record['listing_raw']}"] if not record["listing_date"] else [],
                },
            )
        )

    statuses = Counter(record["source_status"] for record in output)
    if dict(statuses) != _EXPECTED_LISTED_STATUS_COUNTS:
        raise RuntimeError(f"Fermo listed status boundary changed: {dict(statuses)!r}")
    identifier_coverage = sum(bool(record["identifiers"]) for record in output)
    if identifier_coverage != 163:
        raise RuntimeError(f"Fermo listed identifier coverage changed: {identifier_coverage}")

    return ParsedBatch(
        output,
        {
            "parser": LISTED_PARSER_NAME,
            "parser_version": PARSER_VERSION,
            "workbook_shape": shape,
            "source_rows": len(extracted),
            "public_records": len(output),
            "source_ordinal_gaps": [105, 121],
            "status_counts": dict(statuses),
            "identifier_coverage": identifier_coverage,
            "raw_identifier_only": len(_EXPECTED_LISTED_RAW_ONLY),
            "malformed_listing_dates_preserved": len(_EXPECTED_MALFORMED_LISTING_DATES),
        },
    )


def parse_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    rows, shape = _workbook_rows(path, population="applicant")
    _validate_title_and_header(rows, population="applicant")
    extracted = _extract(rows, population="applicant")

    ordinals = tuple(record["source_ordinal"] for record in extracted)
    if ordinals != _EXPECTED_APPLICANT_ORDINALS:
        raise RuntimeError(f"Fermo applicant source-ordinal boundary changed: {ordinals!r}")
    if len(extracted) != _EXPECTED_APPLICANT_ROWS:
        raise RuntimeError(f"Fermo applicant denominator changed: expected {_EXPECTED_APPLICANT_ROWS}, got {len(extracted)}")
    if any(not record["name"] or not record["identifier_raw"] for record in extracted):
        raise RuntimeError("Fermo applicant identity/identifier boundary contains an unexpected blank")

    raw_only = {
        (record["name"], record["identifier_raw"])
        for record in extracted
        if not _structured_identifier(record["identifier_raw"])
    }
    if raw_only != _EXPECTED_APPLICANT_RAW_ONLY:
        raise RuntimeError(f"Fermo applicant reviewed raw-only identifiers changed: {sorted(raw_only)!r}")

    malformed_application = {
        (record["name"], record["application_raw"])
        for record in extracted
        if not record["application_date"]
    }
    if malformed_application != _EXPECTED_MALFORMED_APPLICATION_DATES:
        raise RuntimeError(
            f"Fermo applicant malformed application-date boundary changed: {sorted(malformed_application)!r}"
        )

    output: list[dict[str, Any]] = []
    for record in extracted:
        output.append(
            _record(
                cfg,
                len(output) + 1,
                name=record["name"],
                office=record["office"],
                identifier_raw=record["identifier_raw"],
                activities=record["activities"],
                status="pending",
                outcome_raw="",
                application_date=record["application_date"],
                primary_date_label="Data presentazione istanza",
                source_fields={
                    "requested_activities_source": record["activity_raw"],
                    "application_date_raw_variants": [record["application_raw"]] if record["application_raw"] else [],
                    "malformed_date_pairs": [f"application_date={record['application_raw']}"] if not record["application_date"] else [],
                },
            )
        )

    statuses = Counter(record["source_status"] for record in output)
    if dict(statuses) != _EXPECTED_APPLICANT_STATUS_COUNTS:
        raise RuntimeError(f"Fermo applicant status boundary changed: {dict(statuses)!r}")
    identifier_coverage = sum(bool(record["identifiers"]) for record in output)
    if identifier_coverage != 70:
        raise RuntimeError(f"Fermo applicant identifier coverage changed: {identifier_coverage}")

    return ParsedBatch(
        output,
        {
            "parser": APPLICANT_PARSER_NAME,
            "parser_version": PARSER_VERSION,
            "workbook_shape": shape,
            "source_rows": len(extracted),
            "public_records": len(output),
            "source_ordinal_gaps": [38],
            "status_counts": dict(statuses),
            "identifier_coverage": identifier_coverage,
            "raw_identifier_only": len(_EXPECTED_APPLICANT_RAW_ONLY),
            "malformed_application_dates_preserved": len(_EXPECTED_MALFORMED_APPLICATION_DATES),
        },
    )


PARSERS = {
    LISTED_PARSER_NAME: parse_listed,
    APPLICANT_PARSER_NAME: parse_applicants,
}
