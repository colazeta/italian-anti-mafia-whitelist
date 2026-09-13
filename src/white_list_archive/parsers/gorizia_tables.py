from __future__ import annotations

import hashlib
import re
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record

PARSER_VERSION = "1"
_LISTED_SHA256 = "680f0822f56ce7499a5d3f6c9e524d3d33b631feca51886e57e70fdf9e5fbd6d"
_APPLICANT_SHA256 = "befe7ed054e0cfe9dda48c8dec59f843b569019e40b925d6d5d1fcec19021c7c"
_REFERENCE_DATE = "2026-09-13"

_EXPECTED_LISTED_SECTOR_ROWS = 186
_EXPECTED_LISTED_REGISTRATIONS = 117
_EXPECTED_LISTED_STATUS_COUNTS = {"listed": 86, "renewal_update_in_progress": 31}
_EXPECTED_SECTION_ROWS = {1: 26, 2: 13, 3: 25, 4: 14, 5: 32, 6: 38, 7: 0, 8: 1, 9: 10, 10: 27}
_EXPECTED_UPDATE_VALUES = {
    "": 138,
    "IN AGGIORNAMENTO": 41,
    "21 gennaio 2026": 1,
    "25 settembre 2026": 3,
    "30 luglio 2027": 1,
    "“": 2,
}
_EXPECTED_HEADERS = {
    "Ragione Sociale | Sede legale | Sede secondaria con rappresentanza stabile in Italia | Codice fiscale/Partita IVA | Data d’iscrizione | Data scadenza iscrizione | Aggiornamento in corso",
    "Ragione Sociale | Sede legale | Sede secondaria con rappresentanza stabile in Italia | Codice fiscale/Partita IVA | Data d’ iscrizione | Data scadenza iscrizione | Aggiornamento in corso",
    "Ragione Sociale | Sede legale | Sede secondaria con rappresentanza stabile in Italia | Codice fiscale/Partita IVA | Data d ’ iscrizione | Data scadenza iscrizione | Aggiornamento in corso",
    "Ragione Sociale | Sede legale | Sede secondaria con rappresentanza stabile in Italia | Codice fiscale/Partita IVA | Data di iscrizione | Data scadenza iscrizione | Aggiornamento in corso",
}

_EXPECTED_APPLICANT_IDENTIFIERS = [
    "01270500315",
    "01262160318",
    "01190800316",
    "01270740317",
    "00407990316",
    "01268270319",
    "00557360310",
    "01040110312",
    "01170510315",
    "01195820319",
]
_EXPECTED_APPLICANT_NAMES = [
    "ITALTRCH SRL",
    "METAL X SRL",
    "“L’ANTICA RICETTA SRLS”",
    "EL.NET SOLUTION SRL",
    "C.M.T. SRL",
    "FMGDUE SRL",
    "SI.ECO:SICUREZZA ED ECOLOGIA SRL",
    "SULTAN SRL",
    "SVILUPPO SOLARE SRL",
    "T-RECYCLE SRL",
]
# Row geometry in the approved PDF binds 14.04.2026 to FMGDUE SRL;
# the following SI.ECO applicant row has no application date. Keep this sequence fail-closed.
_EXPECTED_APPLICANT_DATES = [
    "",
    "24.06.2025",
    "21.04.2026",
    "17.02.2026",
    "17.02.2026",
    "14.04.2026",
    "",
    "",
    "01.04.2026",
    "",
]

_ITALIAN_MONTHS = {
    "gennaio": 1,
    "febbraio": 2,
    "marzo": 3,
    "aprile": 4,
    "maggio": 5,
    "giugno": 6,
    "luglio": 7,
    "agosto": 8,
    "settembre": 9,
    "ottobre": 10,
    "novembre": 11,
    "dicembre": 12,
}
_REVIEWED_MALFORMED_DATES = frozenset({"10.12.202", "Dal 10.12.202", "14 agosto 204"})
_REVIEWED_SPLIT_DIGIT_DATES = {
    "2 1 aprile 2026": "2026-04-21",
    "14 luglio 2 027": "2027-07-14",
    "21 ottobre 202 3": "2023-10-21",
}
_REVIEWED_BLANK_EXPIRY_ROWS = frozenset(
    {
        (
            "“ MAROLLI COSTRUZIONI SRL”",
            "MONFALCONE (GO) Viale San Marco, 13/B",
            "C.F./P.I. 01218760310",
            "29 dicembre 2022",
        ),
        (
            "PEVERE LOGISTICA SRL",
            "MONFALCONE (GO) Via Timavo, 63",
            "C.F./P.I. 01282120315",
            "30 luglio 2026",
        ),
    }
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _validate_cfg(cfg: dict[str, Any], *, source_key: str, population_scope: str, sha256: str) -> None:
    if cfg.get("source_key") != source_key:
        raise RuntimeError(f"Gorizia source-key drift: {cfg.get('source_key')!r} != {source_key!r}")
    if cfg.get("authority_key") != "gorizia":
        raise RuntimeError("Gorizia parser bound to a non-Gorizia authority")
    if cfg.get("population_scope") != population_scope:
        raise RuntimeError(
            f"Gorizia population-scope drift for {source_key}: {cfg.get('population_scope')!r} != {population_scope!r}"
        )
    if cfg.get("sha256") != sha256:
        raise RuntimeError(f"Gorizia configured SHA-256 drift for {source_key}: {cfg.get('sha256')!r} != {sha256!r}")
    if cfg.get("reference_date") != _REFERENCE_DATE:
        raise RuntimeError(
            f"Gorizia reference-date drift for {source_key}: {cfg.get('reference_date')!r} != {_REFERENCE_DATE!r}"
        )


def _normalise_italian_date(raw: str) -> str:
    value = _clean(raw)
    if not value:
        return ""
    if value in _REVIEWED_MALFORMED_DATES:
        return ""
    if value in _REVIEWED_SPLIT_DIGIT_DATES:
        return _REVIEWED_SPLIT_DIGIT_DATES[value]
    match = re.fullmatch(r"(\d{1,2})(?:°)?\s+([A-Za-zÀ-ÿ]+)\s+(\d{4})", value)
    if not match:
        raise RuntimeError(f"Gorizia unexpected date lexeme: {value!r}")
    month = _ITALIAN_MONTHS.get(match.group(2).lower())
    if not month:
        raise RuntimeError(f"Gorizia unknown Italian month: {match.group(2)!r}")
    try:
        return date(int(match.group(3)), month, int(match.group(1))).isoformat()
    except ValueError as exc:
        raise RuntimeError(f"Gorizia invalid calendar date: {value!r}") from exc


def _docx_cell_text(cell: Any) -> str:
    values: list[str] = []
    for paragraph in cell.paragraphs:
        text = _clean(paragraph.text)
        if text:
            values.append(text)
    return _clean(" ".join(values))


def _header_signature(row: Any) -> str:
    return " | ".join(_docx_cell_text(cell) for cell in row.cells)


def _listed_sector_rows(path: Path) -> tuple[list[dict[str, Any]], dict[int, int]]:
    document = Document(path)
    if len(document.tables) != 10:
        raise RuntimeError(f"Gorizia listed publisher-shape drift: {len(document.tables)} tables != 10")

    rows: list[dict[str, Any]] = []
    section_counts: defaultdict[int, int] = defaultdict(int)
    for table_index, table in enumerate(document.tables, start=1):
        if not table.rows:
            raise RuntimeError(f"Gorizia listed empty table: {table_index}")
        header = _header_signature(table.rows[0])
        if header not in _EXPECTED_HEADERS:
            raise RuntimeError(f"Gorizia listed header drift in table {table_index}: {header!r}")
        for row_index, row in enumerate(table.rows[1:], start=2):
            cells = [_docx_cell_text(cell) for cell in row.cells]
            if not any(cells):
                continue
            if len(cells) != 7:
                raise RuntimeError(
                    f"Gorizia listed column drift in table {table_index} row {row_index}: {len(cells)} != 7"
                )
            name, office, secondary, identifier_raw, listing_raw, expiry_raw, update_raw = cells
            if not name:
                raise RuntimeError(f"Gorizia listed row without company name in table {table_index} row {row_index}")
            rows.append(
                {
                    "section": table_index,
                    "source_row": row_index,
                    "name": name,
                    "office": office,
                    "secondary": secondary,
                    "identifier_raw": identifier_raw,
                    "listing_raw": listing_raw,
                    "expiry_raw": expiry_raw,
                    "update_raw": update_raw,
                }
            )
            section_counts[table_index] += 1

    observed_sections = {section: section_counts.get(section, 0) for section in range(1, 11)}
    if len(rows) != _EXPECTED_LISTED_SECTOR_ROWS:
        raise RuntimeError(f"Gorizia listed sector-row drift: {len(rows)} != {_EXPECTED_LISTED_SECTOR_ROWS}")
    if observed_sections != _EXPECTED_SECTION_ROWS:
        raise RuntimeError(f"Gorizia listed section-count drift: {observed_sections!r}")
    update_values = dict(Counter(row["update_raw"] for row in rows))
    if update_values != _EXPECTED_UPDATE_VALUES:
        raise RuntimeError(f"Gorizia listed update-value drift: {update_values!r}")
    return rows, observed_sections


def parse_gorizia_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, source_key="gorizia-listed", population_scope="listed", sha256=_LISTED_SHA256)
    if _sha256(path) != _LISTED_SHA256:
        raise RuntimeError("Gorizia listed source bytes drift from approved SHA-256")

    sector_rows, section_counts = _listed_sector_rows(path)
    grouped: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    blank_expiry_rows = 0
    for row in sector_rows:
        if not row["expiry_raw"]:
            signature = (row["name"], row["office"], row["identifier_raw"], row["listing_raw"])
            if signature not in _REVIEWED_BLANK_EXPIRY_ROWS:
                raise RuntimeError(f"Gorizia unreviewed blank expiry row: {signature!r}")
            blank_expiry_rows += 1
        key = (
            row["name"],
            row["office"],
            row["secondary"],
            row["identifier_raw"],
            row["listing_raw"],
            row["expiry_raw"],
        )
        grouped[key].append(row)

    if blank_expiry_rows != len(_REVIEWED_BLANK_EXPIRY_ROWS):
        raise RuntimeError(
            f"Gorizia reviewed blank-expiry population drift: {blank_expiry_rows} != {len(_REVIEWED_BLANK_EXPIRY_ROWS)}"
        )
    if len(grouped) != _EXPECTED_LISTED_REGISTRATIONS:
        raise RuntimeError(
            f"Gorizia listed registration-group drift: {len(grouped)} != {_EXPECTED_LISTED_REGISTRATIONS}"
        )

    records: list[dict[str, Any]] = []
    grouped_source_rows = 0
    for key, members in grouped.items():
        name, office, secondary, identifier_raw, listing_raw, expiry_raw = key
        sections = sorted({member["section"] for member in members})
        update_values = sorted({member["update_raw"] for member in members})
        nonblank_updates = [value for value in update_values if value]
        source_status = "renewal_update_in_progress" if nonblank_updates else "listed"
        source_rows = [[member["section"], member["source_row"]] for member in members]
        grouped_source_rows += len(members)
        listing_date = _normalise_italian_date(listing_raw)
        expiry_date = _normalise_italian_date(expiry_raw)
        records.append(
            _record(
                cfg,
                len(records) + 1,
                name=name,
                office=office,
                secondary=secondary,
                identifier_raw=identifier_raw,
                activities=[f"Sezione {section}" for section in sections],
                status=source_status,
                outcome_raw=" | ".join(nonblank_updates),
                listing_date=listing_date,
                expiry_date=expiry_date,
                primary_date_label="Data iscrizione",
                source_fields={
                    "source_rows": source_rows,
                    "sections": sections,
                    "identifier_raw": identifier_raw,
                    "listing_date_raw": listing_raw,
                    "expiry_date_raw": expiry_raw,
                    "update_values": update_values,
                },
            )
        )

    if grouped_source_rows != _EXPECTED_LISTED_SECTOR_ROWS:
        raise RuntimeError(
            f"Gorizia grouped source-row coverage drift: {grouped_source_rows} != {_EXPECTED_LISTED_SECTOR_ROWS}"
        )
    status_counts = dict(Counter(record["source_status"] for record in records))
    if status_counts != _EXPECTED_LISTED_STATUS_COUNTS:
        raise RuntimeError(
            f"Gorizia listed status-count drift: {status_counts!r} != {_EXPECTED_LISTED_STATUS_COUNTS!r}"
        )
    if len({record["record_locator"] for record in records}) != len(records):
        raise RuntimeError("Gorizia listed record locator collision")

    return ParsedBatch(
        records,
        {
            "parser": "gorizia_listed",
            "parser_version": PARSER_VERSION,
            "sector_rows": len(sector_rows),
            "public_records": len(records),
            "registration_groups": len(records),
            "status_counts": status_counts,
            "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
            "blank_expiry_rows": blank_expiry_rows,
            "section_rows": section_counts,
            "update_values": dict(Counter(row["update_raw"] for row in sector_rows)),
        },
    )


def parse_gorizia_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    _validate_cfg(cfg, source_key="gorizia-applicants", population_scope="applicant", sha256=_APPLICANT_SHA256)
    if _sha256(path) != _APPLICANT_SHA256:
        raise RuntimeError("Gorizia applicant source bytes drift from approved SHA-256")

    with pdfplumber.open(path) as pdf:
        if len(pdf.pages) != 2:
            raise RuntimeError(f"Gorizia applicant page-count drift: {len(pdf.pages)} != 2")
        words: list[dict[str, Any]] = []
        cumulative = 0.0
        for page_number, page in enumerate(pdf.pages, start=1):
            for word in page.extract_words(use_text_flow=True, keep_blank_chars=False):
                words.append({**word, "page": page_number, "global_top": cumulative + float(word["top"])})
            cumulative += float(page.height)

    identifier_anchors: list[dict[str, Any]] = []
    identifier_occurrences: Counter[str] = Counter()
    for word in words:
        text = _clean(word["text"])
        matches = [identifier for identifier in _EXPECTED_APPLICANT_IDENTIFIERS if identifier in text]
        if len(matches) > 1:
            raise RuntimeError(f"Gorizia ambiguous applicant identifier word: {text!r}")
        if matches:
            identifier = matches[0]
            identifier_occurrences[identifier] += 1
            identifier_anchors.append({**word, "identifier": identifier})

    observed_identifiers = [item["identifier"] for item in identifier_anchors]
    if observed_identifiers != _EXPECTED_APPLICANT_IDENTIFIERS or any(
        identifier_occurrences[identifier] != 1 for identifier in _EXPECTED_APPLICANT_IDENTIFIERS
    ):
        raise RuntimeError(
            f"Gorizia applicant identifier-anchor drift: {observed_identifiers!r}; "
            f"occurrences={dict(identifier_occurrences)!r}"
        )

    rows: list[dict[str, Any]] = []
    for index, anchor in enumerate(identifier_anchors):
        current_top = float(anchor["global_top"])
        next_top = (
            float(identifier_anchors[index + 1]["global_top"])
            if index + 1 < len(identifier_anchors)
            else float("inf")
        )
        segment = [word for word in words if current_top - 4 <= float(word["global_top"]) < next_top - 4]
        segment_sorted = sorted(segment, key=lambda item: (float(item["global_top"]), float(item["x0"])))
        name_words = [item for item in segment_sorted if float(item["x0"]) < 225 and item is not anchor]
        name_candidate = _clean(" ".join(_clean(item["text"]) for item in name_words))
        expected_name = _EXPECTED_APPLICANT_NAMES[index]
        if name_candidate.count(expected_name) != 1:
            raise RuntimeError(
                f"Gorizia applicant expected-name anchor drift for {expected_name!r}: {name_candidate!r}"
            )
        date_candidates = [
            _clean(item["text"])
            for item in segment_sorted
            if re.fullmatch(r"\d{1,2}\.\d{1,2}\.\d{4}", _clean(item["text"]))
        ]
        if len(date_candidates) > 1:
            raise RuntimeError(
                f"Gorizia applicant date-anchor ambiguity for {expected_name!r}: {date_candidates!r}"
            )
        rows.append(
            {
                "identifier": observed_identifiers[index],
                "name": expected_name,
                "date_raw": date_candidates[0] if date_candidates else "",
                "page": int(anchor["page"]),
            }
        )

    names = [row["name"] for row in rows]
    dates = [row["date_raw"] for row in rows]
    if names != _EXPECTED_APPLICANT_NAMES:
        raise RuntimeError(f"Gorizia applicant name drift: {names!r}")
    if dates != _EXPECTED_APPLICANT_DATES:
        raise RuntimeError(f"Gorizia applicant application-date drift: {dates!r}")

    records: list[dict[str, Any]] = []
    for row in rows:
        records.append(
            _record(
                cfg,
                len(records) + 1,
                name=row["name"],
                identifier_raw=row["identifier"],
                status="pending",
                application_date=row["date_raw"],
                primary_date_label="Data presentazione istanza",
                source_fields={
                    "page": row["page"],
                    "identifier_raw": row["identifier"],
                    "application_date_raw": row["date_raw"],
                },
            )
        )

    if len(records) != len(_EXPECTED_APPLICANT_IDENTIFIERS):
        raise RuntimeError(f"Gorizia applicant record-count drift: {len(records)} != 10")
    if len({record["record_locator"] for record in records}) != len(records):
        raise RuntimeError("Gorizia applicant record locator collision")

    return ParsedBatch(
        records,
        {
            "parser": "gorizia_applicants",
            "parser_version": PARSER_VERSION,
            "public_records": len(records),
            "status_counts": {"pending": len(records)},
            "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
        },
    )
