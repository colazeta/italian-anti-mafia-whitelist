from __future__ import annotations

import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

import openpyxl

from white_list_archive.parsers.multi_prefecture_tables import ParsedBatch, _clean, _record


LISTED_PARSER_NAME = "matera_listed"
APPLICANT_PARSER_NAME = "matera_applicants"
PARSER_VERSION = "1"

_EXPECTED_LISTED_SHEET = "DPP1059056-20260722-WLIscrizion"
_EXPECTED_APPLICANT_SHEET = "DPP1059056-20260731-WLIscrizion"
_EXPECTED_LISTED_SHAPE = (249, 10)
_EXPECTED_APPLICANT_SHAPE = (91, 7)
_EXPECTED_LISTED_ROWS = 248
_EXPECTED_APPLICANT_ROWS = 90
_EXPECTED_LISTED_STATUSES = {"ISCRITTA": 248}
_EXPECTED_APPLICANT_STATUSES = {"RICHIEDENTE_ISCRIZIONE": 41, "IN_AGGIORNAMENTO": 49}

_LISTED_HEADER = (
    "Prefettura Competente",
    "Codice fiscale",
    "Ragione sociale",
    "Stato iscrizione",
    "Data inizo iscrizione",
    "Data scadenza",
    "Note",
)
_APPLICANT_HEADER = (
    "Prefettura Competente",
    "Data presentazione istanza",
    "Codice fiscale",
    "Ragione sociale",
    "Settori di iscrizione",
    "Stato iscrizione",
    "Protocollo richiesta",
)
_FORMULA_LITERAL = re.compile(r'^="([^"]*)"$')
_SECTION = re.compile(r"^SEZ_(?:I|II|III|IV|V|VI|VII|VIII|IX|X)$")


def _source_literal(value: Any) -> str:
    """Decode only Excel's exact =\"literal\" wrapper; preserve all other text."""
    if value is None:
        return ""
    if isinstance(value, str):
        match = _FORMULA_LITERAL.fullmatch(value)
        if match:
            return _clean(match.group(1))
    return _clean(value)


def _source_date(value: Any) -> str:
    """Normalise only explicit complete dates; never repair missing or malformed digits."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _source_literal(value)
    match = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", raw)
    if match is None:
        return ""
    day, month, year = map(int, match.groups())
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return ""


def _sections(value: Any) -> list[str]:
    raw = _source_literal(value)
    if not raw:
        return []
    parts = [_clean(item).upper() for item in raw.split(",") if _clean(item)]
    if not parts or any(_SECTION.fullmatch(item) is None for item in parts):
        raise RuntimeError(f"Matera applicant section vocabulary changed: {raw!r}")
    if len(parts) != len(set(parts)):
        raise RuntimeError(f"Matera applicant section cell contains duplicates: {raw!r}")
    return parts


def _load(path: Path, *, population: str) -> tuple[list[list[Any]], tuple[int, int]]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    expected_sheet = _EXPECTED_LISTED_SHEET if population == "listed" else _EXPECTED_APPLICANT_SHEET
    expected_shape = _EXPECTED_LISTED_SHAPE if population == "listed" else _EXPECTED_APPLICANT_SHAPE
    if workbook.sheetnames != [expected_sheet]:
        raise RuntimeError(
            f"Matera {population} workbook sheet set changed: expected {[expected_sheet]!r}, got {workbook.sheetnames!r}"
        )
    worksheet = workbook[expected_sheet]
    shape = (worksheet.max_row, worksheet.max_column)
    if shape != expected_shape:
        raise RuntimeError(f"Matera {population} workbook shape changed: expected {expected_shape!r}, got {shape!r}")
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    return rows, shape


def _validate_header(rows: list[list[Any]], *, population: str) -> None:
    expected = _LISTED_HEADER if population == "listed" else _APPLICANT_HEADER
    observed = tuple(_source_literal(value) for value in rows[0][: len(expected)])
    if observed != expected:
        raise RuntimeError(f"Matera {population} header changed: expected {expected!r}, got {observed!r}")


def parse_listed(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    rows, shape = _load(path, population="listed")
    _validate_header(rows, population="listed")
    source_rows = rows[1:]
    if len(source_rows) != _EXPECTED_LISTED_ROWS:
        raise RuntimeError(f"Matera listed denominator changed: expected {_EXPECTED_LISTED_ROWS}, got {len(source_rows)}")

    records: list[dict[str, Any]] = []
    source_statuses: Counter[str] = Counter()
    judicial_control_notes = 0
    for source_row_number, row in enumerate(source_rows, start=2):
        values = list(row) + [None] * max(0, 10 - len(row))
        authority = _source_literal(values[0])
        identifier = _source_literal(values[1])
        name = _source_literal(values[2])
        source_status = _source_literal(values[3])
        listing_raw = _source_literal(values[4])
        expiry_raw = _source_literal(values[5])
        note_values = [_source_literal(value) for value in values[6:10] if _source_literal(value)]
        unique_notes = list(dict.fromkeys(note_values))
        if authority != "MT":
            raise RuntimeError(f"Matera listed authority changed at row {source_row_number}: {authority!r}")
        if not identifier or not name:
            raise RuntimeError(f"Matera listed identity field blank at row {source_row_number}")
        if source_status != "ISCRITTA":
            raise RuntimeError(f"Matera listed status vocabulary changed at row {source_row_number}: {source_status!r}")
        listing_date = _source_date(values[4])
        expiry_date = _source_date(values[5])
        if not listing_date or not expiry_date:
            raise RuntimeError(
                f"Matera listed date became blank/malformed at row {source_row_number}: {listing_raw!r}, {expiry_raw!r}"
            )
        if len(unique_notes) > 1:
            raise RuntimeError(f"Matera listed repeated note columns disagree at row {source_row_number}: {unique_notes!r}")
        note = unique_notes[0] if unique_notes else ""
        if note:
            judicial_control_notes += int("controllo giudiziario" in note.casefold())
        source_statuses[source_status] += 1
        records.append(
            _record(
                cfg,
                len(records) + 1,
                name=name,
                identifier_raw=identifier,
                status="listed",
                outcome_raw=source_status,
                listing_date=listing_date,
                expiry_date=expiry_date,
                primary_date_label="Data iscrizione",
                source_fields={
                    "source_row": source_row_number,
                    "source_authority_code": authority,
                    "source_status_raw": source_status,
                    "listing_date_raw_variants": [listing_raw],
                    "expiry_date_raw_variants": [expiry_raw],
                    "notes": [note] if note else [],
                    "note_cells_raw": note_values,
                },
            )
        )

    if dict(source_statuses) != _EXPECTED_LISTED_STATUSES:
        raise RuntimeError(f"Matera listed status counts changed: {dict(source_statuses)!r}")
    if len(records) != _EXPECTED_LISTED_ROWS:
        raise RuntimeError(f"Matera listed public denominator changed: {len(records)}")
    statuses = Counter(record["source_status"] for record in records)
    return ParsedBatch(
        records,
        {
            "parser": LISTED_PARSER_NAME,
            "parser_version": PARSER_VERSION,
            "workbook_shape": shape,
            "source_rows": len(source_rows),
            "public_records": len(records),
            "source_status_counts": dict(source_statuses),
            "status_counts": dict(statuses),
            "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
            "raw_identifier_only": sum(bool(record["identifier_field_raw"]) and not record["identifiers"] for record in records),
            "listing_date_coverage": sum(bool(record["observed_listing_date"]) for record in records),
            "expiry_date_coverage": sum(bool(record["observed_expiry_date"]) for record in records),
            "rows_with_judicial_control_note": judicial_control_notes,
        },
    )


def parse_applicants(path: Path, cfg: dict[str, Any]) -> ParsedBatch:
    rows, shape = _load(path, population="applicant")
    _validate_header(rows, population="applicant")
    source_rows = rows[1:]
    if len(source_rows) != _EXPECTED_APPLICANT_ROWS:
        raise RuntimeError(f"Matera applicant denominator changed: expected {_EXPECTED_APPLICANT_ROWS}, got {len(source_rows)}")

    records: list[dict[str, Any]] = []
    source_statuses: Counter[str] = Counter()
    true_excel_dates = 0
    for source_row_number, row in enumerate(source_rows, start=2):
        values = list(row) + [None] * max(0, 7 - len(row))
        authority = _source_literal(values[0])
        application_raw = _source_literal(values[1])
        identifier = _source_literal(values[2])
        name = _source_literal(values[3])
        sections_raw = _source_literal(values[4])
        source_status = _source_literal(values[5])
        protocol = _source_literal(values[6])
        if authority != "MT":
            raise RuntimeError(f"Matera applicant authority changed at row {source_row_number}: {authority!r}")
        if not identifier or not name:
            raise RuntimeError(f"Matera applicant identity field blank at row {source_row_number}")
        if source_status not in _EXPECTED_APPLICANT_STATUSES:
            raise RuntimeError(f"Matera applicant status vocabulary changed at row {source_row_number}: {source_status!r}")
        application_date = _source_date(values[1])
        if not application_date:
            raise RuntimeError(f"Matera applicant date became blank/malformed at row {source_row_number}: {application_raw!r}")
        if isinstance(values[1], (date, datetime)):
            true_excel_dates += 1
        canonical_status = "pending" if source_status == "RICHIEDENTE_ISCRIZIONE" else "renewal_update_in_progress"
        source_statuses[source_status] += 1
        records.append(
            _record(
                cfg,
                len(records) + 1,
                name=name,
                identifier_raw=identifier,
                activities=_sections(values[4]),
                status=canonical_status,
                outcome_raw=source_status,
                application_date=application_date,
                primary_date_label="Data presentazione istanza",
                source_fields={
                    "source_row": source_row_number,
                    "source_authority_code": authority,
                    "source_status_raw": source_status,
                    "application_date_raw_variants": [application_raw],
                    "requested_sections_raw": sections_raw,
                    "protocol_request": protocol,
                    "application_date_cell_type": "excel_date" if isinstance(values[1], (date, datetime)) else "formula_literal",
                },
            )
        )

    if dict(source_statuses) != _EXPECTED_APPLICANT_STATUSES:
        raise RuntimeError(f"Matera applicant status counts changed: {dict(source_statuses)!r}")
    if len(records) != _EXPECTED_APPLICANT_ROWS:
        raise RuntimeError(f"Matera applicant public denominator changed: {len(records)}")
    statuses = Counter(record["source_status"] for record in records)
    return ParsedBatch(
        records,
        {
            "parser": APPLICANT_PARSER_NAME,
            "parser_version": PARSER_VERSION,
            "workbook_shape": shape,
            "source_rows": len(source_rows),
            "public_records": len(records),
            "source_status_counts": dict(source_statuses),
            "status_counts": dict(statuses),
            "identifier_coverage": sum(bool(record["identifiers"]) for record in records),
            "raw_identifier_only": sum(bool(record["identifier_field_raw"]) and not record["identifiers"] for record in records),
            "application_date_coverage": sum(bool(record["application_date"]) for record in records),
            "true_excel_date_cells": true_excel_dates,
        },
    )


PARSERS: dict[str, Callable[[Path, dict[str, Any]], ParsedBatch]] = {
    LISTED_PARSER_NAME: parse_listed,
    APPLICANT_PARSER_NAME: parse_applicants,
}
